"""Requests to product managers: what the database lacks, one Excel file per manager.

After a partner file is filled, the cells nobody could fill (a feature value the
database does not have, ТН ВЭД, ОКПД2, a barcode) and the "is this the same
characteristic?" questions become a request. The product's series decides the
manager, from the company's assignment file «Закрепление ПМ.xlsx». Each request
is an Excel file with a yellow "Ответ" column (a dropdown where the partner has
a closed list) and an email draft with the file attached.

The manager sends the file back; the program recognizes it by a hidden mark,
checks every answer and keeps the good ones in the work folder. Next time the
partner file is filled, the answers are used like database values. The program
remembers what it asked, so an open question is not asked again.
"""
from __future__ import annotations

import datetime as dt
import json
import re
import uuid
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from email.message import EmailMessage
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .atomic_files import atomic_write_json
from .filler import normalize_article
from .value_matching import LOGICAL_NO, LOGICAL_YES, Characteristic, canonical, match_value

ASSIGNMENT_RE = re.compile(r"закреплени\w*\s*пм", re.IGNORECASE)
CONTACTS_NAME = "Почты продактов.xlsx"
MARK = "CHINT-MDM-REQUEST"
META_SHEET = "_chint"
LISTS_SHEET = "Списки"
REQUEST_SHEET = "Запрос"
HEADER_ROW = 5
COLUMNS = ["Артикул", "Наименование", "Серия", "Что нужно", "Как заполнять", "Ответ", "Комментарий", "Ключ"]
ARTICLE, NAME, SERIES, LABEL, HINT, ANSWER, COMMENT, KEY = range(1, len(COLUMNS) + 1)
NO_MANAGER = "Без продакта"
REASK_DAYS = 30
NAME_QUESTION = "name:"
YELLOW = PatternFill("solid", fgColor="FFF2CC")
KIND_HINTS = {
    "A": "выберите из списка",
    "N": "число",
    "L": "Да или Нет",
    "R": "от и до, например 100-240",
}
PLAIN_CHECKS = {
    canonical("Код ТН ВЭД"): (re.compile(r"^\d{10}$"), "10 цифр"),
    canonical("Код ОКПД2"): (re.compile(r"^\d{2}\.\d{2}\.\d{2}\.\d{3}$"), "вида 27.12.24.190"),
    canonical("Штрих-код EAN"): (re.compile(r"^\d{13}$"), "13 цифр"),
}


@dataclass(frozen=True)
class Need:
    """One thing to ask: a value for an article, or a yes/no about a pair of names (article empty)."""

    article: str
    name: str
    series: str
    key: str  # ETIM feature code, a column name, or "name:<code>:<attribute>"
    label: str
    kind: str = ""  # N, A, L, R as in RAEK; "" for plain columns
    allowed: tuple[str, ...] = ()


@dataclass(frozen=True)
class Assignment:
    series: str
    manager: str
    line: str
    note: str


@dataclass
class RequestFile:
    manager: str
    path: Path
    items: int
    emails: list[str]
    draft: Path | None = None


@dataclass
class RequestsResult:
    files: list[RequestFile] = field(default_factory=list)
    skipped_open: int = 0
    unassigned_series: Counter = field(default_factory=Counter)
    missing_emails: list[str] = field(default_factory=list)
    assignment: Path | None = None


@dataclass
class AnswersResult:
    manager: str = ""
    accepted: int = 0
    names_kept: int = 0
    names_rejected: int = 0
    rejected: list[tuple[str, str, str, str]] = field(default_factory=list)  # article, what, answer, why
    empty: int = 0


# --- who answers for which series ----------------------------------------------------------------


def series_key(name: str) -> str:
    return re.sub(r"[^0-9A-ZА-Я]", "", str(name or "").upper())


def latest_assignment_file(folders: list[Path]) -> Path | None:
    found = [path for folder in folders if folder.is_dir() for path in folder.glob("*.xlsx") if ASSIGNMENT_RE.search(path.stem)]
    return max(found, key=lambda path: path.stat().st_mtime) if found else None


def load_assignment(path: Path) -> list[Assignment]:
    """Series → manager from «Закрепление ПМ.xlsx»: blocks of three columns (series | Ответственный ПМ | note).

    A row with a name and no manager is a product line heading for the rows below it.
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = [list(row) for row in workbook.worksheets[0].iter_rows(values_only=True)]
    finally:
        workbook.close()
    blocks = []
    for row in rows[:3]:
        for idx, value in enumerate(row):
            if canonical(str(value or "")) == canonical("Ответственный ПМ") and idx > 0 and idx - 1 not in blocks:
                blocks.append(idx - 1)
    result = []
    for start in blocks:
        line = ""
        for row in rows:
            cells = row + [None] * 3
            series, manager, note = (str(cells[start + n] or "").strip() for n in range(3))
            if not series or canonical(manager) == canonical("Ответственный ПМ"):
                continue
            if not manager:
                line = series
                continue
            result.append(Assignment(series, manager, line, note))
    return result


def manager_for(series: str, assignment: list[Assignment]) -> str:
    """Exact series first, then the longest assignment series the product's series starts with ("NB1L-20" → "NB1L")."""
    key = series_key(series)
    if not key:
        return ""
    best, best_len = "", 0
    for item in assignment:
        for variant in {series_key(item.series), series_key(item.series.split()[0]) if item.series.split() else ""}:
            if not variant:
                continue
            if variant == key:
                return item.manager
            if key.startswith(variant) and len(variant) > best_len:
                best, best_len = item.manager, len(variant)
    return best


def people(manager: str) -> list[str]:
    return [name.strip() for name in re.split(r"[,/;]", manager) if name.strip()]


def load_contacts(path: Path) -> dict[str, str]:
    if not path.exists():
        return {}
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return {
            canonical(str(row[0])): str(row[1]).strip()
            for row in workbook.worksheets[0].iter_rows(min_row=2, values_only=True)
            if row and row[0] and len(row) > 1 and row[1] and "@" in str(row[1])
        }
    finally:
        workbook.close()


def update_contacts(path: Path, assignment: list[Assignment]) -> dict[str, str]:
    """The emails table: one row per person of the assignment file; emails typed there are kept."""
    known = load_contacts(path)
    names = sorted({person for item in assignment for person in people(item.manager)})
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Почты"
    sheet.append(["Продакт", "Почта"])
    sheet["A1"].font = sheet["B1"].font = Font(bold=True)
    for name in names:
        sheet.append([name, known.get(canonical(name), "")])
        if not known.get(canonical(name)):
            sheet.cell(sheet.max_row, 2).fill = YELLOW
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 36
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return known


# --- the journal of what was asked ---------------------------------------------------------------


def journal_path(work_dir: Path) -> Path:
    return work_dir / "rules" / "pm_requests.json"


def answers_path(work_dir: Path) -> Path:
    return work_dir / "rules" / "pm_answers.json"


def read_json(path: Path, default: dict) -> dict:
    if not path.exists():
        return default
    return json.loads(path.read_text(encoding="utf-8"))


def open_keys(journal: dict, today: dt.date) -> set[tuple[str, str]]:
    """Questions asked less than REASK_DAYS ago and not answered yet."""
    keys = set()
    for request in journal.get("requests", []):
        if request.get("replaced"):
            continue
        created = dt.date.fromisoformat(request["created"][:10])
        if (today - created).days >= REASK_DAYS:
            continue
        answered = {tuple(item) for item in request.get("answered", [])}
        keys |= {tuple(item) for item in request.get("asked", [])} - answered
    return keys


def load_answers(path: Path | None) -> dict[tuple[str, str], str]:
    """Accepted answers by (article, key), for filling partner files."""
    if path is None or not path.exists():
        return {}
    data = read_json(path, {})
    return {(item["article"], item["key"]): item["value"] for item in data.get("answers", [])}


# --- writing requests ----------------------------------------------------------------------------


def write_requests(
    needs: list[Need],
    assignment: list[Assignment],
    work_dir: Path,
    output_dir: Path,
    reason: str,
    today: dt.date | None = None,
    again: bool = False,
) -> RequestsResult:
    """One request file per manager. With `again`, open questions are asked anew and the earlier files are
    marked replaced (answers sent in them are still accepted)."""
    today = today or dt.date.today()
    result = RequestsResult()
    journal_file = journal_path(work_dir)
    journal = read_json(journal_file, {"version": 1, "requests": []})
    already = set() if again else open_keys(journal, today)
    if again:
        for request in journal["requests"]:
            request.setdefault("replaced", today.isoformat())
    contacts = update_contacts(work_dir / CONTACTS_NAME, assignment) if assignment else load_contacts(work_dir / CONTACTS_NAME)
    by_manager: dict[str, list[Need]] = defaultdict(list)
    seen = set()
    for need in needs:
        key = (need.article, need.key)
        if key in seen:
            continue
        seen.add(key)
        if key in already:
            result.skipped_open += 1
            continue
        manager = manager_for(need.series, assignment)
        if not manager:
            result.unassigned_series[need.series or "серия не указана"] += 1
            manager = NO_MANAGER
        by_manager[manager].append(need)
    for manager, items in sorted(by_manager.items()):
        request_id = uuid.uuid4().hex[:12]
        stem = f"Запрос_{safe_name(manager)}_{today.isoformat()}"
        path = unique(output_dir / f"{stem}.xlsx")
        write_request_file(path, manager, items, request_id, reason, today)
        emails = [contacts[canonical(person)] for person in people(manager) if canonical(person) in contacts]
        missing = [person for person in people(manager) if canonical(person) not in contacts and manager != NO_MANAGER]
        result.missing_emails.extend(person for person in missing if person not in result.missing_emails)
        request = RequestFile(manager, path, len(items), emails)
        if manager != NO_MANAGER:
            request.draft = write_draft(path, emails, len(items), reason)
        result.files.append(request)
        journal["requests"].append(
            {
                "id": request_id,
                "manager": manager,
                "created": dt.datetime.now().isoformat(timespec="seconds"),
                "file": str(path),
                "reason": reason,
                "asked": [[need.article, need.key] for need in items],
                "answered": [],
            }
        )
    journal_file.parent.mkdir(parents=True, exist_ok=True)
    atomic_write_json(journal_file, journal)
    return result


def safe_name(text: str) -> str:
    return re.sub(r"[^\w]+", "_", text).strip("_")[:40] or "продакт"


def unique(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    candidate, version = path, 1
    while candidate.exists():
        version += 1
        candidate = path.with_name(f"{path.stem}_v{version}{path.suffix}")
    return candidate


def write_request_file(path: Path, manager: str, needs: list[Need], request_id: str, reason: str, today: dt.date) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = REQUEST_SHEET
    sheet["A1"] = "Запрос данных о товарах CHINT"
    sheet["A1"].font = Font(bold=True, size=13)
    sheet["A2"] = f"Для: {manager}. {reason}. {today:%d.%m.%Y}."
    sheet["A3"] = (
        "Заполните жёлтую колонку «Ответ» и отправьте файл обратно ответным письмом. "
        "Если значения нет или не знаете, оставьте клетку пустой или напишите в «Комментарий». Не меняйте другие колонки."
    )
    sheet.append([])
    sheet.append(COLUMNS)
    for cell in sheet[HEADER_ROW]:
        cell.font = Font(bold=True)
    lists = workbook.create_sheet(LISTS_SHEET)
    list_ranges: dict[tuple[str, ...], str] = {}
    for need in needs:
        hint = KIND_HINTS.get(need.kind, "")
        if need.key.startswith(NAME_QUESTION):
            hint = "Да или Нет"
        elif canonical(need.label) in PLAIN_CHECKS:
            hint = PLAIN_CHECKS[canonical(need.label)][1]
        sheet.append([need.article or "все товары серии", need.name, need.series, need.label, hint, None, None, need.key])
        row = sheet.max_row
        sheet.cell(row, ANSWER).fill = YELLOW
        choices = need.allowed if need.kind == "A" and need.allowed else ("Да", "Нет") if need.kind == "L" or need.key.startswith(NAME_QUESTION) else ()
        if choices:
            if choices not in list_ranges:
                col = len(list_ranges) + 1
                for idx, value in enumerate(choices, 1):
                    lists.cell(idx, col, value)
                letter = get_column_letter(col)
                list_ranges[choices] = f"'{LISTS_SHEET}'!${letter}$1:${letter}${len(choices)}"
            check = DataValidation(type="list", formula1=list_ranges[choices], allow_blank=True)
            check.error = "Выберите значение из списка"
            check.showErrorMessage = False  # a manager may know better than the list: the program checks the answer itself
            sheet.add_data_validation(check)
            check.add(sheet.cell(row, ANSWER))
    for col, width in zip(range(1, len(COLUMNS) + 1), (12, 60, 14, 50, 22, 28, 30, 12)):
        sheet.column_dimensions[get_column_letter(col)].width = width
    for row in sheet.iter_rows(min_row=HEADER_ROW + 1):
        row[NAME - 1].alignment = row[LABEL - 1].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.column_dimensions[get_column_letter(KEY)].hidden = True
    sheet.freeze_panes = sheet.cell(HEADER_ROW + 1, ANSWER)
    lists.sheet_state = "hidden"
    meta = workbook.create_sheet(META_SHEET)
    for value in (MARK, request_id, manager, today.isoformat()):
        meta.append([value])
    meta.sheet_state = "hidden"
    workbook.save(path)


def write_draft(attachment: Path, emails: list[str], items: int, reason: str) -> Path:
    """An .eml draft with the request attached: Outlook opens it ready to send, Thunderbird via «Изменить как новое»."""
    message = EmailMessage()
    message["To"] = ", ".join(emails)
    message["Subject"] = f"Запрос данных по товарам CHINT ({items} поз.)"
    message["X-Unsent"] = "1"
    message.set_content(
        "Добрый день!\n\n"
        f"{reason}. Для части товаров в нашей базе нет данных, список во вложении.\n"
        "Пожалуйста, заполните жёлтую колонку «Ответ» и пришлите файл обратно ответным письмом.\n"
        "Если значения нет, оставьте клетку пустой.\n\nСпасибо!"
    )
    message.add_attachment(
        attachment.read_bytes(),
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=attachment.name,
    )
    path = attachment.with_suffix(".eml")
    path.write_bytes(bytes(message))
    return path


# --- reading answers -----------------------------------------------------------------------------


def is_request_file(path: Path) -> bool:
    try:
        workbook = load_workbook(path, read_only=True)
    except Exception:
        return False
    try:
        return META_SHEET in workbook.sheetnames and workbook[META_SHEET]["A1"].value == MARK
    finally:
        workbook.close()


def checked_answer(label: str, kind: str, allowed: tuple[str, ...], answer: str) -> tuple[str, str]:
    """(value to keep, why rejected). An answer outside a closed list is matched to the list's wording."""
    plain = PLAIN_CHECKS.get(canonical(label))
    if plain:
        compact = re.sub(r"\s", "", answer)
        return (compact, "") if plain[0].match(compact) else ("", f"нужно {plain[1]}")
    if kind == "L":
        if canonical(answer) in LOGICAL_YES:
            return "Да", ""
        if canonical(answer) in LOGICAL_NO:
            return "Нет", ""
        return "", "нужно Да или Нет"
    if kind in {"N", "R"} and not re.search(r"\d", answer):
        return "", "нужно число"
    if kind == "A" and allowed:
        match = match_value(answer, Characteristic(name=label, values=list(allowed)))
        if match.value and match.confident:
            return match.value, ""
        similar = f", похожие: {', '.join(match.candidates)}" if match.candidates else ""
        return "", f"такого значения нет в списке партнёра{similar}"
    return answer, ""


def import_answers(path: Path, work_dir: Path, allowed_lists: dict[str, tuple[str, ...]] | None = None) -> AnswersResult:
    """Read a request file sent back by a manager and keep the answers that pass the checks."""
    from .raek import load_learned, learned_path  # the yes/no answers about names are RAEK's

    workbook = load_workbook(path, data_only=True)
    try:
        if META_SHEET not in workbook.sheetnames or workbook[META_SHEET]["A1"].value != MARK:
            raise ValueError(f"{path.name}: это не запрос программы (нет отметки).")
        meta = workbook[META_SHEET]
        request_id, manager = str(meta["A2"].value or ""), str(meta["A3"].value or "")
        sheet = workbook[REQUEST_SHEET]
        lists = workbook[LISTS_SHEET] if LISTS_SHEET in workbook.sheetnames else None
        rows = list(sheet.iter_rows(min_row=HEADER_ROW + 1, values_only=True))
        validations = [(str(v.sqref), v.formula1) for v in sheet.data_validations.dataValidation]
    finally:
        workbook.close()
    list_values = read_lists(lists, validations) if lists is not None else {}
    result = AnswersResult(manager=manager)
    store_file = answers_path(work_dir)
    store = read_json(store_file, {"version": 1, "answers": []})
    by_key = {(item["article"], item["key"]): item for item in store["answers"]}
    learned_file = learned_path(work_dir / "rules" / "attribute_mappings.json")
    learned = load_learned(work_dir / "rules" / "attribute_mappings.json")
    answered = []
    now = dt.datetime.now().isoformat(timespec="seconds")
    for offset, row in enumerate(rows):
        row = list(row) + [None] * len(COLUMNS)
        key = str(row[KEY - 1] or "").strip()
        answer = " ".join(str(row[ANSWER - 1] if row[ANSWER - 1] is not None else "").split())
        if not key:
            continue
        if isinstance(row[ANSWER - 1], float) and row[ANSWER - 1].is_integer():
            answer = str(int(row[ANSWER - 1]))
        article = normalize_article(row[ARTICLE - 1]) if key and not key.startswith(NAME_QUESTION) else ""
        label = str(row[LABEL - 1] or "")
        if not answer:
            result.empty += 1
            continue
        if key.startswith(NAME_QUESTION):
            _, code, attribute = key.split(":", 2)
            if canonical(answer) in LOGICAL_YES:
                names = learned["name_aliases"].setdefault(code, [])
                result.names_kept += attribute not in names
                if attribute not in names:
                    names.append(attribute)
            elif canonical(answer) in LOGICAL_NO:
                names = learned["name_rejects"].setdefault(code, [])
                result.names_rejected += attribute not in names
                if attribute not in names:
                    names.append(attribute)
            else:
                result.rejected.append(("", label, answer, "нужно Да или Нет"))
                continue
            answered.append(["", key])
            continue
        kind = KIND_BY_HINT.get(str(row[HINT - 1] or ""), "")
        allowed = (allowed_lists or {}).get(key) or list_values.get(HEADER_ROW + 1 + offset, ())
        value, why = checked_answer(label, kind, tuple(allowed), answer)
        if why:
            result.rejected.append((article, label, answer, why))
            continue
        by_key[(article, key)] = {"article": article, "key": key, "value": value, "label": label, "manager": manager, "file": path.name, "at": now}
        answered.append([article, key])
        result.accepted += 1
    store["answers"] = sorted(by_key.values(), key=lambda item: (item["article"], item["key"]))
    store_file.parent.mkdir(parents=True, exist_ok=True)
    atomic_write_json(store_file, store)
    if result.names_kept or result.names_rejected:
        learned_file.parent.mkdir(parents=True, exist_ok=True)
        atomic_write_json(learned_file, learned)
    journal_file = journal_path(work_dir)
    journal = read_json(journal_file, {"version": 1, "requests": []})
    for request in journal["requests"]:
        if request["id"] == request_id:
            known = {tuple(item) for item in request.get("answered", [])}
            request["answered"] = [list(item) for item in sorted(known | {tuple(item) for item in answered})]
            request["answered_at"] = now
    atomic_write_json(journal_file, journal)
    return result


KIND_BY_HINT = {hint: kind for kind, hint in KIND_HINTS.items()}


def read_lists(lists, validations: list[tuple[str, str]]) -> dict[int, tuple[str, ...]]:
    """Row number → the dropdown values of its answer cell."""
    columns: dict[str, tuple[str, ...]] = {}
    for col in lists.iter_cols(values_only=True):
        values = tuple(str(value) for value in col if value is not None)
        if values:
            columns[get_column_letter(len(columns) + 1)] = values
    result = {}
    for sqref, formula in validations:
        match = re.search(r"\$([A-Z]+)\$1:", formula or "")
        if not match:
            continue
        for ref in sqref.split():
            row = int(re.sub(r"\D", "", ref.split(":")[0]))
            result[row] = columns.get(match.group(1), ())
    return result


def request_lines(result: RequestsResult) -> list[str]:
    if not result.files and not result.skipped_open:
        return []
    lines = []
    for request in result.files:
        to = ", ".join(request.emails) if request.emails else "почты нет"
        lines.append(f"Запрос продакту {request.manager}: {request.items} поз. ({to}) — {request.path.name}")
    if result.skipped_open:
        lines.append(f"Уже спрошено и ждёт ответа: {result.skipped_open} поз., повторно не спрашиваю.")
    if result.unassigned_series:
        series = ", ".join(f"{name} ({count})" for name, count in result.unassigned_series.most_common(6))
        lines.append(f"Серий нет в файле закрепления, вынесены в «{NO_MANAGER}»: {series}")
    if result.missing_emails:
        lines.append(f"Нет почты: {', '.join(result.missing_emails)}. Впишите её в «{CONTACTS_NAME}» в рабочей папке.")
    return lines


def answer_lines(result: AnswersResult) -> list[str]:
    lines = [f"Ответы {result.manager}: принято {result.accepted}, пустых {result.empty}, не принято {len(result.rejected)}."]
    if result.names_kept or result.names_rejected:
        lines.append(f"Пары названий: подтверждено {result.names_kept}, отклонено {result.names_rejected}.")
    for article, label, answer, why in result.rejected[:10]:
        lines.append(f"  {article} {label}: «{answer}» — {why}")
    return lines
