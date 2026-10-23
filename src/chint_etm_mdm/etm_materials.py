"""Materials for ETM cards after they are created: photos, documents, certificates, descriptions.

ETM's OneTeam cabinet takes them in four separate forms (CSV, Win-1251, ";"). The program
fills them the way the manager did by hand:

- photos: the file name of the product's main photo on ensmas.ru and a direct link to it
  (ETM downloads it itself);
- technical information: the product's documents on ensmas.ru (passport, manual, catalog,
  drawings), the type from ETM's list taken from the file name; the files themselves are
  downloaded into "Файлы для загрузки" under the names written in the form;
- certificates: the number and the registry link from the ЧЗ file, the dates from the
  certificates table, the type and accreditation codes from the number, the PDF from the site;
- descriptions: one text per series, kept by the user in «Описания серий.xlsx» (taken from
  the catalogs), cleaned to ETM's rules.

What is missing or doubtful goes to the "Проверить" sheet of Итог.xlsx, not into the upload.
"""

from __future__ import annotations

import csv
import datetime as dt
import hashlib
import re
import shutil
import urllib.parse
import urllib.request
from collections.abc import Callable, Iterable
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .certificates import (
    TABLE_NAME as CERTIFICATES_TABLE,
    as_date,
    is_eaeu_document,
    is_refusal_letter,
    letter_date,
    number_key,
    read_chz,
    table_dates,
)
from .filler import normalize_article
from .net import open_url
from .site_info import USER_AGENT, SiteCache, SiteRecord, site_lines, update_site_cache
from .upload_names import cp1251_text, series_label, series_names

MANUFACTURER_CODE = "2019"
BRAND = "CHINT"
ADD = "Добавление материалов"
CSV_ENCODING = "cp1251"
UPLOAD_DIR = "Загрузить в ЭТМ"
CHECK_DIR = "На проверку"
FILES_DIR = "Файлы для загрузки"
SUMMARY_NAME = "Итог.xlsx"
DESCRIPTIONS_TABLE = "Описания серий.xlsx"
MAX_DOWNLOAD = 200 * 1024 * 1024

PHOTO_HEADERS = [
    "Статус", "Код производителя", "Расширенный артикул", "Бренд",
    "Основное Изображение", "Дополнительные Изображения", "Ссылка на скачивание файла",
]
DOC_HEADERS = [
    "Статус", "Код Производителя", "Расширенный Артикул ", "Бренд",
    "Название дополнительной технической информации", "Наименование файла", "Ссылка на скачивание файла",
]
CERT_HEADERS = [
    "Код производителя", "Расширенный артикул", "Номер сертификата", "Код аккредитации", "Код типа сертификата",
    "Начало срока действия", "Окончание срока действия", "Имя файла", "Ссылка на запись документа в реестре",
]
DESCRIPTION_FIELDS = [
    "Описание", "Преимущества", "УТП (Уникальное торговое предложение)", "Комплектация",
    "Сфера применения", "Дополнительная информация",
]
DESCRIPTION_HEADERS = ["Статус", "Код производителя", "Расширенный артикул", "Бренд", *DESCRIPTION_FIELDS]
INFO_FIELDS = ["Откуда", "Товаров", "Пример товара"]
TABLE_HEADERS = ["Серия", *DESCRIPTION_FIELDS, *INFO_FIELDS]
YELLOW = PatternFill("solid", fgColor="FFF2CC")
LATIN = str.maketrans("ABCEHKMOPTXY", "АВСЕНКМОРТХУ")
FORBIDDEN_IN_FILE_NAME = re.compile(r"[:/?#\[\]@!$&'()*+,;=\\\"<>|]")


@dataclass(frozen=True)
class CheckRow:
    article: str
    form: str
    problem: str
    what_to_do: str = ""


@dataclass
class MaterialsResult:
    output_dir: Path
    summary_path: Path
    articles: list[str]
    files: dict[str, Path] = field(default_factory=dict)  # form -> CSV written
    counts: dict[str, int] = field(default_factory=dict)  # form -> rows
    downloaded: int = 0
    checks: list[CheckRow] = field(default_factory=list)
    site_lines: list[str] = field(default_factory=list)
    # Cards in ETM without ТН ВЭД/ОКПД2: (article, name, missing fields, values the ЧЗ file has)
    codes: list[tuple[str, str, list[str], dict[str, str]]] = field(default_factory=list)


# --- documents -------------------------------------------------------------------------------------


def file_name_from_url(url: str) -> str:
    return urllib.parse.unquote(Path(urllib.parse.urlsplit(url).path).name)


def quoted_url(url: str) -> str:
    """The site's links carry Cyrillic and spaces as is; a request needs them encoded."""
    parts = urllib.parse.urlsplit(url)
    return urllib.parse.urlunsplit(parts._replace(path=urllib.parse.quote(urllib.parse.unquote(parts.path), safe="/")))


def document_type(name: str) -> str:
    """ETM's "Название дополнительной технической информации" for a file, or "" when unknown."""
    text = name.casefold()
    suffix = Path(text).suffix.lstrip(".")
    if suffix in {"dwg", "dxf"}:
        return f"2D чертеж (формат {suffix})"
    if suffix in {"stp", "step", "3ds", "ifc"}:
        return f"3D модель (формат {'stp' if suffix == 'step' else suffix})"
    if suffix != "pdf":
        return ""
    if "руководство" in text and "паспорт" in text:
        return "Руководство по эксплуатации и паспорт"
    if "паспорт" in text:
        return "Паспорт"
    if "руководство" in text or "инструкция" in text:
        return "Руководство по эксплуатации"
    if "каталог" in text:
        return "Каталог производителя"
    if "листовка" in text:
        return "Листовка"
    if "брошюра" in text:
        return "Брошюра"
    if "схема" in text:
        return "Схема подключения"
    if "чертеж" in text or "чертёж" in text:
        return "Чертеж"
    return "Техническая информация"


def is_certificate_document(document: dict[str, str]) -> bool:
    return "сертификат" in str(document.get("section", "")).casefold()


def download(url: str, target: Path, cache_dir: Path | None, fetch: Callable[[str], bytes] | None = None) -> bool:
    """Put the file at `url` to `target`; a file downloaded before comes from the cache. False on failure."""
    cached = None
    if cache_dir is not None:
        cached = cache_dir / hashlib.sha1(url.encode("utf-8")).hexdigest()[:16] / file_name_from_url(url)
        if cached.is_file():
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(cached, target)
            return True
    try:
        data = (fetch or fetch_bytes)(url)
    except (OSError, ValueError):
        return False
    if not data:
        return False
    target.parent.mkdir(parents=True, exist_ok=True)
    temp = target.with_name(target.name + ".part")
    temp.write_bytes(data)
    temp.replace(target)
    if cached is not None:
        cached.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(target, cached)
    return True


def fetch_bytes(url: str) -> bytes:
    request = urllib.request.Request(quoted_url(url), headers={"User-Agent": USER_AGENT})
    with open_url(request, 60) as response:
        data = response.read(MAX_DOWNLOAD + 1)
    if len(data) > MAX_DOWNLOAD:
        raise ValueError("файл больше 200 МБ")
    return data


# --- certificates ----------------------------------------------------------------------------------


def certificate_type(number: str) -> str:
    """ETM's certificate type code from the document number; "" when it cannot be told."""
    compact = re.sub(r"\s+", " ", str(number).upper().translate(LATIN))
    if not is_eaeu_document(number):
        return "ПОТКАЗ" if ("ПИСЬМО" in compact or "ОТКАЗ" in compact or is_refusal_letter(number)) else ""
    russian = " RU " in f" {compact} " or "N RU" in compact
    if re.search(r"\bД-", compact):
        return "ДТРЕАЭС" if russian else "ДТРЕАЭСбезРФ"
    if re.search(r"\bС-", compact):
        return "СТРЕАЭС" if russian else "СТРЕАЭСбезРФ"
    return ""


def accreditation_code(number: str) -> str:
    """The body's code in the number: "ЕАЭС N RU Д-CN.РА04.В.00526/25" -> "РА04" (Cyrillic, as ETM wants)."""
    found = re.search(r"-[A-ZА-Я]{2}\.([A-ZА-Я]{2}\d{2})\.", str(number).upper())
    return found.group(1).translate(LATIN) if found else ""


def certificate_file_name(number: str) -> str:
    """As the manager named them: spaces, dots and slashes become "_"."""
    name = re.sub(r"[\s./]+", "_", str(number).strip())
    return FORBIDDEN_IN_FILE_NAME.sub("_", name) + ".pdf"


def plus_ten_years(day: dt.date) -> dt.date:
    try:
        return day.replace(year=day.year + 10)
    except ValueError:  # 29 February
        return day.replace(year=day.year + 10, day=28)


# --- descriptions ----------------------------------------------------------------------------------


def clean_text(text: str) -> str:
    """ETM's rules for description texts: one line, no ";", no degree sign, quotes as two apostrophes."""
    text = str(text or "")
    text = re.sub(r"[   ]", " ", text)
    text = re.sub(r"[​-‍﻿\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
    text = re.sub(r"\s*[\r\n]+\s*", " ", text)
    text = text.replace(";", ",")
    text = re.sub(r"\s*°\s*[CС]\b", " градусов Цельсия", text)
    text = text.replace("°", " градусов")
    text = re.sub("[«»“”„\"]", "''", text)
    text = text.replace(":", " -")
    return re.sub(r" {2,}", " ", cp1251_text(text)).strip()


def description_problems(fields: dict[str, str]) -> list[str]:
    problems = []
    description = fields.get("Описание", "")
    if not 250 <= len(description) <= 1000:
        problems.append(f"описание {len(description)} символов, ЭТМ просит от 250 до 1000")
    advantages = [item for item in fields.get("Преимущества", "").split("|") if item.strip()]
    if advantages and not 2 <= len(advantages) <= 7:
        problems.append(f"преимуществ {len(advantages)}, ЭТМ просит от 2 до 7")
    if any(len(item.strip()) > 120 for item in advantages):
        problems.append("пункт преимуществ длиннее 120 символов")
    if len(fields.get("УТП (Уникальное торговое предложение)", "")) > 75:
        problems.append("УТП длиннее 75 символов")
    for name in DESCRIPTION_FIELDS:
        if len(fields.get(name, "")) > 1000:
            problems.append(f"«{name}» длиннее 1000 символов")
        if re.search(r"[~=]", fields.get(name, "")):
            problems.append(f"в «{name}» есть «~» или «=»: для тока пишите AC и DC")
    return problems


def read_descriptions(path: Path) -> dict[str, dict[str, str]]:
    """Series (upper case) -> texts from «Описания серий.xlsx»."""
    if not path.exists():
        return {}
    workbook = load_workbook(path, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        return {}
    header = [str(cell or "").strip() for cell in rows[0]]
    found = {}
    for row in rows[1:]:
        values = {name: str(row[idx] or "").strip() for idx, name in enumerate(header) if idx < len(row) and name}
        series = values.get("Серия", "")
        if series:
            found[series.upper()] = {name: values.get(name, "") for name in DESCRIPTION_FIELDS + INFO_FIELDS}
    return found


def write_descriptions(path: Path, texts: dict[str, dict[str, str]], counts: dict[str, int], examples: dict[str, str]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Описания серий"
    sheet.append(TABLE_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for series in sorted(texts):
        values = texts[series]
        sheet.append(
            [
                series,
                *(values.get(name, "") for name in DESCRIPTION_FIELDS),
                values.get("Откуда", ""),
                counts.get(series) or as_count(values.get("Товаров")),
                examples.get(series) or values.get("Пример товара", ""),
            ]
        )
        if not values.get("Описание"):
            sheet.cell(sheet.max_row, 2).fill = YELLOW
    widths = [14, 80, 60, 30, 40, 40, 40, 22, 9, 50]
    for column, width in zip("ABCDEFGHIJ", widths):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "B2"
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(path.stem + ".saving.xlsx")
    workbook.save(temp)
    temp.replace(path)


def as_count(value: object) -> int | str:
    text = str(value or "").strip()
    return int(float(text)) if re.fullmatch(r"\d+(\.0)?", text) else ""


def update_descriptions_table(path: Path, series_of: dict[str, str], names: dict[str, str]) -> int:
    """Add a row for every series of these products that the table does not have yet; returns how many.

    The file is rewritten only when a series is added, so a table open in Excel stays as it is."""
    texts = read_descriptions(path)
    counts: dict[str, int] = {}
    examples: dict[str, str] = {}
    for article, series in series_of.items():
        key = series.strip().upper()
        if key:
            counts[key] = counts.get(key, 0) + 1
            examples.setdefault(key, names.get(article, article))
    added = 0
    for key in counts:
        if key not in texts:
            texts[key] = {name: "" for name in DESCRIPTION_FIELDS}
            added += 1
    if added or not path.exists():
        write_descriptions(path, texts, counts, examples)
    return added


def import_description_file(csv_path: Path, table_path: Path, series_of: Callable[[list[str]], dict[str, str]]) -> int:
    """Take texts from an ETM description file sent earlier (per product) into the series table.

    Only empty series rows are filled: what the user wrote in the table stays.
    """
    raw = csv_path.read_bytes()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode(CSV_ENCODING)
    rows = list(csv.reader(text.splitlines(), delimiter=";"))
    if not rows:
        return 0
    header = [cell.strip() for cell in rows[0]]
    article_idx = next((idx for idx, name in enumerate(header) if name.casefold() == "расширенный артикул"), None)
    if article_idx is None or "Описание" not in header:
        raise ValueError(f"{csv_path.name}: это не файл описаний ЭТМ (нет колонок «Расширенный артикул» и «Описание»).")
    by_article = {}
    for row in rows[1:]:
        if article_idx < len(row) and normalize_article(row[article_idx]):
            by_article[normalize_article(row[article_idx])] = {
                name: (row[header.index(name)] if header.index(name) < len(row) else "").strip()
                for name in DESCRIPTION_FIELDS
                if name in header
            }
    series = series_of(sorted(by_article))
    texts = read_descriptions(table_path)
    taken = 0
    for article, values in by_article.items():
        key = series.get(article, "").strip().upper()
        if not key or not values.get("Описание"):
            continue
        current = texts.setdefault(key, {name: "" for name in DESCRIPTION_FIELDS})
        if current.get("Описание"):
            continue
        current.update({name: values.get(name, "") for name in DESCRIPTION_FIELDS})
        current["Откуда"] = csv_path.name
        taken += 1
    write_descriptions(table_path, texts, {}, {})
    return taken


# --- the run ---------------------------------------------------------------------------------------


def write_csv(path: Path, header: list[str], rows: list[list[object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding=CSV_ENCODING, errors="replace", newline="") as handle:
        writer = csv.writer(handle, delimiter=";", lineterminator="\r\n")
        writer.writerow(header)
        writer.writerows([[cp1251_text(value) if isinstance(value, str) else value for value in row] for row in rows])


def csv_name(kind: str, series: Iterable[str], when: dt.date) -> str:
    return f"upload_{kind}_{series_label(series)}_{when.isoformat()}.csv"


def make_materials(
    db_path: Path,
    articles: list[str],
    work_dir: Path,
    output_dir: Path,
    codes_path: Path | None,
    ask_site: bool = True,
    progress: Callable[[str], None] | None = None,
    fetch_page: Callable[[str], str] | None = None,
    fetch_file: Callable[[str], bytes] | None = None,
    today: dt.date | None = None,
    site_delay: float | None = None,
    wanted: dict[str, set[str]] | None = None,
    codes_missing: dict[str, list[str]] | None = None,
) -> MaterialsResult:
    """The four ETM forms for the products, the files to upload with them and a summary.

    `wanted` (form -> articles, e.g. from an ETM report) limits each form to the products that lack
    it; "2D/3D" asks for drawings and models only. `codes_missing` (article -> code names) lists
    products whose card in ETM lacks ТН ВЭД or ОКПД2.
    """
    today = today or dt.date.today()
    say = progress or (lambda _text: None)
    if wanted is not None:
        articles = sorted({article for members in wanted.values() for article in members})
    articles = list(dict.fromkeys(normalize_article(article) for article in articles if normalize_article(article)))

    def wants(form: str, article: str) -> bool:
        return wanted is None or article in wanted.get(form, ())

    result = MaterialsResult(output_dir, output_dir / SUMMARY_NAME, articles)
    upload_dir, check_dir, files_dir = (output_dir / name for name in (UPLOAD_DIR, CHECK_DIR, FILES_DIR))
    series = series_names(db_path, articles)
    names = product_names(db_path, articles)

    cache = SiteCache(work_dir / "cache" / "ensmas_site.json")
    if ask_site:
        say("Спрашиваю сайт ensmas.ru о товарах, которых нет в памяти: по 2–3 секунды на товар...")
        options: dict = {"fetch": fetch_page} if fetch_page else {}
        if site_delay is not None:
            options["delay"] = site_delay
        on_site = [article for article in articles if any(wants(form, article) for form in ("Фото", "Тех. информация", "2D/3D", "Сертификаты"))]
        result.site_lines = site_lines(update_site_cache(cache, on_site, progress=say, **options))
    records: dict[str, SiteRecord | None] = {article: cache.found(article) for article in articles}
    file_cache = work_dir / "cache" / "ensmas_files"
    all_series = [series.get(article, "") for article in articles]

    # Photos: the main photo on the site, by link.
    photo_rows = []
    for article in articles:
        if not wants("Фото", article):
            continue
        record = records[article]
        if record is None or not record.image_url:
            result.checks.append(CheckRow(article, "Фото", "на ensmas.ru нет фото товара", "добавьте фото в ЭТМ вручную"))
            continue
        name = file_name_from_url(record.image_url)  # Cyrillic, "+" and spaces are fine: ETM takes the link
        photo_rows.append([ADD, MANUFACTURER_CODE, article, BRAND, name, "", record.image_url])

    # Technical information: the documents on the site, files downloaded.
    doc_rows = []
    saved: dict[str, bool] = {}
    for article in articles:
        all_documents = wants("Тех. информация", article)
        models_only = not all_documents and wants("2D/3D", article)
        record = records[article]
        if record is None or not (all_documents or models_only):
            continue
        for document in record.documents:
            if is_certificate_document(document) or not document.get("url"):
                continue
            name = file_name_from_url(document["url"])
            kind = document_type(name)
            if models_only and not kind.startswith(("2D", "3D")):
                continue  # the card has its documents, only drawings and models are missing
            if not kind:
                result.checks.append(CheckRow(article, "Тех. информация", f"файл «{name}»: не понятно, какой это тип для ЭТМ", "посмотрите файл и загрузите его сами"))
                continue
            if name not in saved:
                say(f"Скачиваю {name}...")
                saved[name] = download(document["url"], files_dir / name, file_cache, fetch_file)
                result.downloaded += saved[name]
            if not saved[name]:
                result.checks.append(CheckRow(article, "Тех. информация", f"не удалось скачать «{name}»", document["url"]))
                continue
            doc_rows.append([ADD, MANUFACTURER_CODE, article, "", kind, name, ""])

    # Certificates: number and registry link from ЧЗ, dates from the table, the PDF from the site.
    cert_rows = []
    chz = {row["article"]: row for row in read_chz(codes_path)}
    dates = table_dates(work_dir / CERTIFICATES_TABLE)
    pdf_by_number: dict[str, str] = {}
    for article in articles:
        record = records[article]
        row = chz.get(article)
        if record and row and row.get("number"):
            pdf = next((doc["url"] for doc in record.documents if is_certificate_document(doc) and doc.get("url", "").lower().endswith(".pdf")), "")
            if pdf:
                pdf_by_number.setdefault(number_key(str(row["number"])), pdf)
    for article in articles:
        if not wants("Сертификаты", article):
            continue
        row = chz.get(article)
        number = " ".join(str((row or {}).get("number") or "").split())
        if not number or number.casefold() in {"не требуется", "-"}:
            if wanted is not None and number:
                result.checks.append(
                    CheckRow(article, "Сертификаты", f"в файле ЧЗ «{number}», а ЭТМ просит документ", "приложите отказное письмо (тип ПОТКАЗ) или документ")
                )
            elif codes_path is None:
                result.checks.append(CheckRow(article, "Сертификаты", "нет файла ЧЗ", "сохраните свежий файл ЧЗ в «Загрузки»"))
            elif row is None:
                result.checks.append(CheckRow(article, "Сертификаты", "артикула нет в файле ЧЗ"))
            continue
        kind = certificate_type(number)
        if not kind:
            result.checks.append(CheckRow(article, "Сертификаты", f"не понятен тип документа «{number}»", "укажите код типа по листу ЭТМ «Типы сертификатов»"))
            continue
        key = number_key(number)
        start, end = dates.get(key, (None, None))
        start = start or as_date(row.get("start")) or letter_date(number)
        end = end or as_date(row.get("end"))
        if start is None:
            result.checks.append(CheckRow(article, "Сертификаты", f"нет даты начала для «{number}»", "впишите даты в «Сертификаты.xlsx» (вкладка РАЭК)"))
            continue
        end = end or plus_ten_years(start)  # ETM's rule for a document without an end date
        link = str(row.get("link") or "").strip()
        accreditation = accreditation_code(number) if kind in {"ДТРЕАЭС", "СТРЕАЭС"} else ""
        if kind in {"ДТРЕАЭС", "СТРЕАЭС", "ДТРЕАЭСбезРФ", "СТРЕАЭСбезРФ"} and not link:
            result.checks.append(CheckRow(article, "Сертификаты", f"нет ссылки на реестр для «{number}»", "ЭТМ требует ссылку для этого типа"))
            continue
        name = certificate_file_name(number)
        if key not in saved:
            pdf = pdf_by_number.get(key, "")
            saved[key] = bool(pdf) and download(pdf, files_dir / name, file_cache, fetch_file)
            result.downloaded += saved[key]
            if not saved[key]:
                result.checks.append(CheckRow(article, "Сертификаты", f"нет PDF документа «{number}» на сайте", f"положите файл {name} в папку «{FILES_DIR}» сами"))
        cert_rows.append(
            [MANUFACTURER_CODE, article, number, accreditation, kind, f"{start:%d.%m.%Y}", f"{end:%d.%m.%Y}", name, link]
        )

    # Descriptions: one text per series from the table.
    table = work_dir / DESCRIPTIONS_TABLE
    try:
        update_descriptions_table(table, {article: value for article, value in series.items() if wants("Описание", article)}, names)
    except PermissionError:
        result.checks.append(
            CheckRow("", "Описание", f"«{DESCRIPTIONS_TABLE}» открыт в Excel: новые серии в него не добавлены", "закройте файл и соберите материалы заново")
        )
    texts = read_descriptions(table)
    description_rows, held_rows = [], []
    for article in articles:
        if not wants("Описание", article):
            continue
        key = series.get(article, "").strip().upper()
        values = texts.get(key, {})
        if not values.get("Описание"):
            result.checks.append(CheckRow(article, "Описание", f"нет описания серии «{series.get(article, '') or 'серия не указана'}»", f"впишите его в «{DESCRIPTIONS_TABLE}» из каталога"))
            continue
        fields = {name: clean_text(values.get(name, "")) for name in DESCRIPTION_FIELDS}
        line = [ADD, MANUFACTURER_CODE, article, BRAND, *(fields[name] for name in DESCRIPTION_FIELDS)]
        problems = description_problems(fields)
        if problems:
            held_rows.append(line)
            result.checks.append(CheckRow(article, "Описание", "; ".join(problems), f"поправьте в «{DESCRIPTIONS_TABLE}»: файл описаний в папке «{CHECK_DIR}»"))
        else:
            description_rows.append(line)

    for form, kind, header, rows, folder in (
        ("Фото", "img", PHOTO_HEADERS, photo_rows, upload_dir),
        ("Тех. информация", "doc", DOC_HEADERS, doc_rows, upload_dir),
        ("Сертификаты", "cert", CERT_HEADERS, cert_rows, upload_dir),
        ("Описание", "desc", DESCRIPTION_HEADERS, description_rows, upload_dir),
        ("Описание на проверку", "desc", DESCRIPTION_HEADERS, held_rows, check_dir),
    ):
        result.counts[form] = len(rows)
        if rows:
            path = folder / csv_name(kind, all_series, today)
            if folder == check_dir:
                path = path.with_name(f"{path.stem}_check.csv")
            write_csv(path, header, rows)
            result.files[form] = path
    if codes_missing:
        known = load_codes(codes_path)
        for article, fields in sorted(codes_missing.items()):
            result.codes.append((article, names.get(article, ""), fields, {field: known.get((article, field), "") for field in fields}))
    write_summary(result, names)
    return result


def load_codes(codes_path: Path | None) -> dict[tuple[str, str], str]:
    if codes_path is None:
        return {}
    from .etm_categories import load_codes_file

    try:
        return load_codes_file(codes_path)
    except (OSError, ValueError):
        return {}


def product_names(db_path: Path, articles: list[str]) -> dict[str, str]:
    from .db import fetch_products
    from .export_templates import clean_product_name

    return {article: clean_product_name(product.full_name or product.name) for article, product in fetch_products(db_path, articles).items()}


def todo_lines(result: MaterialsResult) -> list[str]:
    lines = []
    forms = [form for form in ("Фото", "Тех. информация", "Сертификаты", "Описание") if form in result.files]
    if forms:
        lines.append(
            f"1. Загрузите в ЭТМ файлы из папки «{UPLOAD_DIR}» — каждый в свою форму личного кабинета: "
            + ", ".join(f"{form.lower()} ({result.counts[form]})" for form in forms)
            + "."
        )
    if any(result.counts.get(form) for form in ("Тех. информация", "Сертификаты")):
        lines.append(
            f"2. К формам «Тех. информация» и «Сертификаты» приложите файлы из папки «{FILES_DIR}»: "
            "их имена уже вписаны в формы."
        )
    if "Описание на проверку" in result.files:
        lines.append(f"3. Описания в папке «{CHECK_DIR}» не проходят правила ЭТМ (лист «Проверить»): поправьте и соберите заново.")
    if result.checks:
        lines.append(f"Что не получилось сделать само: лист «Проверить» ({len(result.checks)} строк).")
    if result.codes:
        known = sum(1 for _a, _n, fields, values in result.codes if all(values.get(field) for field in fields))
        lines.append(
            f"У {len(result.codes)} карточек в ЭТМ нет ТН ВЭД или ОКПД2 (лист «Коды в карточке»); у {known} коды есть "
            "в файле ЧЗ. Коды — поля самой карточки, их меняют через карточку товара в ЭТМ."
        )
    if not lines:
        lines.append("Делать ничего не нужно.")
    lines.append("")
    lines.append("Фото ЭТМ заберёт сам по ссылке на ensmas.ru, загружать их файлами не нужно.")
    lines.append(f"Описания серий хранятся в рабочей папке, в «{DESCRIPTIONS_TABLE}»: впишите туда тексты из каталогов один раз.")
    return lines


def write_summary(result: MaterialsResult, names: dict[str, str]) -> None:
    workbook = Workbook()
    todo = workbook.active
    todo.title = "Что делать"
    for line in todo_lines(result):
        todo.append([line])
    todo.column_dimensions["A"].width = 120
    check = workbook.create_sheet("Проверить")
    check.append(["Артикул", "Название", "Форма", "Что не так", "Что сделать"])
    for cell in check[1]:
        cell.font = Font(bold=True)
    for row in result.checks:
        check.append([row.article, names.get(row.article, ""), row.form, row.problem, row.what_to_do])
    for column, width in zip("ABCDE", (14, 60, 18, 70, 70)):
        check.column_dimensions[column].width = width
    check.freeze_panes = "A2"
    check.auto_filter.ref = check.dimensions
    if result.codes:
        codes = workbook.create_sheet("Коды в карточке")
        codes.append(["Артикул", "Название", "Нет в ЭТМ", "Код ТН ВЭД из файла ЧЗ", "Код ОКПД2 из файла ЧЗ"])
        for cell in codes[1]:
            cell.font = Font(bold=True)
        for article, name, fields, values in result.codes:
            codes.append([article, name, ", ".join(fields), values.get("Код ТН ВЭД", ""), values.get("Код ОКПД2", "")])
        for column, width in zip("ABCDE", (14, 60, 24, 24, 24)):
            codes.column_dimensions[column].width = width
        codes.freeze_panes = "A2"
        codes.auto_filter.ref = codes.dimensions
    result.output_dir.mkdir(parents=True, exist_ok=True)
    workbook.save(result.summary_path)


def materials_lines(result: MaterialsResult) -> list[str]:
    lines = [f"Товаров: {len(result.articles)}"]
    for form in ("Фото", "Тех. информация", "Сертификаты", "Описание"):
        lines.append(f"{form}: {result.counts.get(form, 0)} строк")
    if result.counts.get("Описание на проверку"):
        lines.append(f"Описаний на проверку: {result.counts['Описание на проверку']}")
    lines.append(f"Скачано файлов: {result.downloaded}")
    if result.checks:
        lines.append(f"Проверить: {len(result.checks)} (лист «Проверить» в Итог.xlsx)")
    lines.extend(result.site_lines)
    lines.append(f"Папка: {result.output_dir.name}")
    return lines


# --- ETM's reports on what cards lack ---------------------------------------------------------------

GOODS_FLAGS = {
    "наличие фото": "Фото",
    "наличие файлов тех.информации": "Тех. информация",
    "наличие описания": "Описание",
    "наличие 2d/3d модели": "2D/3D",
}
CODE_FLAGS = {"наличие тнвэд": "Код ТН ВЭД", "наличие окпд 2": "Код ОКПД2"}


@dataclass
class ReportNeeds:
    """What ETM's reports say the cards lack: form -> articles, and the codes missing in cards."""

    wanted: dict[str, set[str]] = field(default_factory=dict)
    codes_missing: dict[str, list[str]] = field(default_factory=dict)
    reasons: dict[str, str] = field(default_factory=dict)  # article -> ETM's reason about certificates
    files: list[str] = field(default_factory=list)

    def add(self, other: "ReportNeeds") -> None:
        for form, members in other.wanted.items():
            self.wanted.setdefault(form, set()).update(members)
        for article, fields in other.codes_missing.items():
            self.codes_missing.setdefault(article, [])
            self.codes_missing[article] += [item for item in fields if item not in self.codes_missing[article]]
        self.reasons.update(other.reasons)
        self.files += other.files


def report_rows(path: Path) -> list[list[str]]:
    """The first sheet of an XLSX that has a header, or a CSV (Win-1251 or UTF-8, ";")."""
    if path.suffix.lower() == ".csv":
        raw = path.read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode(CSV_ENCODING)
        return [[cell.strip() for cell in row] for row in csv.reader(text.splitlines(), delimiter=";")]
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            rows = [["" if cell is None else str(cell).strip() for cell in row] for row in sheet.iter_rows(values_only=True)]
            rows = [row for row in rows if any(row)]
            if rows:
                return rows
    finally:
        workbook.close()
    return []


def read_etm_report(path: Path) -> ReportNeeds:
    """ETM's goods list («Наличие Фото», «Наличие описания»...) or its certificates report («Причина»)."""
    rows = report_rows(path)
    if not rows:
        raise ValueError(f"{path.name}: файл пустой.")
    header = [" ".join(cell.split()).casefold() for cell in rows[0]]
    article_idx = next((idx for idx, name in enumerate(header) if name == "расширенный артикул"), None)
    if article_idx is None:
        raise ValueError(f"{path.name}: нет колонки «Расширенный артикул» — это не отчёт ЭТМ о товарах.")
    needs = ReportNeeds(files=[path.name])

    def cell(row: list[str], idx: int | None) -> str:
        return row[idx] if idx is not None and idx < len(row) else ""

    if any(name in header for name in GOODS_FLAGS):
        supplied_idx = header.index("поставляемый") if "поставляемый" in header else None
        for row in rows[1:]:
            article = normalize_article(cell(row, article_idx))
            if not article or cell(row, supplied_idx).casefold() == "нет":
                continue  # products ETM no longer takes are left alone
            for name, form in GOODS_FLAGS.items():
                if name in header and cell(row, header.index(name)).casefold() == "нет":
                    needs.wanted.setdefault(form, set()).add(article)
            for name, code in CODE_FLAGS.items():
                if name in header and cell(row, header.index(name)).casefold() == "нет":
                    needs.codes_missing.setdefault(article, []).append(code)
        return needs
    if "причина" in header:
        reason_idx = header.index("причина")
        for row in rows[1:]:
            article = normalize_article(cell(row, article_idx))
            if article and cell(row, reason_idx):
                needs.wanted.setdefault("Сертификаты", set()).add(article)
                needs.reasons[article] = cell(row, reason_idx)
        return needs
    raise ValueError(
        f"{path.name}: не знаю такой отчёт. Подходят список товаров ЭТМ (колонки «Наличие Фото», «Наличие описания»...) "
        "и отчёт по сертификатам (колонка «Причина»)."
    )


def report_lines(needs: ReportNeeds) -> list[str]:
    lines = [f"Отчёты: {', '.join(needs.files)}"]
    labels = {"Фото": "без фото", "Тех. информация": "без тех. информации", "2D/3D": "без 2D/3D", "Описание": "без описания", "Сертификаты": "с замечаниями по сертификатам"}
    for form, label in labels.items():
        if needs.wanted.get(form):
            lines.append(f"Товаров {label}: {len(needs.wanted[form])}")
    if needs.codes_missing:
        lines.append(f"Карточек без ТН ВЭД или ОКПД2: {len(needs.codes_missing)}")
    return lines
