"""New product cards for ETM through category templates.

ETM's OneTeam cabinet gives one upload template per product class ("81 класс"):
the usual card columns plus the class's required characteristics ("Конфиг:..."),
most of them with a drop-down list on the "Списки" sheet. The user keeps the
downloaded templates in one folder (templates saved to the inbox or "Загрузки"
are picked up from there). For the products the program picks each product's
class, fills a copy of that class's template and writes the CSV that ETM takes.
A run is one folder: Итог.xlsx (what to do, every product's state, the empty
fields, the values to check, the classes whose template is missing), the files
ready to upload, the cards to check first and the working files. The empty
fields become questions for the product managers; their answers fill the next run.
"""
from __future__ import annotations

import csv
import datetime as dt
import json
import re
import shutil
import sqlite3
from collections import Counter
from contextlib import closing
from copy import copy
from dataclasses import asdict, dataclass, field, replace
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .atomic_files import atomic_write_json
from .db import ProductRecord, connect, fetch_products
from .export_templates import PartnerTemplate, clean_product_name, load_partner_template, partner_template_text
from .filler import FillReportRow, FillResult, fill_xlsx_template, normalize_article, normalize_header
from .pm_requests import Need, load_answers
from .upload_names import check_path_for, report_path_for, series_names, upload_csv_path
from .value_matching import canonical

TEMPLATE_SHEET = "Шаблон"
LISTS_SHEET = "Списки"
CLASS_HEADER = "81 класс"
CONFIG_PREFIX = "Конфиг:"
FILE_CLASS_RE = re.compile(r"upload_goods_category_template_([0-9A-Za-z]+)", re.IGNORECASE)
CODE_HEADERS = {"part number": "article", "tariff code": "Код ТН ВЭД", "окпд 2 (9 зн.)": "Код ОКПД2"}
CSV_ENCODING = "cp1251"


@dataclass(frozen=True)
class CategoryTemplate:
    class_code: str
    path: Path
    dropdowns: dict[str, tuple[str, ...]]


@dataclass(frozen=True)
class ClassChoice:
    code: str
    basis: str
    sure: bool


@dataclass
class CategoryRun:
    output_dir: Path
    summary_path: Path
    classes: dict[str, ClassChoice] = field(default_factory=dict)
    results: dict[str, tuple[FillResult, Path]] = field(default_factory=dict)
    missing_templates: dict[str, list[str]] = field(default_factory=dict)
    unclassified: list[str] = field(default_factory=list)
    not_found: list[str] = field(default_factory=list)
    codes_file: Path | None = None
    articles: list[str] = field(default_factory=list)
    missing_rows: list[FillReportRow] = field(default_factory=list)
    check_rows: list[FillReportRow] = field(default_factory=list)
    # (class, canonical field) -> ETM's closed list of values for it
    allowed: dict[tuple[str, str], tuple[str, ...]] = field(default_factory=dict)
    needs: list[Need] = field(default_factory=list)
    copied_templates: list[Path] = field(default_factory=list)


def cell_text(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return normalize_header(value)


def read_category_template(path: Path) -> CategoryTemplate:
    workbook = load_workbook(path)
    if TEMPLATE_SHEET not in workbook.sheetnames:
        raise ValueError(f"В файле {path.name} нет листа «{TEMPLATE_SHEET}».")
    sheet = workbook[TEMPLATE_SHEET]
    headers = {cell.column_letter: normalize_header(cell.value) for cell in sheet[1]}
    class_letter = next((letter for letter, name in headers.items() if name == CLASS_HEADER), None)
    class_code = cell_text(sheet[f"{class_letter}2"].value) if class_letter else ""
    if not class_code:
        found = FILE_CLASS_RE.search(path.stem)
        class_code = found.group(1) if found else ""
    if not class_code:
        raise ValueError(f"В шаблоне {path.name} не указан 81 класс.")
    lists = workbook[LISTS_SHEET] if LISTS_SHEET in workbook.sheetnames else None
    dropdowns: dict[str, tuple[str, ...]] = {}
    for validation in sheet.data_validations.dataValidation:
        if validation.type != "list" or not validation.formula1 or lists is None:
            continue
        sheet_name, _, area = validation.formula1.partition("!")
        if sheet_name.strip("'") != LISTS_SHEET:
            continue
        values = tuple(
            text for row in lists[area.replace("$", "")] for cell in row if (text := cell_text(cell.value))
        )
        for cell_range in validation.sqref.ranges:
            header = headers.get(re.sub(r"\d", "", str(cell_range).split(":")[0]), "")
            if header.startswith(CONFIG_PREFIX) and values:
                dropdowns[canonical(header.removeprefix(CONFIG_PREFIX))] = values
    return CategoryTemplate(class_code, path, dropdowns)


def category_library(folder: Path) -> dict[str, Path]:
    """Class code -> template file, for every category template in the folder."""
    library: dict[str, Path] = {}
    for path in sorted(folder.glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        found = FILE_CLASS_RE.search(path.stem)
        if found:
            library[found.group(1)] = path
            continue
        try:
            library[read_category_template(path).class_code] = path
        except (ValueError, KeyError, OSError):
            continue
    return library


def new_articles_since(db_path: Path, since: dt.date) -> list[str]:
    """Articles whose price-list introduction date ("Введено в ассортимент") is after `since`."""
    with closing(connect(db_path)) as conn:
        rows = conn.execute("SELECT article, introduced_at FROM new_products").fetchall()
    result = []
    for article, introduced in rows:
        try:
            date = dt.datetime.strptime(str(introduced or "").strip()[:10], "%d.%m.%Y").date()
        except ValueError:
            continue
        if date > since:
            result.append(str(article))
    return sorted(result)


def series_family(series: str) -> str:
    return re.split(r"[-\s/(]", series.strip(), maxsplit=1)[0].upper()


def suggest_classes(
    db_path: Path, products: dict[str, ProductRecord], with_template: frozenset[str] = frozenset()
) -> dict[str, ClassChoice]:
    """The class each product already has, else the usual class of its series in ETM's catalog.

    When the series is split between classes, a class whose template the user has wins over a more
    frequent one without a template (ETM's catalog mixes in service positions); such a choice is
    marked for a check.
    """
    choices = {
        article: ClassChoice(product.class81_code, "уже есть в базе", True)
        for article, product in products.items()
        if product.class81_code
    }
    pending = [article for article in products if article not in choices]
    if not pending:
        return choices
    try:
        with closing(connect(db_path)) as conn:
            known = conn.execute(
                """
                SELECT DISTINCT p.article, COALESCE(p.series_name, ''), COALESCE(p.equipment_group, ''), ig.class81_code
                FROM products p JOIN ipro_goods ig ON ig.article = p.article
                WHERE COALESCE(ig.class81_code, '') <> ''
                """
            ).fetchall()
            placeholders = ",".join("?" for _ in pending)
            own = conn.execute(
                f"SELECT article, COALESCE(series_name, ''), COALESCE(equipment_group, '') FROM products WHERE article IN ({placeholders})",
                pending,
            ).fetchall()
    except sqlite3.OperationalError:
        return choices
    by_series: dict[str, Counter[str]] = {}
    by_family: dict[str, Counter[str]] = {}
    by_group: dict[str, Counter[str]] = {}
    for _, series, group, code in known:
        code = cell_text(code)
        if series:
            by_series.setdefault(series.strip().upper(), Counter())[code] += 1
            by_family.setdefault(series_family(series), Counter())[code] += 1
        if group:
            by_group.setdefault(group.strip(), Counter())[code] += 1
    for article, series, group in own:
        for counts, basis, sure in (
            (by_series.get(series.strip().upper()), f"как у серии {series}", True),
            (by_family.get(series_family(series)) if series else None, f"как у серий {series_family(series)}", True),
            (by_group.get(group.strip()) if group else None, f"как у группы «{group}»", False),
        ):
            if counts:
                total = sum(counts.values())
                code, votes = counts.most_common(1)[0]
                if code not in with_template:
                    usable = [(item, count) for item, count in counts.most_common() if item in with_template]
                    if usable:
                        code, votes = usable[0]
                        sure = False
                choices[str(article)] = ClassChoice(code, f"{basis} ({votes} из {total})", sure and votes / total >= 0.8)
                break
    return choices


CODES_FILE_RE = re.compile(r"чз|честн", re.IGNORECASE)
NAME_DATE_RE = re.compile(r"(\d{2})\.(\d{2})\.(\d{4})")


def codes_file_date(path: Path) -> tuple[dt.date, float]:
    """The date in the name ("..._19.10.2026.xlsx"), else the file's own date; newer sorts higher."""
    found = NAME_DATE_RE.search(path.stem)
    mtime = path.stat().st_mtime
    if found:
        try:
            return dt.date(int(found.group(3)), int(found.group(2)), int(found.group(1))), mtime
        except ValueError:
            pass
    return dt.date.fromtimestamp(mtime), mtime


def latest_codes_file(folders: list[Path], chosen: Path | None = None) -> Path | None:
    """The newest "Честный знак" product list among the chosen file and the files in `folders`.

    The product manager sends a new version now and then; the user only saves it
    to Downloads and the program takes the newest one on the next run.
    """
    found = [chosen] if chosen and chosen.is_file() else []
    for folder in folders:
        if folder.is_dir():
            found += [
                path
                for path in folder.glob("*.xlsx")
                if CODES_FILE_RE.search(path.stem) and not path.name.startswith("~$")
            ]
    return max(found, key=codes_file_date, default=None)


def load_codes_file(path: Path, headers: dict[str, str] = CODE_HEADERS) -> dict[tuple[str, str], str]:
    """ТН ВЭД and ОКПД2 by article from the "Честный знак" product list (sheet "Price-list").

    `headers` maps the file's column names (lower case) to the names the values are kept under.
    """
    # keep_links=False: the company file links to another workbook (150 MB), reading it takes a minute.
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        sheet = workbook["Price-list"] if "Price-list" in workbook.sheetnames else workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        columns: dict[str, int] = {}
        for row in rows:
            names = [" ".join(str(cell or "").split()).casefold() for cell in row]
            if "part number" in names:
                columns = {headers[name]: idx for idx, name in enumerate(names) if name in headers}
                break
        if "article" not in columns:
            raise ValueError("В файле кодов нет колонки «Part number».")
        values: dict[tuple[str, str], str] = {}
        for row in rows:
            article = normalize_article(row[columns["article"]] if columns["article"] < len(row) else "")
            if not article:
                continue
            for header, idx in columns.items():
                value = cell_text(row[idx]) if header != "article" and idx < len(row) else ""
                if value and value not in {"0", "-"}:
                    values[(article, header)] = value
        return values
    finally:
        workbook.close()


@dataclass(frozen=True)
class SeriesDefault:
    """A value the whole series shares, for a characteristic the database has no data on."""

    classes: frozenset[str]
    series: str
    characteristic: str
    value: str
    basis: str = ""
    if_name: str = ""
    unless_name: str = ""

    def applies(self, class_code: str, series: str, name: str) -> bool:
        if self.classes and class_code not in self.classes:
            return False
        if series_family(series) != series_family(self.series) and series.strip().upper() != self.series.upper():
            return False
        if self.if_name and not re.search(self.if_name, name, re.IGNORECASE):
            return False
        return not (self.unless_name and re.search(self.unless_name, name, re.IGNORECASE))


def load_series_defaults(partner: str = "etm") -> list[SeriesDefault]:
    raw = json.loads(partner_template_text(partner)).get("series_defaults", [])
    return [
        SeriesDefault(
            classes=frozenset(str(code) for code in item.get("classes", [])),
            series=item["series"],
            characteristic=item["characteristic"],
            value=item["value"],
            basis=item.get("basis", ""),
            if_name=item.get("if_name", ""),
            unless_name=item.get("unless_name", ""),
        )
        for item in raw
    ]


def series_default_values(
    defaults: list[SeriesDefault],
    category: CategoryTemplate,
    articles: list[str],
    products: dict[str, ProductRecord],
    series: dict[str, str],
) -> dict[tuple[str, str], tuple[str, str]]:
    values: dict[tuple[str, str], tuple[str, str]] = {}
    for article in articles:
        product = products[article]
        name = product.full_name or product.name
        for rule in defaults:
            if rule.applies(category.class_code, series.get(article, ""), name):
                source = f"series_default:{rule.series}" + (f" ({rule.basis})" if rule.basis else "")
                values[(article, CONFIG_PREFIX + rule.characteristic)] = (rule.value, source)
    return values


def prepare_input(category: CategoryTemplate, articles: list[str], target: Path) -> None:
    """A copy of the category template with one row per article (style of ETM's first row)."""
    workbook = load_workbook(category.path)
    sheet = workbook[TEMPLATE_SHEET]
    workbook.active = workbook.sheetnames.index(TEMPLATE_SHEET)
    headers = [normalize_header(cell.value) for cell in sheet[1]]
    columns = {name: idx + 1 for idx, name in enumerate(headers)}
    class_value = sheet.cell(2, columns[CLASS_HEADER]).value if CLASS_HEADER in columns else category.class_code
    for row in range(2, sheet.max_row + 1):
        for col in range(1, len(headers) + 1):
            sheet.cell(row, col).value = None
    for offset, article in enumerate(articles):
        row = offset + 2
        for col in range(1, len(headers) + 1):
            cell = sheet.cell(row, col)
            if row > 6:
                cell._style = copy(sheet.cell(2, col)._style)
        for header in ("Артикул", "Расширенный артикул"):
            if header in columns:
                sheet.cell(row, columns[header]).value = article
        if CLASS_HEADER in columns:
            sheet.cell(row, columns[CLASS_HEADER]).value = class_value
    workbook.save(target)


def csv_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "ДА" if value else ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else f"{value:.9f}".rstrip("0").rstrip(".")
    return str(value)


def export_upload_csv(xlsx_path: Path, csv_path: Path, articles: frozenset[str] | None = None) -> int:
    """The "Шаблон" sheet as ETM takes it: Win-1251, ";", a dot in numbers; only rows with an article.

    With `articles`, only those products' rows are written.
    """
    workbook = load_workbook(xlsx_path, data_only=True)
    sheet = workbook[TEMPLATE_SHEET]
    rows = list(sheet.iter_rows(values_only=True))
    headers = [normalize_header(value) for value in rows[0]]
    width = max(idx for idx, name in enumerate(headers) if name) + 1
    article_idx = headers.index("Артикул")
    written = 0
    with csv_path.open("w", encoding=CSV_ENCODING, errors="replace", newline="") as handle:
        writer = csv.writer(handle, delimiter=";", lineterminator="\r\n")
        writer.writerow(headers[:width])
        for row in rows[1:]:
            article = normalize_article(row[article_idx])
            if article and (articles is None or article in articles):
                writer.writerow([csv_text(value) for value in row[:width]])
                written += 1
    return written


UPLOAD_DIR = "Загрузить в ЭТМ"
CHECK_DIR = "На проверку"
TECH_DIR = "Служебное"
SUMMARY_NAME = "Итог.xlsx"
MISSING_STATUSES = {"missing_value", "value_not_allowed"}
CHECK_STATUSES = {"filled_suggested", "filled_check"}
TEMPLATE_FILE_GLOB = "upload_goods_category_template_*.xlsx"


def parse_articles(text: str) -> list[str]:
    """Articles typed or pasted by the user: spaces, commas, semicolons or new lines between them."""
    seen: dict[str, None] = {}
    for part in re.split(r"[\s,;]+", text or ""):
        article = normalize_article(part.strip("'\""))  # "'586772" as Excel copies a text cell
        if article:
            seen.setdefault(article, None)
    return list(seen)


def collect_templates(folders: list[Path], templates_dir: Path) -> list[Path]:
    """Copy class templates the user downloaded (to the inbox or "Загрузки") into the templates folder.

    A file already there is replaced only by a newer download. Returns the copied files.
    """
    templates_dir.mkdir(parents=True, exist_ok=True)
    copied: list[Path] = []
    for folder in folders:
        if not folder.is_dir() or folder.resolve() == templates_dir.resolve():
            continue
        for path in sorted(folder.glob(TEMPLATE_FILE_GLOB)):
            if path.name.startswith("~$"):
                continue
            target = templates_dir / path.name
            try:
                if target.exists() and target.stat().st_mtime >= path.stat().st_mtime:
                    continue
                shutil.copy2(path, target)
            except OSError:
                continue
            copied.append(target)
    return copied


def upload_file_name(made: tuple[FillResult, Path] | None, article: str) -> str:
    if not made:
        return ""
    result, csv_path = made
    if result.check_path and article in result.check_articles:
        return f"{CHECK_DIR}/{result.check_path.name}"
    return f"{UPLOAD_DIR}/{csv_path.name}"


def field_label(header: str) -> str:
    return header.removeprefix(CONFIG_PREFIX).strip()


def card_status(run: CategoryRun, article: str, missing: int) -> str:
    choice = run.classes.get(article)
    if choice is None:
        return "Класс не подобран: укажите класс вручную"
    made = run.results.get(choice.code)
    if made is None:
        return "Нет шаблона класса: скачайте его"
    result, _csv = made
    if result.check_path and article in result.check_articles:
        status = "Проверить перед загрузкой (лист «Проверить»)"
    else:
        status = "Можно загружать"
    if missing:
        status += f"; не заполнено полей: {missing} (лист «Дозаполнить»)"
    return status


def write_summary(run: CategoryRun, products: dict[str, ProductRecord], library: dict[str, Path]) -> None:
    """One workbook to look at: what to do, every product's state, empty fields, values to check."""
    missing_by_article = Counter(row.article for row in run.missing_rows)
    workbook = Workbook()
    todo = workbook.active
    todo.title = "Что делать"
    for line in todo_lines(run):
        todo.append([line])
    todo.column_dimensions["A"].width = 120

    sheet = workbook.create_sheet("Товары")
    sheet.append(["Артикул", "Название", "81 класс", "Как выбран класс", "Состояние", "Файл для загрузки"])
    for article in sorted(products):
        product = products[article]
        choice = run.classes.get(article)
        code = choice.code if choice else ""
        basis = choice.basis if choice else "класс не найден, укажите вручную"
        if choice and not choice.sure:
            basis += " — проверьте класс"
        sheet.append(
            [
                article,
                clean_product_name(product.full_name or product.name),
                code,
                basis,
                card_status(run, article, missing_by_article[article]),
                upload_file_name(run.results.get(code), article),
            ]
        )
    for row in sheet.iter_rows(min_row=2, min_col=6, max_col=6):
        cell = row[0]
        if cell.value:
            cell.hyperlink = str(cell.value)  # relative to Итог.xlsx: a click opens the file
            cell.style = "Hyperlink"
    for article in run.not_found:
        sheet.append([article, "", "", "", "Артикула нет в базе", ""])

    empty = workbook.create_sheet("Дозаполнить")
    empty.append(["Артикул", "Название", "81 класс", "Поле", "Почему пусто", "Допустимые значения"])
    for row in run.missing_rows:
        product = products.get(row.article)
        choice = run.classes.get(row.article)
        allowed = run.allowed.get((choice.code if choice else "", canonical(field_label(row.column))), ())
        empty.append(
            [
                row.article,
                clean_product_name(product.full_name or product.name) if product else "",
                choice.code if choice else "",
                field_label(row.column),
                row.note,
                ", ".join(allowed[:15]) + (" …" if len(allowed) > 15 else ""),
            ]
        )

    check = workbook.create_sheet("Проверить")
    check.append(["Артикул", "Название", "81 класс", "Поле", "Значение", "Почему проверить"])
    for row in run.check_rows:
        product = products.get(row.article)
        choice = run.classes.get(row.article)
        check.append(
            [
                row.article,
                clean_product_name(product.full_name or product.name) if product else "",
                choice.code if choice else "",
                field_label(row.column),
                row.value,
                row.note,
            ]
        )
    for article, choice in sorted(run.classes.items()):
        if not choice.sure and choice.code in run.results:
            product = products.get(article)
            check.append(
                [
                    article,
                    clean_product_name(product.full_name or product.name) if product else "",
                    choice.code,
                    "81 класс",
                    choice.code,
                    f"класс подобран неуверенно ({choice.basis})",
                ]
            )

    templates = workbook.create_sheet("Нет шаблона")
    templates.append(["81 класс", "Товаров", "Например", "Что сделать"])
    for code, members in sorted(run.missing_templates.items()):
        examples = "; ".join(
            clean_product_name(products[article].full_name or products[article].name) for article in members[:3]
        )
        templates.append(
            [
                code,
                len(members),
                examples,
                f"Скачайте в личном кабинете ЭТМ шаблон класса {code} (upload_goods_category_template_{code}.xlsx) "
                "в «Загрузки» или папку входящих и нажмите «Ещё раз для тех же товаров»",
            ]
        )

    widths = {"A": 14, "B": 60, "C": 12, "D": 36, "E": 40, "F": 60}
    for page in (sheet, empty, check):
        page.freeze_panes = "A2"
        page.auto_filter.ref = page.dimensions
        for letter, width in widths.items():
            page.column_dimensions[letter].width = width
    for letter, width in {"A": 12, "B": 10, "C": 70, "D": 90}.items():
        templates.column_dimensions[letter].width = width
    workbook.save(run.summary_path)


def needs_for_managers(run: CategoryRun, products: dict[str, ProductRecord], series: dict[str, str]) -> list[Need]:
    """The empty fields as questions for the product managers, with ETM's list where it has one."""
    needs = []
    for row in run.missing_rows:
        product = products.get(row.article)
        choice = run.classes.get(row.article)
        allowed = run.allowed.get((choice.code if choice else "", canonical(field_label(row.column))), ())
        needs.append(
            Need(
                article=row.article,
                name=clean_product_name(product.full_name or product.name) if product else "",
                series=series.get(row.article, ""),
                key=row.column,
                label=field_label(row.column),
                kind="A" if allowed else "",
                allowed=allowed,
            )
        )
    return needs


def fill_new_products(
    db_path: Path,
    articles: list[str],
    templates_dir: Path,
    output_dir: Path,
    rules_path: Path | None = None,
    manual_values_path: Path | None = None,
    codes_path: Path | None = None,
    template: PartnerTemplate | None = None,
    answers_file: Path | None = None,
) -> CategoryRun:
    """Cards for the products: one folder with Итог.xlsx, the files to upload, the ones to check first
    and the working files. Answers product managers sent back fill what the base lacks."""
    template = template or load_partner_template("etm")
    products = fetch_products(db_path, articles)
    output_dir.mkdir(parents=True, exist_ok=True)
    upload_dir, check_dir, tech_dir = (output_dir / name for name in (UPLOAD_DIR, CHECK_DIR, TECH_DIR))
    run = CategoryRun(output_dir=output_dir, summary_path=output_dir / SUMMARY_NAME)
    run.articles = list(articles)
    run.not_found = sorted(set(articles) - set(products))
    library = category_library(templates_dir)
    run.classes = suggest_classes(db_path, products, frozenset(library))
    run.unclassified = sorted(article for article in products if article not in run.classes)
    codes = load_codes_file(codes_path) if codes_path else {}
    run.codes_file = codes_path
    answers = load_answers(answers_file)
    defaults = load_series_defaults(template.code)
    series = series_names(db_path, sorted(products))

    by_class: dict[str, list[str]] = {}
    for article, choice in run.classes.items():
        by_class.setdefault(choice.code, []).append(article)
    for code, members in sorted(by_class.items()):
        members.sort()
        if code not in library:
            run.missing_templates[code] = members
            continue
        category = read_category_template(library[code])
        for header_key, values in category.dropdowns.items():
            run.allowed[(code, header_key)] = values
        tech_dir.mkdir(parents=True, exist_ok=True)
        input_path = tech_dir / f"новинки_{code}.xlsx"
        prepare_input(category, members, input_path)
        member_set = set(members)
        given = {
            (article, key): (value, "pm_answer") for (article, key), value in answers.items() if article in member_set
        }
        fallback = {**series_default_values(defaults, category, members, products, series), **codes, **given}
        result = fill_xlsx_template(
            db_path,
            input_path,
            tech_dir,
            rules_path,
            manual_values_path,
            template=replace(template, dropdowns=category.dropdowns),
            fallback_values=fallback,
        )
        # The name is made unique against the working files, which every class keeps.
        name = upload_csv_path(tech_dir, "new", (series.get(article, "") for article in members), code).name
        upload_dir.mkdir(parents=True, exist_ok=True)
        csv_path = upload_dir / name
        xlsx_path = result.output_path.rename(tech_dir / f"{csv_path.stem}{result.output_path.suffix}")
        report_path = result.report_path.rename(report_path_for(tech_dir / csv_path.name))
        result = replace(result, output_path=xlsx_path, report_path=report_path)
        held = result.check_articles | {article for article in members if not run.classes[article].sure}
        safe_rows = export_upload_csv(result.output_path, csv_path, frozenset(members) - held)
        if not safe_rows:
            csv_path.unlink()  # nothing ready: no empty file to upload by mistake
        check_path = None
        if held:
            check_dir.mkdir(parents=True, exist_ok=True)
            check_path = check_path_for(check_dir / csv_path.name)
            export_upload_csv(result.output_path, check_path, held)
        result = replace(result, check_path=check_path, check_rows=len(held), safe_rows=safe_rows, check_articles=frozenset(held))
        input_path.unlink()
        run.results[code] = (result, csv_path)
        run.missing_rows.extend(row for row in result.rows if row.status in MISSING_STATUSES and row.column)
        run.check_rows.extend(row for row in result.rows if row.status in CHECK_STATUSES and row.column)
    run.needs = needs_for_managers(run, products, series)
    write_summary(run, products, library)
    save_needs(run)
    return run


NEEDS_FILE = "запросы_продактам.json"
RUN_FOLDER_PREFIX = "ЭТМ карточки "


def save_needs(run: CategoryRun) -> None:
    """The empty fields of the run, so the managers can be asked after the program was closed."""
    path = run.output_dir / TECH_DIR / NEEDS_FILE
    path.parent.mkdir(parents=True, exist_ok=True)
    items = [{**asdict(need), "allowed": list(need.allowed)} for need in run.needs]
    atomic_write_json(path, {"version": 1, "needs": items})


def last_run_needs(output_dir: Path) -> tuple[Path, list[Need]] | None:
    """The newest cards run in the output folder and its empty fields, or None."""
    runs = sorted(
        (folder for folder in output_dir.glob(f"{RUN_FOLDER_PREFIX}*") if (folder / TECH_DIR / NEEDS_FILE).is_file()),
        key=lambda folder: folder.stat().st_mtime,
    )
    if not runs:
        return None
    try:
        data = json.loads((runs[-1] / TECH_DIR / NEEDS_FILE).read_text(encoding="utf-8"))
        needs = [Need(**{**item, "allowed": tuple(item.get("allowed", ()))}) for item in data.get("needs", [])]
    except (OSError, ValueError, TypeError):
        return None
    return runs[-1], needs


def short_lines(run: CategoryRun) -> list[str]:
    """The few numbers for the message after a run; the details are in Итог.xlsx."""
    ready = sum(result.safe_rows for result, _ in run.results.values())
    held = sum(result.check_rows for result, _ in run.results.values())
    lines = [f"Товаров: {len(run.articles)}", f"Можно загружать сразу: {ready}"]
    if held:
        lines.append(f"Посмотреть перед загрузкой: {held}")
    if run.missing_rows:
        lines.append(f"Пустых полей: {len(run.missing_rows)} у {len({row.article for row in run.missing_rows})} товаров")
    waiting = sum(len(members) for members in run.missing_templates.values())
    if waiting:
        lines.append(f"Ждут шаблона класса: {waiting} (классов: {len(run.missing_templates)})")
    if run.unclassified:
        lines.append(f"Класс не подобран: {len(run.unclassified)}")
    if run.not_found:
        lines.append(f"Нет в базе: {len(run.not_found)}")
    lines.append("")
    lines.append("Что делать дальше — на первом листе Итог.xlsx (он уже открыт).")
    return lines


def todo_lines(run: CategoryRun) -> list[str]:
    """The "Что делать" sheet: the steps in order, only those that apply."""
    ready = [(result, path) for result, path in run.results.values() if result.safe_rows]
    held = sum(result.check_rows for result, _ in run.results.values())
    lines = []
    step = 1
    if ready:
        lines.append(
            f"{step}. Загрузите в ЭТМ файлы из папки «{UPLOAD_DIR}»: {len(ready)} файлов, "
            f"{sum(result.safe_rows for result, _ in ready)} карточек. Один файл — один класс ЭТМ."
        )
        step += 1
    if held:
        lines.append(
            f"{step}. {held} карточек в папке «{CHECK_DIR}»: посмотрите лист «Проверить» (что и почему). "
            "Если всё верно, загрузите и эти файлы."
        )
        step += 1
    if run.missing_templates:
        codes = ", ".join(sorted(run.missing_templates))
        count = sum(len(members) for members in run.missing_templates.values())
        lines.append(
            f"{step}. Для {count} товаров нет шаблона класса ({codes}): скачайте шаблоны в личном кабинете ЭТМ "
            "в «Загрузки» или папку входящих (лист «Нет шаблона») и нажмите «Ещё раз для тех же товаров» — "
            "программа сама заберёт их в папку шаблонов."
        )
        step += 1
    if run.missing_rows:
        products = len({row.article for row in run.missing_rows})
        lines.append(
            f"{step}. Пустых полей: {len(run.missing_rows)} у {products} товаров (лист «Дозаполнить»). "
            "Нажмите «Запросить у продактов»: каждому продакту уйдёт Excel с этими полями. Когда ответы вернутся, "
            "примите их и нажмите «Ещё раз для тех же товаров» — карточки соберутся с ответами."
        )
        step += 1
    if run.unclassified:
        lines.append(f"{step}. Класс не подобран у {len(run.unclassified)} товаров (лист «Товары»): укажите класс вручную.")
        step += 1
    if run.not_found:
        lines.append(f"{step}. Нет в базе: {len(run.not_found)} артикулов. Обновите базу из прайса или проверьте артикулы.")
        step += 1
    if not lines:
        lines.append("Делать ничего не нужно.")
    lines.append("")
    lines.append(f"Папка «{TECH_DIR}» — рабочие копии и подробные отчёты, обычно туда смотреть не нужно.")
    if run.codes_file:
        lines.append(f"Коды ТН ВЭД и ОКПД2 взяты из файла {run.codes_file.name}.")
    else:
        lines.append("Файла ЧЗ с кодами не нашлось: сохраните его в «Загрузки», и программа возьмёт самый свежий.")
    return lines


def run_lines(run: CategoryRun) -> list[str]:
    """What the user needs to know after a run, in plain words."""
    lines = []
    for code, (result, csv_path) in sorted(run.results.items()):
        line = f"Класс {code}: {result.found_articles} товаров, можно грузить сразу {result.safe_rows}"
        if result.safe_rows:
            line += f" — файл {csv_path.name}"
        if result.check_path:
            line += f"; на проверку {result.check_rows} — файл {result.check_path.name}"
        if result.missing_cells:
            line += f"; не заполнено полей: {result.missing_cells}"
        lines.append(line)
    for code, members in sorted(run.missing_templates.items()):
        lines.append(
            f"Нет шаблона класса {code} ({len(members)} товаров): скачайте его в личном кабинете ЭТМ "
            "в «Загрузки» и нажмите «Ещё раз для тех же товаров»."
        )
    if run.unclassified:
        lines.append(f"Не удалось подобрать класс: {len(run.unclassified)} товаров, список в итоге.")
    if run.not_found:
        lines.append(f"Нет в базе: {len(run.not_found)} артикулов.")
    if run.missing_rows:
        lines.append(f"Пустых полей: {len(run.missing_rows)} — можно спросить у продактов (кнопка «Запросить у продактов»).")
    if run.codes_file:
        lines.append(f"Коды ТН ВЭД и ОКПД2 взяты из файла {run.codes_file.name}.")
    else:
        lines.append("Файла ЧЗ с кодами не нашлось: сохраните его в «Загрузки», и программа возьмёт самый свежий.")
    lines.append(f"Главный файл: {run.summary_path.name} в папке {run.output_dir.name}")
    return lines
