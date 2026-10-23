"""RAEK archive: the manager's zip in, the same zip with empty cells filled out.

The RAEK manager sends a zip of Excel files, one per kind of data (country,
EAN, ТН ВЭД, ОКПД2, series, certificate, links, ETIM features...). Every sheet
has the same first columns (RAEC_ID, our article in "Код поставщика", name...)
and the data columns after them. The program fills only empty data cells and
keeps everything else as RAEK sent it, then writes:

- <archive>_заполнено_<date>.zip with the same file names, ready to send back;
- <archive>_отчёт_<date>.xlsx: what was filled, what to check, what is empty
  and why, and the ETIM feature ↔ database attribute pairs the program used.

Values follow the format of RAEK's own latest files: text, a decimal comma,
ranges as "24...24" ("-...-" when the range does not apply, e.g. a DC voltage
for an AC relay), yes/no as "true"/"false", listed values in RAEK's wording.
"""
from __future__ import annotations

import datetime as dt
import json
import re
import zipfile
from collections import Counter
from contextlib import closing
from dataclasses import dataclass, field
from importlib import resources
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from . import partners
from .atomic_files import atomic_write_json
from .certificates import is_eaeu_document, number_key, register_documents, table_dates
from .db import ProductRecord, connect, fetch_products
from .etm_categories import load_codes_file
from .filler import normalize_article
from .pm_requests import NAME_QUESTION, Need, load_answers
from .site_info import FOUND as SITE_FOUND
from .site_info import SiteCache, SiteRecord, SiteResult, site_lines, update_site_cache
from .upload_names import series_names
from .value_matching import (
    LOGICAL_NO,
    LOGICAL_YES,
    Characteristic,
    canonical,
    current_kind_from_text,
    current_kind_from_value,
    format_number,
    looks_like_file,
    match_value,
    name_candidates,
    only_number,
)

HEADER_START = "RAEC_ID"
ARTICLE_HEADER = "Код поставщика"
NAME_HEADER = "Название от поставщика"
FIRST_DATA_COLUMN = 6  # A..F are RAEK's own columns
FEATURE_CODE_RE = re.compile(r"^[ER]F\d{6}$")
FEATURE_LABEL_RE = re.compile(r"^\s*([NALR])\s*-\s*(.+)$", re.DOTALL)
ALLOWED_RE = re.compile(r"^\s*[ER]V\d{6}\s*-\s*(.+?)\s*$")
CHZ_HEADERS = {
    "part number": "article",
    "tariff code": "tnved",
    "окпд 2 (9 зн.)": "okpd2",
    "шк / gtin": "ean",
    "серия": "series",
    "сертификат 1": "cert_number",
    "ссылка на реестр 1": "cert_link",
    "to be marked with honest sign (y/n)": "honest_sign",
}
CERT_NUMBER_ATTRIBUTES = ("Номер сертификата",)
CERT_FROM_ATTRIBUTES = ("Дата начала",)
CERT_TO_ATTRIBUTES = ("Дата окончания",)
CURRENT_KIND_ATTRIBUTES = ("Род тока включения", "Род тока", "Род тока катушки управления", "Тип напряжения (тока)")
DIMENSIONS_RE = re.compile(r"габарит", re.IGNORECASE)
DIMENSION_LETTERS = {"ш": "width", "в": "height", "г": "depth", "д": "length"}
DIMENSION_FEATURES = {"EF000008": "width", "EF000040": "height", "EF000049": "depth"}
# RAEK abbreviates ETIM names; our database has them in full.
ABBREVIATIONS = {
    "номин": "номинальный",
    "макс": "максимальная",
    "мин": "минимальная",
    "вход": "входных",
    "выход": "выходных",
    "перемен": "переменного",
    "постоян": "постоянного",
    "кол-во": "количество",
    "ac": "переменного тока",
    "dc": "постоянного тока",
}
ABBREVIATION_RE = re.compile(r"(?<![\w-])(" + "|".join(re.escape(item) for item in ABBREVIATIONS) + r")\.?(?![\w-])", re.IGNORECASE)

FILLED, CHECK, EMPTY = "filled", "check", "empty"
STATUS_LABELS = {FILLED: "Заполнено", CHECK: "Заполнено, проверить", EMPTY: "Нечем заполнить"}


@dataclass(frozen=True)
class Cell:
    status: str
    value: object = None
    note: str = ""


@dataclass
class RaekRow:
    file: str
    sheet: str
    raec_id: str
    article: str
    name: str
    column: str
    status: str
    value: str
    note: str


@dataclass
class RaekResult:
    output_path: Path
    report_path: Path
    files: int = 0
    products: int = 0
    filled: int = 0
    check: int = 0
    empty: int = 0
    kept: int = 0
    rows: list[RaekRow] = field(default_factory=list)
    names: list[list[object]] = field(default_factory=list)
    skipped_files: list[str] = field(default_factory=list)
    codes_file: Path | None = None
    site: SiteResult | None = None
    needs: list[Need] = field(default_factory=list)  # what to ask product managers


@dataclass(frozen=True)
class Feature:
    code: str
    kind: str  # N number, A listed value, L yes/no, R range
    name: str
    allowed: tuple[str, ...]


@dataclass
class Sources:
    products: dict[str, ProductRecord]
    chz: dict[tuple[str, str], str]
    series: dict[str, str]
    coded: dict[str, dict[str, str]]  # article -> ETIM feature code -> value (ETIM workbook import)
    settings: dict
    learned: dict
    site: dict[str, SiteRecord] = field(default_factory=dict)  # article -> what ensmas.ru answered
    cert_dates: dict[str, tuple] = field(default_factory=dict)  # document key -> (from, to) from the certificates table
    answers: dict[tuple[str, str], str] = field(default_factory=dict)  # (article, feature code or column) -> manager's answer
    chz_articles: frozenset[str] = field(init=False)

    def __post_init__(self) -> None:
        self.chz_articles = frozenset(article for article, _ in self.chz)


def load_settings() -> dict:
    return json.loads(resources.files(partners).joinpath("raek.json").read_text(encoding="utf-8"))


def learned_path(rules_path: Path | None) -> Path | None:
    return rules_path.with_name("raek_learned.json") if rules_path else None


def load_learned(rules_path: Path | None) -> dict:
    path = learned_path(rules_path)
    if path is None or not path.exists():
        return {"name_aliases": {}, "name_rejects": {}}
    data = json.loads(path.read_text(encoding="utf-8"))
    return {"name_aliases": data.get("name_aliases", {}), "name_rejects": data.get("name_rejects", {})}


def text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def comma(number: float) -> str:
    return format_number(number).replace(".", ",")


SIGNED_NUMBER_RE = re.compile(r"(?<![\d.,])[-+]?\d+(?:[.,]\d+)?")  # "-40…+70" is signed, "525-690" is not


def numbers_in(value: str) -> list[float]:
    return [float(item.replace(",", ".")) for item in SIGNED_NUMBER_RE.findall(value.replace("−", "-"))]


def expand(name: str) -> str:
    return ABBREVIATION_RE.sub(lambda found: ABBREVIATIONS[found.group(1).casefold()], name)


def coded_attributes(db_path: Path, articles: list[str]) -> dict[str, dict[str, str]]:
    """Values the ETIM workbook import stored with their ETIM feature code."""
    if not articles:
        return {}
    placeholders = ",".join("?" for _ in articles)
    result: dict[str, dict[str, str]] = {}
    with closing(connect(db_path)) as conn:
        rows = conn.execute(
            f"""
            SELECT article, attribute_code, COALESCE(value_normalized, value_raw, '')
            FROM product_attribute_values
            WHERE article IN ({placeholders}) AND COALESCE(attribute_code, '') <> ''
            ORDER BY confidence DESC, id DESC
            """,
            articles,
        ).fetchall()
    for article, code, value in rows:
        if str(value).strip():
            result.setdefault(str(article), {}).setdefault(str(code).strip(), str(value).strip())
    return result


def read_features(sheet, header_row: int) -> dict[int, Feature]:
    """ETIM feature columns of a required_features sheet: code row, "<обязательное>\\nN - name" row, allowed values row."""
    features: dict[int, Feature] = {}
    if header_row < 3:
        return features
    for col in range(FIRST_DATA_COLUMN + 1, sheet.max_column + 1):
        code = text(sheet.cell(header_row, col).value)
        if not FEATURE_CODE_RE.match(code):
            continue
        label = text(sheet.cell(header_row - 1, col).value)
        label = "\n".join(line for line in label.splitlines() if not line.strip().startswith("<"))
        found = FEATURE_LABEL_RE.match(label)
        kind, name = (found.group(1), found.group(2).strip()) if found else ("A", label)
        allowed = tuple(
            item.group(1)
            for line in text(sheet.cell(header_row - 2, col).value).splitlines()
            if (item := ALLOWED_RE.match(line))
        )
        features[col] = Feature(code, kind, name, allowed)
    return features


def product_name(product: ProductRecord) -> str:
    return product.full_name or product.name


def current_kind(product: ProductRecord) -> str:
    for name in CURRENT_KIND_ATTRIBUTES:
        kind = current_kind_from_value((product.attributes or {}).get(name, ""))
        if kind:
            return kind
    return current_kind_from_text(product_name(product))[0]


def dimensions(product: ProductRecord) -> dict[str, float]:
    """Width/height/depth from "Габаритные ... (ШхВхГ), мм" = "27,8×21,8×44"."""
    for attribute, value in (product.attributes or {}).items():
        if not DIMENSIONS_RE.search(attribute):
            continue
        order = re.search(r"\(([ШВГДшвгд])\s*[xх×*]\s*([ШВГДшвгд])\s*[xх×*]\s*([ШВГДшвгд])\)", attribute)
        values = numbers_in(value)
        if order and len(values) == 3:
            return {DIMENSION_LETTERS[letter.casefold()]: number for letter, number in zip(order.groups(), values)}
    return {}


def certificate_type(number: str) -> str:
    """RAEK's wording: "Декларация ЕАС" for "...Д-...", "Сертификат ЕАС" for "...С-..."."""
    compact = number.replace(" ", "").upper()
    if re.search(r"Д-", compact):
        return "Декларация ЕАС"
    if re.search(r"[СC]-", compact):
        return "Сертификат ЕАС"
    return ""


def mandatory_certification(article: str, sources: Sources) -> Cell:
    """"Подлежит обязательному подтверждению соответствия" from the ЧЗ file's "Сертификат 1" column."""
    document = sources.chz.get((article, "cert_number"), "")
    if not document:
        site = site_record(article, sources)
        if site and is_eaeu_document(site.cert_number):
            return Cell(CHECK, "Да", f"В файле ЧЗ данных нет, на сайте ensmas.ru у товара {site.cert_number}")
    if not sources.chz:
        return Cell(EMPTY, note="Нет файла ЧЗ: обязательна ли сертификация, неизвестно")
    if article not in sources.chz_articles:
        return Cell(EMPTY, note="Артикула нет в файле ЧЗ: обязательна ли сертификация, неизвестно")
    if not document:
        return Cell(EMPTY, note="В файле ЧЗ у артикула не указан сертификат")
    if is_eaeu_document(document):
        return Cell(FILLED, "Да")
    if "не треб" in document.casefold():
        return Cell(FILLED, "Нет")
    return Cell(CHECK, "Нет", f"В файле ЧЗ вместо сертификата ЕАЭС документ «{document}», похоже на отказное письмо")


def site_record(article: str, sources: Sources) -> SiteRecord | None:
    record = sources.site.get(article)
    return record if record and record.status == SITE_FOUND else None


def site_missing_note(article: str, sources: Sources, what: str) -> str:
    record = sources.site.get(article)
    if record is None:
        return f"{what}: сайт ensmas.ru по этому товару ещё не проверялся"
    if record.status != SITE_FOUND:
        return f"{what}: {record.note or 'товара нет на сайте ensmas.ru'}"
    return f"{what}: на странице товара на ensmas.ru этого нет"


def site_column(article: str, sources: Sources, attribute: str, what: str) -> Cell:
    record = site_record(article, sources)
    value = getattr(record, attribute) if record else ""
    return Cell(FILLED, value) if value else Cell(EMPTY, note=site_missing_note(article, sources, what))


def same_certificate(left: str, right: str) -> bool:
    def norm(value: str) -> str:
        return re.sub(r"[^0-9A-ZА-Я]", "", value.upper().replace("C", "С").replace("A", "А").replace("B", "В").replace("H", "Н").replace("P", "Р"))

    return bool(left and right) and norm(left) == norm(right)


def as_datetime(value: dt.date | None) -> dt.datetime | None:
    return dt.datetime.combine(value, dt.time()) if value else None


def as_date(value: str) -> dt.datetime | None:
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d.%m.%Y"):
        try:
            return dt.datetime.strptime(value.strip(), fmt)
        except ValueError:
            continue
    return None


def ean13(product: ProductRecord, chz_value: str) -> Cell:
    """The product's own barcode from the ЧЗ file. A pack barcode is a different code and is never used."""
    if re.fullmatch(r"\d{13}", chz_value):
        return Cell(FILLED, chz_value)
    return Cell(EMPTY, note="Нет штрих-кода товара в файле ЧЗ")


ASKED_COLUMNS = {canonical(name) for name in ("Штрих-код EAN", "Код ТН ВЭД", "Код ОКПД2", "Серия")}
NO_VALUE = "В базе нет значения"


def column_value(header: str, article: str, product: ProductRecord, sources: Sources) -> Cell | None:
    """Plain (non-ETIM) columns. None for a column the program does not know."""
    chz = sources.chz
    key = canonical(header)
    constants = {canonical(name): value for name, value in sources.settings.get("constants", {}).items()}
    if key in constants:
        return Cell(FILLED, constants[key])
    if key in {canonical(name) for name in sources.settings.get("leave_empty", [])}:
        return Cell(FILLED, None)
    answer = sources.answers.get((article, key), "")
    if key in ASKED_COLUMNS and answer:
        return Cell(FILLED, answer, "Ответ продакта")
    if key == canonical("Штрих-код EAN"):
        return ean13(product, chz.get((article, "ean"), ""))
    if key == canonical("Код ТН ВЭД"):
        value = chz.get((article, "tnved")) or product.tnved_code
        return Cell(FILLED, value) if value else Cell(EMPTY, note="Нет кода ни в файле ЧЗ, ни в базе")
    if key == canonical("Код ОКПД2"):
        value = chz.get((article, "okpd2")) or (product.okpd2_code if product.okpd2_code.count(".") == 3 else "")
        return Cell(FILLED, value) if value else Cell(EMPTY, note="Нет 9-значного кода ни в файле ЧЗ, ни в базе")
    if key == canonical("Серия"):
        value = sources.series.get(article) or chz.get((article, "series"), "")
        return Cell(FILLED, value) if value else Cell(EMPTY, note="В базе нет серии")
    if key == canonical("Подлежит обязательному подтверждению соответствия"):
        return mandatory_certification(article, sources)
    if key == canonical("Подлежит маркировке Честный Знак"):
        mark = chz.get((article, "honest_sign"), "").casefold()
        if mark in {"yes", "y", "да"}:
            return Cell(FILLED, "Да")
        if mark in {"no", "n", "нет"}:
            return Cell(FILLED, "Нет")
        return Cell(EMPTY, note="В файле ЧЗ нет отметки о маркировке")
    if key == canonical("Изображение"):
        return site_column(article, sources, "image_url", "Нет картинки")
    if key == canonical("Ссылка на сайте поставщика"):
        return site_column(article, sources, "url", "Нет ссылки")
    cert = certificate(article, product, sources)
    if key in cert:
        return cert[key]
    hints = {canonical(name): note for name, note in sources.settings.get("not_yet", {}).items()}
    if key in hints:
        return Cell(EMPTY, note=hints[key])
    return None


def certificate(article: str, product: ProductRecord, sources: Sources) -> dict[str, Cell]:
    attributes = product.attributes or {}
    in_chz = sources.chz.get((article, "cert_number"), "")
    from_chz = in_chz if is_eaeu_document(in_chz) else ""
    from_db = next((attributes[name] for name in CERT_NUMBER_ATTRIBUTES if attributes.get(name)), "")
    site = site_record(article, sources)
    from_site = site.cert_number if site and is_eaeu_document(site.cert_number) and not in_chz else ""
    number = from_chz or from_db or from_site
    if not number:
        reason = f"Сертификата ЕАЭС нет: в файле ЧЗ «{in_chz}»" if in_chz else "Нет номера сертификата ни в файле ЧЗ, ни в базе"
        missing = Cell(EMPTY, note=reason)
        return {canonical(name): missing for name in ("Тип сертификата", "Номер сертификата", "Сертификат от", "Сертификат до", "Ссылка на реестр")}
    table_start, table_end = sources.cert_dates.get(number_key(number), (None, None))
    dates_match = not from_chz or not from_db or same_certificate(from_chz, from_db)
    start = as_date(next((attributes[name] for name in CERT_FROM_ATTRIBUTES if attributes.get(name)), "")) if dates_match else None
    end = as_date(next((attributes[name] for name in CERT_TO_ATTRIBUTES if attributes.get(name)), "")) if dates_match else None
    if table_start or table_end:
        start, end = as_datetime(table_start), as_datetime(table_end)
    no_dates = "Дат нет: впишите их один раз в таблицу сертификатов (вкладка «Справочники», «Сертификаты» → «Открыть таблицу»)"
    link = sources.chz.get((article, "cert_link"), "")
    kind = certificate_type(number)
    return {
        canonical("Тип сертификата"): Cell(FILLED, kind) if kind else Cell(EMPTY, note=f"Не понял вид документа по номеру {number}"),
        canonical("Номер сертификата"): Cell(FILLED, number),
        canonical("Сертификат от"): Cell(FILLED, start) if start else Cell(EMPTY, note=no_dates),
        canonical("Сертификат до"): Cell(FILLED, end) if end else Cell(EMPTY, note=no_dates),
        canonical("Ссылка на реестр"): Cell(FILLED, link) if link else Cell(EMPTY, note="Ссылки на реестр ФСА нет в файле ЧЗ"),
        canonical("Файл"): Cell(FILLED, None),
    }


class FeatureMatcher:
    """Which database attribute holds an ETIM feature, decided once per feature and sheet."""

    def __init__(self, sources: Sources, products: list[ProductRecord]) -> None:
        self.sources = sources
        self.attributes = [product.attributes or {} for product in products]
        self.cache: dict[str, tuple[str, bool]] = {}

    def attribute_for(self, feature: Feature) -> tuple[str, bool]:
        """(attribute name, sure). Sure for verified pairs and same names; a similar name needs a check."""
        if feature.code in self.cache:
            return self.cache[feature.code]
        verified = list(self.sources.settings.get("name_aliases", {}).get(feature.code, []))
        verified += self.sources.learned["name_aliases"].get(feature.code, [])
        rejected = set(self.sources.settings.get("name_rejects", {}).get(feature.code, []))
        rejected |= set(self.sources.learned["name_rejects"].get(feature.code, []))
        present = {name for attributes in self.attributes for name, value in attributes.items() if value}
        choice = next(((name, True) for name in verified if name in present), None)
        if choice is None:
            candidates = [
                item for item in name_candidates(expand(feature.name), self.attributes) if item[0] not in rejected
            ]
            if candidates:
                name = candidates[0][0]
                sure = canonical(expand(name)) == canonical(expand(feature.name))
                choice = (name, sure)
        self.cache[feature.code] = choice or ("", False)
        return self.cache[feature.code]


def feature_value(feature: Feature, product: ProductRecord, sources: Sources, matcher: FeatureMatcher) -> Cell:
    article = product.article
    dims = dimensions(product)
    if feature.code in DIMENSION_FEATURES and DIMENSION_FEATURES[feature.code] in dims:
        return Cell(FILLED, comma(dims[DIMENSION_FEATURES[feature.code]]))
    raw, sure, source = "", True, ""
    coded = sources.coded.get(article, {}).get(feature.code, "")
    if coded:
        raw, source = coded, "ETIM"
    else:
        attribute, sure = matcher.attribute_for(feature)
        raw = (product.attributes or {}).get(attribute, "") if attribute else ""
        source = attribute
    kind_needed = range_current_kind(feature.name)
    if feature.kind == "R" and kind_needed:
        own = current_kind(product)
        if own and kind_needed not in own.split("/"):
            return Cell(FILLED, "-...-")
    if not raw or looks_like_file(raw):
        answer = sources.answers.get((article, feature.code), "")
        if not answer:
            return Cell(EMPTY, note=NO_VALUE + (f" (искали «{source}»)" if source else ""))
        cell = format_feature(feature, answer, sources)
        return Cell(cell.status, cell.value, "Ответ продакта" + (f"; {cell.note}" if cell.note else ""))
    cell = format_feature(feature, raw, sources)
    if cell.status == FILLED and not sure:
        return Cell(CHECK, cell.value, f"Взято из «{source}» по похожему названию")
    return cell


def range_current_kind(name: str) -> str:
    words = canonical(expand(name))
    if "переменного" in words and "постоянного" not in words:
        return "AC"
    if "постоянного" in words and "переменного" not in words:
        return "DC"
    return ""


def format_feature(feature: Feature, raw: str, sources: Sources) -> Cell:
    if feature.kind == "L":
        key = canonical(raw)
        if key in LOGICAL_YES:
            return Cell(FILLED, "true")
        if key in LOGICAL_NO:
            return Cell(FILLED, "false")
        return Cell(EMPTY, note=f"Значение «{raw}» не похоже на да/нет")
    if feature.kind == "R":
        values = numbers_in(raw)
        if len(values) == 1:
            return Cell(FILLED, f"{comma(values[0])}...{comma(values[0])}")
        if len(values) == 2 and re.search(r"\d\s*(?:-|–|…|\.\.\.)\s*[-+]?\d", raw):
            return Cell(FILLED, f"{comma(values[0])}...{comma(values[1])}")
        return Cell(EMPTY, note=f"Не понял диапазон «{raw}»")
    if feature.kind == "N":
        number = only_number(raw)
        if number is not None:
            return Cell(FILLED, comma(number))
        values = numbers_in(raw)
        if len(values) == 2 and re.search(r"\d\s*(?:-|–|…|\.\.\.)\s*\d", raw) and "макс" in canonical(feature.name):
            return Cell(CHECK, comma(max(values)), f"Взят верхний предел диапазона «{raw}»")
        return Cell(EMPTY, note=f"В базе «{raw}», а нужно одно число")
    if not feature.allowed:
        return Cell(FILLED, raw)
    learned = {canonical(key): value for key, value in sources.settings.get("values", {}).get(feature.code, {}).items()}
    match = match_value(raw, Characteristic(name=feature.name, values=list(feature.allowed)), learned)
    if match.value and match.confident:
        return Cell(FILLED, match.value)
    if match.value:
        return Cell(CHECK, match.value, f"В базе «{raw}», подобрано по похожему")
    hint = f"; похожие: {', '.join(match.candidates)}" if match.candidates else ""
    return Cell(EMPTY, note=f"В базе «{raw}», такого значения нет в списке РАЭК{hint}")


def find_header_row(sheet) -> int | None:
    for row in range(1, min(sheet.max_row, 10) + 1):
        if text(sheet.cell(row, 1).value) == HEADER_START:
            return row
    return None


def fill_sheet(file_name: str, sheet, sources: Sources, result: RaekResult) -> bool:
    header_row = find_header_row(sheet)
    if header_row is None:
        return False
    headers = {col: text(sheet.cell(header_row, col).value) for col in range(1, sheet.max_column + 1)}
    article_col = next((col for col, name in headers.items() if name == ARTICLE_HEADER), 2)
    name_col = next((col for col, name in headers.items() if name == NAME_HEADER), 3)
    features = read_features(sheet, header_row)
    rows = [
        row for row in range(header_row + 1, sheet.max_row + 1) if normalize_article(sheet.cell(row, article_col).value)
    ]
    products = [sources.products[a] for row in rows if (a := normalize_article(sheet.cell(row, article_col).value)) in sources.products]
    matcher = FeatureMatcher(sources, products)
    for row in rows:
        article = normalize_article(sheet.cell(row, article_col).value)
        product = sources.products.get(article)
        for col in range(FIRST_DATA_COLUMN + 1, sheet.max_column + 1):
            header = headers.get(col, "")
            if not header:
                continue
            target = sheet.cell(row, col)
            column = features[col].name if col in features else header
            if text(target.value):
                result.kept += 1
                continue
            if product is None:
                cell = Cell(EMPTY, note="Артикула нет в базе")
            elif col in features:
                cell = feature_value(features[col], product, sources, matcher)
            else:
                cell = column_value(header, article, product, sources) or Cell(EMPTY, note="Программа пока не знает, чем заполнять эту колонку")
            if cell.status in {FILLED, CHECK} and cell.value is None:
                continue  # a column RAEK leaves empty
            if cell.status in {FILLED, CHECK}:
                target.value = cell.value
                if isinstance(cell.value, dt.datetime):
                    target.number_format = "DD.MM.YYYY"
            setattr(result, cell.status, getattr(result, cell.status) + 1)
            if product is not None and cell.status == EMPTY:
                need = manager_need(article, text(sheet.cell(row, name_col).value), header, features.get(col), cell, sources)
                if need:
                    result.needs.append(need)
            if cell.status != FILLED or cell.note:
                result.rows.append(
                    RaekRow(
                        file_name,
                        sheet.title,
                        text(sheet.cell(row, 1).value),
                        article,
                        text(sheet.cell(row, name_col).value),
                        f"{features[col].code} {column}" if col in features else column,
                        STATUS_LABELS[cell.status],
                        text(cell.value.strftime("%d.%m.%Y") if isinstance(cell.value, dt.datetime) else cell.value),
                        cell.note,
                    )
                )
    for col, feature in features.items():
        attribute, sure = matcher.cache.get(feature.code, ("", True))
        if attribute and not sure:
            users = sum(1 for product in products if (product.attributes or {}).get(attribute))
            examples = sorted({(product.attributes or {}).get(attribute, "") for product in products} - {""})[:3]
            result.names.append([sheet.title, feature.code, feature.name, attribute, users, "; ".join(examples), ""])
            series = Counter(series_of(product.article, sources) for product in products if (product.attributes or {}).get(attribute))
            result.needs.append(
                Need(
                    "",
                    f"Характеристика РАЭК «{feature.name}» — в базе «{attribute}», примеры: {'; '.join(examples)}",
                    series.most_common(1)[0][0] if series else "",
                    f"{NAME_QUESTION}{feature.code}:{attribute}",
                    "Это одна и та же характеристика?",
                )
            )
    return True


def series_of(article: str, sources: Sources) -> str:
    return sources.series.get(article) or sources.chz.get((article, "series"), "")


def manager_need(article: str, name: str, header: str, feature: Feature | None, cell: Cell, sources: Sources) -> Need | None:
    """An empty cell a product manager can answer: a feature the database lacks, or a code nobody has."""
    if feature is not None:
        if not cell.note.startswith(NO_VALUE):
            return None
        return Need(article, name, series_of(article, sources), feature.code, feature.name, feature.kind, feature.allowed)
    if canonical(header) not in ASKED_COLUMNS:
        return None
    return Need(article, name, series_of(article, sources), canonical(header), header)


def zip_articles(path: Path) -> list[str]:
    articles: set[str] = set()
    with zipfile.ZipFile(path) as archive:
        for info in archive.infolist():
            if not info.filename.lower().endswith(".xlsx"):
                continue
            workbook = load_workbook(BytesIO(archive.read(info)), read_only=True)
            try:
                for sheet in workbook.worksheets:
                    header: list[str] | None = None
                    for values in sheet.iter_rows(values_only=True):
                        cells = [text(value) for value in values]
                        if header is None:
                            if cells and cells[0] == HEADER_START:
                                header = cells
                            continue
                        idx = header.index(ARTICLE_HEADER) if ARTICLE_HEADER in header else 1
                        if idx < len(cells) and normalize_article(cells[idx]):
                            articles.add(normalize_article(cells[idx]))
            finally:
                workbook.close()
    return sorted(articles)


def output_paths(zip_path: Path, output_dir: Path) -> tuple[Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    stem = f"{zip_path.stem}_заполнено_{dt.date.today().isoformat()}"
    candidate, version = stem, 1
    while (output_dir / f"{candidate}.zip").exists():
        version += 1
        candidate = f"{stem}_v{version}"
    return output_dir / f"{candidate}.zip", output_dir / f"{candidate.replace('_заполнено_', '_отчёт_')}.xlsx"


def fill_raek_zip(
    db_path: Path,
    zip_path: Path,
    output_dir: Path,
    codes_path: Path | None = None,
    rules_path: Path | None = None,
    site_cache: Path | None = None,
    ask_site: bool = False,
    fetch=None,
    progress=None,
    certificates_table: Path | None = None,
    answers_path: Path | None = None,
    cert_folder: Path | None = None,
) -> RaekResult:
    """Fill the archive. With `ask_site`, ensmas.ru is asked first about articles it has no recent answer for."""
    articles = zip_articles(zip_path)
    cache = SiteCache(site_cache)
    site_result = None
    if ask_site:
        options = {"fetch": fetch} if fetch else {}
        site_result = update_site_cache(cache, articles, progress=progress, **options)
    sources = Sources(
        products=fetch_products(db_path, articles),
        chz=load_codes_file(codes_path, CHZ_HEADERS) if codes_path else {},
        series=series_names(db_path, articles),
        coded=coded_attributes(db_path, articles),
        settings=load_settings(),
        learned=load_learned(rules_path),
        site=cache.records,
        cert_dates=table_dates(certificates_table, register_documents(cert_folder)),
        answers=load_answers(answers_path),
    )
    output_path, report_path = output_paths(zip_path, output_dir)
    result = RaekResult(output_path, report_path, products=len(articles), codes_file=codes_path, site=site_result)
    with zipfile.ZipFile(zip_path) as source, zipfile.ZipFile(output_path, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info)
            if info.filename.lower().endswith(".xlsx"):
                workbook = load_workbook(BytesIO(data))
                touched = [fill_sheet(Path(info.filename).name, sheet, sources, result) for sheet in workbook.worksheets]
                if any(touched):
                    buffer = BytesIO()
                    workbook.save(buffer)
                    data = buffer.getvalue()
                    result.files += 1
                else:
                    result.skipped_files.append(info.filename)
            else:
                result.skipped_files.append(info.filename)
            target.writestr(info, data)
    write_report(result)
    return result


NAMES_HEADERS = ["Лист", "Код ETIM", "Характеристика РАЭК", "Взято из базы", "Товаров", "Примеры", "Верно? (да/нет)"]


@dataclass
class DataSummary:
    """What the archive says about products and their data, not about cells."""

    products: int
    complete: int
    not_in_db: int
    missing: list[tuple[str, int, str]]  # data, products without it, main reason
    to_check: list[tuple[str, int, str]]  # data, products where it should be checked, main reason


def data_label(column: str) -> str:
    return re.sub(r"^EF\d+\s+", "", column)


def data_summary(result: RaekResult) -> DataSummary:
    not_in_db = {row.article for row in result.rows if row.note == "Артикула нет в базе"}
    groups: dict[str, dict[str, tuple[set[str], Counter]]] = {STATUS_LABELS[EMPTY]: {}, STATUS_LABELS[CHECK]: {}}
    incomplete = set(not_in_db)
    for row in result.rows:
        if row.status not in groups or row.article in not_in_db:
            continue
        incomplete.add(row.article)
        articles, reasons = groups[row.status].setdefault(data_label(row.column), (set(), Counter()))
        articles.add(row.article)
        reasons[row.note.split(" (")[0].split(";")[0]] += 1

    def ordered(status: str) -> list[tuple[str, int, str]]:
        items = [(label, len(articles), reasons.most_common(1)[0][0]) for label, (articles, reasons) in groups[status].items()]
        return sorted(items, key=lambda item: (-item[1], item[0]))

    return DataSummary(
        result.products,
        max(result.products - len(incomplete), 0),
        len(not_in_db),
        ordered(STATUS_LABELS[EMPTY]),
        ordered(STATUS_LABELS[CHECK]),
    )


def write_report(result: RaekResult) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Итог"
    totals = data_summary(result)
    summary.append(["Товаров в архиве", totals.products])
    summary.append(["Все данные есть и проверять нечего", totals.complete])
    if totals.not_in_db:
        summary.append(["Товаров нет в базе", totals.not_in_db])
    for title, items in (("Каких данных нет", totals.missing), ("Что проверить", totals.to_check)):
        if not items:
            continue
        summary.append([])
        summary.append([title, "Товаров", "Почему"])
        summary.cell(summary.max_row, 1).font = Font(bold=True)
        for item in items:
            summary.append(list(item))
    summary.append([])
    if result.codes_file:
        summary.append(["Файл ЧЗ", result.codes_file.name])
    for name in result.skipped_files:
        summary.append(["Файл без данных РАЭК, оставлен как есть", name])
    sheets = [summary]
    for title, statuses in (("Проверить", {STATUS_LABELS[CHECK]}), ("Нечем заполнить", {STATUS_LABELS[EMPTY]})):
        sheet = workbook.create_sheet(title)
        sheet.append(["Файл", "Лист", "RAEC_ID", "Артикул", "Название", "Колонка", "Значение", "Почему"])
        for row in result.rows:
            if row.status in statuses:
                sheet.append([row.file, row.sheet, row.raec_id, row.article, row.name, row.column, row.value, row.note])
        sheets.append(sheet)
    if result.names:
        sheet = workbook.create_sheet("Названия")
        sheet.append(NAMES_HEADERS)
        for item in result.names:
            sheet.append(item)
        sheets.append(sheet)
    for sheet in sheets:
        sheet.freeze_panes = "A2"
        for cells in sheet.columns:
            width = max(len(str(cell.value or "")) for cell in cells)
            sheet.column_dimensions[cells[0].column_letter].width = min(max(width + 2, 10), 70)
    workbook.save(result.report_path)


def remember_name_answers(report_path: Path, rules_path: Path) -> tuple[int, int]:
    """Read "да"/"нет" from the "Названия" sheet of a RAEK report: (pairs kept, pairs rejected)."""
    workbook = load_workbook(report_path, read_only=True, data_only=True)
    try:
        if "Названия" not in workbook.sheetnames:
            return 0, 0
        rows = list(workbook["Названия"].iter_rows(min_row=2, values_only=True))
    finally:
        workbook.close()
    path = learned_path(rules_path)
    assert path is not None
    learned = load_learned(rules_path)
    kept = rejected = 0
    for row in rows:
        if len(row) < 7:
            continue
        code, attribute, answer = text(row[1]), text(row[3]), canonical(text(row[6]))
        if not code or not attribute:
            continue
        if answer in LOGICAL_YES:
            names = learned["name_aliases"].setdefault(code, [])
            if attribute not in names:
                names.append(attribute)
                kept += 1
        elif answer in LOGICAL_NO:
            names = learned["name_rejects"].setdefault(code, [])
            if attribute not in names:
                names.append(attribute)
                rejected += 1
    path.parent.mkdir(parents=True, exist_ok=True)
    atomic_write_json(path, learned)
    return kept, rejected


def is_raek_zip(path: Path) -> bool:
    try:
        with zipfile.ZipFile(path) as archive:
            names = [Path(name).name for name in archive.namelist()]
    except (OSError, zipfile.BadZipFile):
        return False
    return any(re.search(r"_(required_features|COUNTRY_ID|THVED|EAN|SERIES_ID|certificate|MANUFACTURER_LINK)\.xlsx$", name, re.IGNORECASE) for name in names)


def latest_raek_zip(folders: list[Path]) -> Path | None:
    """The newest RAEK archive among the zips in `folders` (by file date)."""
    found = [path for folder in folders if folder.is_dir() for path in folder.glob("*.zip") if "_заполнено_" not in path.name]
    for path in sorted(found, key=lambda item: item.stat().st_mtime, reverse=True):
        if is_raek_zip(path):
            return path
    return None


def result_lines(result: RaekResult) -> list[str]:
    totals = data_summary(result)
    lines = [f"Товаров в архиве: {totals.products}. Все данные есть у {totals.complete}."]
    if totals.not_in_db:
        lines.append(f"Товаров нет в базе: {totals.not_in_db}.")
    if totals.missing:
        lines.append("Каких данных нет:")
        lines.extend(f"  {label} — у {count} тов. ({reason})" for label, count, reason in totals.missing[:8])
        if len(totals.missing) > 8:
            lines.append(f"  и ещё {len(totals.missing) - 8} видов данных, см. отчёт.")
    if totals.to_check:
        lines.append("Что проверить:")
        lines.extend(f"  {label} — у {count} тов. ({reason})" for label, count, reason in totals.to_check[:5])
    lines.append("Данные, которые у РАЭК уже были, программа не трогала.")
    if result.names:
        lines.append(f"Проверьте лист «Названия» в отчёте ({len(result.names)} строк): да/нет, и программа запомнит.")
    if result.needs:
        lines.append(f"Можно спросить у продактов: {len(result.needs)} поз. (кнопка «Запросить у продактов»).")
    if result.codes_file:
        lines.append(f"Коды и штрих-коды взяты из файла {result.codes_file.name}.")
    if result.site:
        lines.extend(site_lines(result.site))
    lines.append(f"Архив для РАЭК: {result.output_path.name}")
    lines.append(f"Отчёт: {result.report_path.name}")
    return lines
