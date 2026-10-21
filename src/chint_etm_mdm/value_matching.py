"""Matching product characteristic values to a partner's allowed values.

ETM publishes a reference ("Справочник") with the allowed values of each
characteristic. A value outside that list is rejected by ETM, so before a value
goes into a closed-list column it is matched to one of the allowed values:

- exact or after harmless normalization (case, ё/е, spaces, decimal comma);
- by a match the user confirmed earlier (learned matches);
- yes/no wording for logical characteristics;
- the same number ("230 В" -> 230);
- by words, when exactly one allowed value is the closest fit by words
  ("Переменный ток (AC)" -> "Переменный (AC)"); such a value needs a check.
"""
from __future__ import annotations

import difflib
import json
import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from .atomic_files import atomic_write_json

LOGICAL_YES = {"да", "yes", "true", "1", "есть", "имеется", "наличие", "y"}
LOGICAL_NO = {"нет", "no", "false", "0", "отсутствует", "без", "n", "-"}
WORD_SYNONYMS = {
    "перемен": "переменный",
    "переменного": "переменный",
    "переменное": "переменный",
    "пост": "постоянный",
    "постоян": "постоянный",
    "постоянного": "постоянный",
    "постоянное": "постоянный",
    "светодиод": "led",
    "светодиодный": "led",
    "светодиодная": "led",
    "светодиодные": "led",
}
NUMBER_RE = re.compile(r"[-+]?\d+(?:[.,]\d+)?")
UNIT_WORDS = {
    "в", "v", "а", "a", "ма", "ma", "ка", "ka", "мм", "mm", "см", "м", "m", "вт", "w", "квт", "kw",
    "гц", "hz", "кг", "kg", "г", "шт", "°c", "с", "мс", "ms", "ч",
}
PATTERN_RE = re.compile(r"^0[0/\-.,]*$")
UNIT_SUFFIX_RE = re.compile(r"(?:,\s*[^,()]{1,12}|\([^()]{1,12}\))\s*$")
WORD_RE = re.compile(r"[0-9a-zа-я]+")


def canonical(value: str) -> str:
    text = str(value).replace("\xa0", " ").replace("ё", "е").replace("Ё", "Е").casefold()
    text = re.sub(r"(?<=\d),(?=\d)", ".", text)
    text = re.sub(r"\s*([/,;:()+\-])\s*", r"\1", text)
    return " ".join(text.split()).strip(" .")


def words(value: str) -> set[str]:
    return {WORD_SYNONYMS.get(word, word) for word in WORD_RE.findall(canonical(value))}


def only_number(value: str) -> float | None:
    """Return the number if the value is a single number, optionally with a unit."""
    text = canonical(value)
    numbers = NUMBER_RE.findall(text)
    if len(numbers) != 1:
        return None
    rest = NUMBER_RE.sub("", text, count=1).replace(" ", "").strip(".")
    if rest and rest not in UNIT_WORDS and rest != "~":
        return None
    return float(numbers[0].replace(",", "."))


def format_number(number: float) -> str:
    if number.is_integer():
        return str(int(number))
    return f"{number:.6f}".rstrip("0").rstrip(".")


def is_number_pattern(values: list[str]) -> bool:
    """ETM writes number formats like "000" or "000/000" instead of a list of values."""
    return bool(values) and all(PATTERN_RE.match(value) for value in values)


def clean_open_value(name: str, raw: str) -> str:
    """For an open characteristic whose name carries a unit ("Ток, А", "Ток (А)"), keep only the number."""
    suffix = UNIT_SUFFIX_RE.search(name)
    if not suffix:
        return raw
    unit = canonical(suffix.group(0)).strip(" ,()")
    text = canonical(raw)
    numbers = NUMBER_RE.findall(text)
    if len(numbers) != 1:
        return raw
    rest = NUMBER_RE.sub("", text, count=1).replace(" ", "").strip(".")
    if rest and rest != unit and rest not in UNIT_WORDS:
        return raw
    return format_number(float(numbers[0].replace(",", ".")))


def format_allowed(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


@dataclass
class Characteristic:
    name: str
    code: str = ""
    formats: list[str] = field(default_factory=list)
    values: list[str] = field(default_factory=list)
    comment: str = ""

    @property
    def closed(self) -> bool:
        return bool(self.values) and not is_number_pattern(self.values)


@dataclass(frozen=True)
class MatchResult:
    value: str
    kind: str
    candidates: tuple[str, ...] = ()

    @property
    def confident(self) -> bool:
        return self.kind in {"exact", "normalized", "learned", "logical", "number", "open"}


def word_match(raw_words: set[str], allowed: list[str]) -> str:
    """The one allowed value that says the same in other words, or "".

    Values with digits are left to the number rule: "230-240" is not "230".
    """
    if any(char.isdigit() for word in raw_words for char in word):
        return ""
    allowed_words = {value: words(value) for value in allowed}
    inside = [value for value, item in allowed_words.items() if item and item <= raw_words]
    if inside:
        best = max(len(allowed_words[value]) for value in inside)
        top = [value for value in inside if len(allowed_words[value]) == best]
        return top[0] if len(top) == 1 else ""
    around = [value for value, item in allowed_words.items() if raw_words <= item]
    if around:
        best = min(len(allowed_words[value]) for value in around)
        top = [value for value in around if len(allowed_words[value]) == best]
        return top[0] if len(top) == 1 else ""
    return ""


def match_value(
    raw: str,
    characteristic: Characteristic,
    learned: dict[str, str] | None = None,
) -> MatchResult:
    allowed = characteristic.values
    if not allowed or is_number_pattern(allowed):
        return MatchResult(clean_open_value(characteristic.name, raw), "open")
    if raw in allowed:
        return MatchResult(raw, "exact")
    by_canonical = {canonical(value): value for value in allowed}
    key = canonical(raw)
    if key in by_canonical:
        return MatchResult(by_canonical[key], "normalized")
    learned_value = (learned or {}).get(key)
    if learned_value and learned_value in allowed:
        return MatchResult(learned_value, "learned")
    if set(by_canonical) <= {"да", "нет"}:
        if key in LOGICAL_YES and "да" in by_canonical:
            return MatchResult(by_canonical["да"], "logical")
        if key in LOGICAL_NO and "нет" in by_canonical:
            return MatchResult(by_canonical["нет"], "logical")
    number = only_number(raw)
    numeric = [(value, only_number(value)) for value in allowed]
    numeric = [(value, item) for value, item in numeric if item is not None]
    if number is not None:
        same = [value for value, item in numeric if item == number]
        if len(same) == 1:
            return MatchResult(same[0], "number")
    raw_words = words(raw)
    if raw_words:
        by_words = word_match(raw_words, allowed)
        if by_words:
            return MatchResult(by_words, "words")
    if number is not None and numeric:
        nearest = sorted(numeric, key=lambda pair: abs(pair[1] - number))[:3]
        return MatchResult("", "none", tuple(value for value, _ in nearest))
    close = difflib.get_close_matches(key, list(by_canonical), n=3, cutoff=0.5)
    if not close and raw_words:
        close = [canonical(value) for value in allowed if raw_words & words(value)][:3]
    return MatchResult("", "none", tuple(by_canonical[item] for item in close))


class ValueLists:
    """Allowed values of one partner plus the matches the user confirmed."""

    def __init__(self, reference_path: Path, learned_path: Path) -> None:
        self.reference_path = reference_path
        self.learned_path = learned_path
        self.characteristics = load_reference(reference_path)
        self.learned = load_learned(learned_path)
        self.name_aliases = load_name_aliases(learned_path)

    def characteristic(self, name: str) -> Characteristic | None:
        return self.characteristics.get(canonical(name))

    def aliases_for(self, name: str) -> list[str]:
        return self.name_aliases.get(canonical(name), [])

    def match(self, name: str, raw: str) -> MatchResult:
        characteristic = self.characteristic(name) or Characteristic(name=name)
        return match_value(raw, characteristic, self.learned.get(canonical(name)))


def load_reference(path: Path) -> dict[str, Characteristic]:
    if not path.exists():
        return {}
    data = json.loads(path.read_text(encoding="utf-8"))
    result: dict[str, Characteristic] = {}
    for item in data.get("characteristics", []):
        characteristic = Characteristic(
            name=item["name"],
            code=str(item.get("code", "")),
            formats=list(item.get("formats", [])),
            values=[str(value) for value in item.get("values", [])],
            comment=item.get("comment", ""),
        )
        result[canonical(characteristic.name)] = characteristic
    return result


def load_learned(path: Path) -> dict[str, dict[str, str]]:
    if not path.exists():
        return {}
    data = json.loads(path.read_text(encoding="utf-8"))
    learned: dict[str, dict[str, str]] = {}
    for item in data.get("matches", []):
        name = canonical(item.get("characteristic", ""))
        raw = canonical(item.get("raw", ""))
        value = str(item.get("value", "")).strip()
        if name and raw and value:
            learned.setdefault(name, {})[raw] = value
    return learned


def load_name_aliases(path: Path) -> dict[str, list[str]]:
    """Partner characteristic name -> our attribute names the user confirmed for it."""
    if not path.exists():
        return {}
    data = json.loads(path.read_text(encoding="utf-8"))
    aliases: dict[str, list[str]] = {}
    for item in data.get("name_aliases", []):
        name = canonical(item.get("characteristic", ""))
        source = str(item.get("source_attribute", "")).strip()
        if name and source and source not in aliases.setdefault(name, []):
            aliases[name].append(source)
    return aliases


def remember_name_aliases(path: Path, aliases: list[tuple[str, str]]) -> int:
    """Save (partner characteristic, our attribute name) pairs; return how many are new."""
    data: dict = {"version": 1, "matches": []}
    if path.exists():
        data = json.loads(path.read_text(encoding="utf-8"))
    items = data.setdefault("name_aliases", [])
    known = {(canonical(item.get("characteristic", "")), item.get("source_attribute", "")) for item in items}
    added = 0
    for characteristic, source in aliases:
        key = (canonical(characteristic), source)
        if not key[0] or not source or key in known:
            continue
        items.append({"characteristic": characteristic, "source_attribute": source})
        known.add(key)
        added += 1
    path.parent.mkdir(parents=True, exist_ok=True)
    atomic_write_json(path, data)
    return added


def stems(name: str) -> set[str]:
    return {word[:5] for word in words(UNIT_SUFFIX_RE.sub("", name)) if len(word) >= 4 and not word.isdigit()}


def name_unit(name: str) -> str:
    suffix = UNIT_SUFFIX_RE.search(name)
    return canonical(suffix.group(0)).strip(" ,()") if suffix else ""


def name_candidates(
    name: str, products_attributes: list[dict[str, str]], limit: int = 3
) -> list[tuple[str, int, tuple[str, ...]]]:
    """Our attribute names that may mean the partner's characteristic `name`.

    Ranked by shared word stems, then same unit, then how many of the products have it.
    Returns (attribute name, products with it, up to 3 example values).
    """
    target = stems(name)
    unit = name_unit(name)
    counts: dict[str, int] = {}
    examples: dict[str, list[str]] = {}
    for attributes in products_attributes:
        for attribute, value in attributes.items():
            if not value:
                continue
            counts[attribute] = counts.get(attribute, 0) + 1
            bucket = examples.setdefault(attribute, [])
            if value not in bucket and len(bucket) < 3:
                bucket.append(value)
    scored = []
    for attribute, count in counts.items():
        overlap = len(target & stems(attribute))
        if not overlap:
            continue
        same_unit = 1 if unit and unit == name_unit(attribute) else 0
        scored.append(((overlap, same_unit, count), attribute))
    scored.sort(key=lambda item: (item[0], item[1]), reverse=True)
    return [(attribute, counts[attribute], tuple(examples[attribute])) for _, attribute in scored[:limit]]


def remember_matches(path: Path, matches: list[tuple[str, str, str]]) -> int:
    """Save (characteristic, raw value, allowed value) triples; return how many are new."""
    data = {"version": 1, "matches": []}
    if path.exists():
        data = json.loads(path.read_text(encoding="utf-8"))
        data.setdefault("matches", [])
    by_key = {
        (canonical(item.get("characteristic", "")), canonical(item.get("raw", ""))): item
        for item in data["matches"]
    }
    added = 0
    for characteristic, raw, value in matches:
        key = (canonical(characteristic), canonical(raw))
        if not key[0] or not key[1]:
            continue
        item = by_key.get(key)
        if item is None:
            item = {"characteristic": characteristic, "raw": raw}
            data["matches"].append(item)
            by_key[key] = item
            added += 1
        item["value"] = value
    path.parent.mkdir(parents=True, exist_ok=True)
    atomic_write_json(path, data)
    return added


def read_etm_reference_workbook(path: Path) -> list[Characteristic]:
    """Read the ETM OneTeam "Справочник" export: one row per allowed value."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Справочник"] if "Справочник" in workbook.sheetnames else workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = [str(cell or "").strip().casefold() for cell in next(rows, ())]

        def column(*names: str) -> int:
            for idx, title in enumerate(header):
                if any(title.startswith(name) for name in names):
                    return idx
            raise ValueError(f"В справочнике нет колонки «{names[0]}».")

        code_idx = column("код характеристики")
        name_idx = column("название характеристики", "наименование характеристики")
        format_idx = column("формат")
        value_idx = column("возможные значения")
        comment_idx = column("комментарий")
        result: dict[str, Characteristic] = {}
        for row in rows:
            if not row or row[name_idx] in (None, ""):
                continue
            name = str(row[name_idx]).strip()
            characteristic = result.setdefault(canonical(name), Characteristic(name=name))
            code = format_allowed(row[code_idx]) if row[code_idx] is not None else ""
            if code and not characteristic.code:
                characteristic.code = code
            kind = str(row[format_idx] or "").strip()
            if kind and kind not in characteristic.formats:
                characteristic.formats.append(kind)
            value = row[value_idx]
            if value not in (None, ""):
                text = format_allowed(value)
                if text not in characteristic.values:
                    characteristic.values.append(text)
            comment = str(row[comment_idx] or "").strip() if comment_idx < len(row) else ""
            if comment and not characteristic.comment:
                characteristic.comment = comment
        return list(result.values())
    finally:
        workbook.close()


def import_etm_reference(source: Path, target: Path) -> tuple[int, int]:
    """Convert the ETM reference XLSX into the program's JSON; return (characteristics, values)."""
    characteristics = read_etm_reference_workbook(source)
    if not characteristics:
        raise ValueError("В справочнике не найдено ни одной характеристики.")
    target.parent.mkdir(parents=True, exist_ok=True)
    atomic_write_json(
        target,
        {
            "version": 1,
            "source_file": source.name,
            "characteristics": [
                {
                    "name": item.name,
                    "code": item.code,
                    "formats": item.formats,
                    "values": item.values,
                    "comment": item.comment,
                }
                for item in characteristics
            ],
        },
    )
    return len(characteristics), sum(len(item.values) for item in characteristics)
