"""Certificate dates kept once per certificate, not per product.

The whole CHINT range shares a few dozen ЕАЭС certificates and declarations,
while partners (RAEK) want the dates "from" and "to" on every product. The
program keeps a table in the work folder, «Сертификаты.xlsx»: one row per
document, with the dates it already knows (the ЧЗ file, when it has date
columns; the database) and empty cells for the user to fill once. On the
next update the user's dates are kept and new documents are added. Partner
files take the dates from this table for every product with that document.

Dates filled earlier in another file (for example a partner's certificate form
with a row per product) are taken over with «Взять даты из файла...».
"""
from __future__ import annotations

import datetime as dt
import re
from collections import Counter, defaultdict
from contextlib import closing
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .db_utils import connect_database
from .filler import normalize_article

TABLE_NAME = "Сертификаты.xlsx"
SHEET = "Сертификаты"
HEADERS = ["Номер документа", "Вид", "Действует с", "Действует до", "Откуда даты", "Товаров", "Серии", "Пример товара", "Ссылка на реестр"]
NUMBER, KIND, START, END, SOURCE, COUNT, SERIES, EXAMPLE, LINK = range(len(HEADERS))
FROM_USER, FROM_CHZ, FROM_DB, FROM_FILE = "вписано вручную", "файл ЧЗ", "база", "взято из файла"
FROM_NUMBER = "дата из номера письма"
EXPIRING_DAYS = 90
LATIN = str.maketrans("ABCEHKMOPTXY", "АВСЕНКМОРТХУ")
FILL_EMPTY = PatternFill("solid", fgColor="FFF2CC")
FILL_EXPIRED = PatternFill("solid", fgColor="F4CCCC")


def number_key(number: str) -> str:
    """The same document written with Latin look-alike letters or other spacing gives the same key."""
    return re.sub(r"[^0-9А-ЯЁ]", "", str(number or "").upper().translate(LATIN))


def is_eaeu_document(number: str) -> bool:
    """A certificate or declaration of the Eurasian union (ЕАЭС/ТС), not a refusal letter or "не требуется"."""
    compact = str(number or "").upper().translate(LATIN)
    return "ЕАЭС" in compact or bool(re.match(r"\s*ТС\b", compact))


def is_refusal_letter(number: str) -> bool:
    """A refusal letter (the product needs no certificate): "24/11/0005", "29.05.2026-15ССП"..."""
    text = str(number or "").strip()
    return bool(re.search(r"\d", text)) and not is_eaeu_document(text) and not re.match(r"(?i)\s*(росс|poсс|pocc|гост)", text)


def letter_date(number: str) -> dt.date | None:
    """The date some letters carry in the number: "29.05.2026-15ССП" -> 29.05.2026."""
    found = re.search(r"(?<!\d)(\d{2})\.(\d{2})\.(\d{4})(?!\d)", str(number or ""))
    if not found:
        return None
    try:
        return dt.date(int(found.group(3)), int(found.group(2)), int(found.group(1)))
    except ValueError:
        return None


def document_kind(number: str) -> str:
    if is_refusal_letter(number):
        return "Отказное письмо"
    compact = str(number or "").replace(" ", "").upper().translate(LATIN)
    if "Д-" in compact:
        return "Декларация ЕАС"
    if "С-" in compact:
        return "Сертификат ЕАС"
    return ""


def as_date(value: object) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y", "%d/%m/%Y"):
        try:
            return dt.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


@dataclass
class Document:
    number: str
    start: dt.date | None = None
    end: dt.date | None = None
    source: str = ""
    link: str = ""
    articles: set[str] = field(default_factory=set)
    series: Counter = field(default_factory=Counter)
    example: str = ""


@dataclass
class TableResult:
    path: Path
    documents: int = 0
    with_dates: int = 0
    without_dates: int = 0
    added: int = 0
    from_file: int = 0
    expiring: list[tuple[str, dt.date]] = field(default_factory=list)
    expired: list[tuple[str, dt.date]] = field(default_factory=list)


def chz_date_columns(names: list[str]) -> dict[str, int]:
    """Date columns of the first certificate in the ЧЗ file, once someone adds them.

    Accepted headers: «Сертификат 1 от»/«Сертификат 1 до», «Дата начала 1»/«Дата окончания 1»,
    «Действует с»/«Действует до» and the like, without a number or with 1.
    """
    columns: dict[str, int] = {}
    for idx, name in enumerate(names):
        if not name or "ссылка" in name or re.search(r"[2-9]", name):
            continue
        words = set(re.findall(r"[а-яё]+", name))
        if "start" not in columns and (words & {"от", "с"} or re.search(r"начал|выдач", name)) and re.search(r"серт|дат|действ|срок", name):
            columns["start"] = idx
        elif "end" not in columns and (words & {"до", "по"} or re.search(r"оконч|истеч", name)) and re.search(r"серт|дат|действ|срок", name):
            columns["end"] = idx
    return columns


def read_chz(path: Path | None) -> list[dict[str, object]]:
    """Rows of the ЧЗ file: article, series, name, first certificate, its registry link and dates if the file has them."""
    if path is None:
        return []
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        sheet = workbook["Price-list"] if "Price-list" in workbook.sheetnames else workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        columns: dict[str, int] = {}
        for row in rows:
            names = [" ".join(str(cell or "").split()).casefold() for cell in row]
            if "part number" in names:
                wanted = {"part number": "article", "russian description": "name", "серия": "series", "сертификат 1": "number", "ссылка на реестр 1": "link"}
                columns = {wanted[name]: idx for idx, name in enumerate(names) if name in wanted}
                columns.update(chz_date_columns(names))
                break
        if "article" not in columns or "number" not in columns:
            return []
        result = []
        for row in rows:
            item = {key: (row[idx] if idx < len(row) else None) for key, idx in columns.items()}
            item["article"] = normalize_article(item["article"])
            if item["article"]:
                result.append(item)
        return result
    finally:
        workbook.close()


def database_documents(db_path: Path) -> dict[str, tuple[str, dt.date | None, dt.date | None, set[str]]]:
    """Certificate number and dates the database holds per product, grouped by document."""
    with closing(connect_database(db_path, readonly=True)) as conn:
        rows = conn.execute(
            """
            SELECT article, attribute_name, COALESCE(value_normalized, value_raw) AS value
            FROM product_attribute_values
            WHERE attribute_name IN ('Номер сертификата', 'Дата начала', 'Дата окончания')
            ORDER BY confidence DESC, id DESC
            """
        ).fetchall()
    per_article: dict[str, dict[str, str]] = defaultdict(dict)
    for row in rows:
        per_article[str(row["article"])].setdefault(str(row["attribute_name"]), str(row["value"] or ""))
    grouped: dict[str, tuple[str, Counter, set[str]]] = {}
    for article, values in per_article.items():
        number = values.get("Номер сертификата", "")
        if not is_eaeu_document(number):
            continue
        key = number_key(number)
        _, dates, articles = grouped.setdefault(key, (number, Counter(), set()))
        dates[(as_date(values.get("Дата начала")), as_date(values.get("Дата окончания")))] += 1
        articles.add(article)
    result = {}
    for key, (number, dates, articles) in grouped.items():
        known = [(pair, count) for pair, count in dates.most_common() if pair[0] or pair[1]]
        start, end = known[0][0] if known else (None, None)
        result[key] = (number, start, end, articles)
    return result


def read_dates_file(path: Path) -> dict[str, Document]:
    """Documents and dates from any table with a certificate number column and date columns.

    Fits partner forms filled earlier: a row per product with «Номер сертификата», «Начало действия
    сертификата», «Окончание действия сертификата». Rows of one document give one entry.
    """
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    documents: dict[str, Document] = {}
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            columns: dict[str, int] = {}
            for _, row in zip(range(15), rows):
                names = [" ".join(str(cell or "").split()).casefold() for cell in row]
                number = next((i for i, name in enumerate(names) if re.search(r"номер\s+(сертификата|документа|декларации)", name)), None)
                if number is None:
                    continue
                columns = {"number": number}
                for idx, name in enumerate(names):
                    if "start" not in columns and re.search(r"начал|выдач|действует с|\bот\b", name) and idx != number:
                        columns["start"] = idx
                    elif "end" not in columns and re.search(r"оконч|истеч|действует до|\bдо\b", name) and idx != number:
                        columns["end"] = idx
                    elif "article" not in columns and name.startswith("артикул"):
                        columns["article"] = idx
                    elif "link" not in columns and "реестр" in name:
                        columns["link"] = idx
                break
            if "start" not in columns and "end" not in columns:
                continue
            for row in rows:
                cells = {key: (row[idx] if idx < len(row) else None) for key, idx in columns.items()}
                number = str(cells["number"] or "").strip()
                if not is_eaeu_document(number):
                    continue
                doc = documents.setdefault(number_key(number), Document(number, source=FROM_FILE))
                start, end = as_date(cells.get("start")), as_date(cells.get("end"))
                if (start or end) and not (doc.start or doc.end):
                    doc.start, doc.end = start, end
                if cells.get("article"):
                    doc.articles.add(normalize_article(cells["article"]))
                doc.link = doc.link or str(cells.get("link") or "").strip()
    finally:
        workbook.close()
    return {key: doc for key, doc in documents.items() if doc.start or doc.end}


def read_table(path: Path) -> dict[str, Document]:
    """Documents of an existing table, with the dates written in it."""
    if not path.exists():
        return {}
    workbook = load_workbook(path, data_only=True)
    try:
        sheet = workbook[SHEET] if SHEET in workbook.sheetnames else workbook.worksheets[0]
        documents = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            row = list(row) + [None] * (len(HEADERS) - len(row))
            number = str(row[NUMBER] or "").strip()
            if not number:
                continue
            documents[number_key(number)] = Document(
                number=number, start=as_date(row[START]), end=as_date(row[END]), source=str(row[SOURCE] or "").strip()
            )
        return documents
    finally:
        workbook.close()


def table_dates(path: Path | None) -> dict[str, tuple[dt.date | None, dt.date | None]]:
    """Dates by document key, for partner files."""
    if path is None:
        return {}
    return {key: (doc.start, doc.end) for key, doc in read_table(path).items() if doc.start or doc.end}


def update_table(
    db_path: Path, codes_path: Path | None, path: Path, today: dt.date | None = None, dates_file: Path | None = None
) -> TableResult:
    """Create or refresh the table. Dates already in the table always win: the user may have typed them.

    With `dates_file`, dates from that file fill the documents the user has not dated by hand.
    """
    today = today or dt.date.today()
    old = read_table(path)
    documents: dict[str, Document] = {}
    for item in read_chz(codes_path):
        number = str(item.get("number") or "").strip()
        if not (is_eaeu_document(number) or is_refusal_letter(number)):
            continue
        doc = documents.setdefault(number_key(number), Document(number))
        doc.articles.add(item["article"])
        if item.get("series"):
            doc.series[str(item["series"]).strip()] += 1
        doc.example = doc.example or str(item.get("name") or "")
        doc.link = doc.link or str(item.get("link") or "")
        start, end = as_date(item.get("start")), as_date(item.get("end"))
        if (start or end) and not (doc.start or doc.end):
            doc.start, doc.end, doc.source = start, end, FROM_CHZ
        elif not (doc.start or doc.end) and letter_date(number):
            doc.start, doc.source = letter_date(number), FROM_NUMBER
    for key, (number, start, end, articles) in database_documents(db_path).items():
        doc = documents.setdefault(key, Document(number))
        doc.articles |= articles
        if (start or end) and not (doc.start or doc.end):
            doc.start, doc.end, doc.source = start, end, FROM_DB
    result = TableResult(path)
    for key, previous in old.items():
        doc = documents.setdefault(key, Document(previous.number))
        if not (previous.start or previous.end):
            continue
        if previous.source == FROM_FILE:
            doc.start, doc.end, doc.source = previous.start, previous.end, FROM_FILE
        elif (previous.start, previous.end) != (doc.start, doc.end) or previous.source == FROM_USER:
            doc.start, doc.end, doc.source = previous.start, previous.end, FROM_USER
    if dates_file is not None:
        for key, found in read_dates_file(dates_file).items():
            doc = documents.setdefault(key, Document(found.number))
            doc.articles |= found.articles
            doc.link = doc.link or found.link
            if doc.source == FROM_USER:
                continue
            if (doc.start, doc.end) != (found.start, found.end):
                result.from_file += 1
            doc.start, doc.end, doc.source = found.start, found.end, FROM_FILE
    result.added = len(set(documents) - set(old))
    write_table(path, documents, today, result)
    return result


def write_table(path: Path, documents: dict[str, Document], today: dt.date, result: TableResult) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    ordered = sorted(documents.values(), key=lambda doc: (doc.start is not None or doc.end is not None, -len(doc.articles), doc.number))
    for doc in ordered:
        sheet.append(
            [
                doc.number,
                document_kind(doc.number),
                doc.start,
                doc.end,
                doc.source,
                len(doc.articles),
                ", ".join(name for name, _ in doc.series.most_common(5)),
                doc.example,
                doc.link,
            ]
        )
        row = sheet.max_row
        for col in (START, END):
            sheet.cell(row, col + 1).number_format = "DD.MM.YYYY"
        if not (doc.start or doc.end):
            result.without_dates += 1
            for col in (START, END):
                sheet.cell(row, col + 1).fill = FILL_EMPTY
        else:
            result.with_dates += 1
        if doc.end and doc.end < today:
            result.expired.append((doc.number, doc.end))
            sheet.cell(row, END + 1).fill = FILL_EXPIRED
        elif doc.end and (doc.end - today).days <= EXPIRING_DAYS:
            result.expiring.append((doc.number, doc.end))
            sheet.cell(row, END + 1).fill = FILL_EXPIRED
    result.documents = len(ordered)
    dates = DataValidation(type="date", operator="greaterThan", formula1="DATE(2000,1,1)", allow_blank=True)
    dates.error = "Впишите дату, например 21.03.2025"
    sheet.add_data_validation(dates)
    dates.add(f"C2:D{max(sheet.max_row, 2)}")
    for col, width in zip("ABCDEFGHI", (34, 16, 13, 13, 16, 9, 24, 50, 50)):
        sheet.column_dimensions[col].width = width
    sheet.freeze_panes = "A2"
    note = workbook.create_sheet("Как заполнять")
    for line in (
        "Впишите даты «Действует с» и «Действует до» в жёлтые клетки, сохраните файл и закройте его.",
        "Даты достаточно вписать один раз на документ: программа сама поставит их всем товарам с этим сертификатом.",
        "Вписанные даты программа не перезаписывает. Новые сертификаты из файла ЧЗ она добавляет сама при обновлении таблицы.",
        "Красным отмечены сертификаты, которые истекли или истекают в ближайшие 90 дней.",
    ):
        note.append([line])
    note.column_dimensions["A"].width = 110
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f".{path.name}.tmp")
    workbook.save(tmp)
    tmp.replace(path)


def table_lines(result: TableResult) -> list[str]:
    lines = [
        f"Таблица сертификатов: {result.documents} документов, с датами {result.with_dates}, без дат {result.without_dates}"
        + (f", новых {result.added}." if result.added else ".")
    ]
    if result.from_file:
        lines.append(f"Из выбранного файла взяты даты для {result.from_file} документов. Вписанные вручную даты не тронуты.")
    if result.without_dates:
        lines.append("Впишите даты в жёлтые клетки, сохраните файл: дальше программа поставит их сама.")
    for number, end in sorted(result.expired, key=lambda item: item[1]):
        lines.append(f"Истёк {end:%d.%m.%Y}: {number}")
    for number, end in sorted(result.expiring, key=lambda item: item[1]):
        lines.append(f"Истекает {end:%d.%m.%Y}: {number}")
    lines.append(f"Файл: {result.path}")
    return lines
