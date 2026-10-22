"""ETIM import: every sheet is read, and the base and the report never disagree."""

from __future__ import annotations

import sqlite3
import tempfile
import unittest
from contextlib import closing
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

from chint_etm_mdm import etim_importer


def make_base(path: Path, articles: list[str]) -> None:
    with closing(sqlite3.connect(path)) as conn, conn:
        conn.execute("CREATE TABLE products (article TEXT PRIMARY KEY)")
        conn.executemany("INSERT INTO products VALUES (?)", [(article,) for article in articles])


def write_etim(path: Path, sheets: list[tuple[str, list[list[object]]]]) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets:
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
    workbook.save(path)
    workbook.close()
    return path


SHORT = ("Кратко", [["Артикул", "Наименование"], ["A-1", "Выключатель"]])
DETAILED = ("Подробно", [["Артикул", "Длина, мм", "Ширина, мм", "Высота, мм", "Цвет"], ["A-1", 100, 50, 80, "белый"]])


def rows(path: Path, query: str) -> list[tuple]:
    with closing(sqlite3.connect(path)) as conn:
        return conn.execute(query).fetchall()


def has_table(path: Path, name: str) -> bool:
    return bool(rows(path, f"SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = '{name}'"))


class ScanAllSheetsTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def scan(self, sheets, articles=("A-1",), notes=None):
        path = write_etim(self.root / "etim.xlsx", sheets)
        return etim_importer.scan_etim_workbook(path, list(articles), notes)

    def test_short_sheet_first_does_not_hide_the_detailed_one(self) -> None:
        found = self.scan([SHORT, DETAILED])["A-1"]
        self.assertEqual({"length_mm": 100.0, "width_mm": 50.0, "height_mm": 80.0}, found.dimensions)
        self.assertEqual("Подробно", found.dimension_sheets["length_mm"])
        self.assertEqual("Выключатель", found.attributes["Наименование"])
        self.assertEqual("белый", found.attributes["Цвет"])
        self.assertEqual("Кратко", found.attribute_sheets["Наименование"])
        self.assertEqual("Подробно", found.attribute_sheets["Цвет"])

    def test_sheet_order_does_not_change_the_result(self) -> None:
        first = self.scan([SHORT, DETAILED])["A-1"]
        second = self.scan([DETAILED, SHORT])["A-1"]
        self.assertEqual(first.dimensions, second.dimensions)
        self.assertEqual(first.attributes, second.attributes)

    def test_analog_is_not_taken_for_the_product(self) -> None:
        notes = etim_importer.ScanNotes()
        found = self.scan(
            [("Лист", [["Артикул", "Артикул аналога", "Длина, мм"], ["B-9", "A-1", 300]])], notes=notes
        )
        self.assertEqual({}, found)
        self.assertEqual({"B-9": "Лист"}, notes.unknown_articles)

    def test_different_sizes_in_the_file_are_not_chosen_silently(self) -> None:
        notes = etim_importer.ScanNotes()
        other = ("Ещё", [["Артикул", "Длина, мм", "Ширина, мм"], ["A-1", 120, 50]])
        found = self.scan([DETAILED, other], notes=notes)["A-1"]
        self.assertNotIn("length_mm", found.dimensions)
        self.assertEqual(50.0, found.dimensions["width_mm"])
        lengths = [conflict for conflict in notes.conflicts if conflict.field == "Длина, мм"]
        self.assertEqual(1, len(lengths))
        self.assertEqual({"100", "120"}, {value for value, _sheet, _row in lengths[0].values})
        self.assertEqual("", lengths[0].kept)

    def test_same_value_twice_is_not_a_conflict(self) -> None:
        notes = etim_importer.ScanNotes()
        found = self.scan([DETAILED, ("Копия", DETAILED[1])], notes=notes)["A-1"]
        self.assertEqual(100.0, found.dimensions["length_mm"])
        self.assertEqual([], notes.conflicts)


class ImportSafetyTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.db = self.root / "base" / "chint_mdm.sqlite"
        self.db.parent.mkdir()
        make_base(self.db, ["A-1", "A-2"])
        self.reports = self.root / "reports"

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def run_import(self, sheets, make_backup: bool = False):
        path = write_etim(self.root / "etim.xlsx", sheets)
        return etim_importer.import_etim_workbook(path, self.db, make_backup=make_backup, report_dir=self.reports)

    def test_import_reads_both_sheets_and_reports_extras(self) -> None:
        other = ("Ещё", [["Артикул", "Высота, мм"], ["A-1", 90], ["Z-404", 10]])
        result = self.run_import([SHORT, DETAILED, other])
        self.assertEqual(1, result.found_articles)
        self.assertEqual(1, result.unknown_articles)
        self.assertEqual(1, result.file_conflicts)
        self.assertEqual(
            [(100.0, 50.0, None)],
            rows(self.db, "SELECT length_mm, width_mm, height_mm FROM product_dimensions_resolved"),
        )
        self.assertEqual(
            [("length_mm", "Подробно"), ("width_mm", "Подробно")],
            rows(self.db, "SELECT dimension_kind, source_sheet FROM product_dimensions_raw ORDER BY dimension_kind"),
        )
        report = load_workbook(result.report_path)
        self.assertIn(["Z-404", "Ещё"], [list(row) for row in report["Нет в базе"].iter_rows(min_row=2, values_only=True)])
        heights = [row for row in report["Разное в файле"].iter_rows(min_row=2, values_only=True) if row[1] == "Высота, мм"]
        self.assertEqual(2, len(heights))
        self.assertEqual({"не записано"}, {row[5] for row in heights})
        self.assertEqual(result.import_id, etim_importer.report_import_id(report))
        report.close()
        self.assertEqual([], list(self.reports.glob("*.saving.xlsx")))
        lines = etim_importer.result_lines(result)
        self.assertTrue(any("Разные значения в самом файле: 1" in line for line in lines))
        self.assertTrue(any("нет в базе: 1" in line for line in lines))

    def test_nothing_found_is_said_plainly(self) -> None:
        result = self.run_import([("Лист", [["Артикул", "Длина, мм"], ["Z-1", 10]])])
        self.assertEqual(0, result.found_articles)
        self.assertTrue(etim_importer.result_lines(result)[0].startswith("Внимание: в файле не найдено"))

    def test_report_that_cannot_be_saved_leaves_the_base_unchanged(self) -> None:
        with mock.patch.object(etim_importer, "write_etim_report", side_effect=PermissionError("файл открыт")):
            with self.assertRaisesRegex(RuntimeError, "База не изменена"):
                self.run_import([DETAILED])
        self.assertEqual([], rows(self.db, "SELECT * FROM dimension_imports"))
        self.assertEqual([], rows(self.db, "SELECT * FROM product_dimensions_resolved"))
        self.assertEqual([], rows(self.db, "SELECT * FROM product_attribute_values"))
        self.assertEqual([], list(self.reports.glob("*.xlsx")))

    def test_report_name_failure_after_commit_still_gives_the_report(self) -> None:
        with mock.patch.object(etim_importer.os, "replace", side_effect=PermissionError("занято")):
            result = self.run_import([DETAILED])
        self.assertIn("временным именем", result.report_warning)
        self.assertTrue(result.report_path.name.endswith(".saving.xlsx"))
        self.assertTrue(result.report_path.exists())
        self.assertEqual(1, len(rows(self.db, "SELECT * FROM dimension_imports")))
        report = load_workbook(result.report_path)
        self.assertEqual(result.import_id, etim_importer.report_import_id(report))
        report.close()
        self.assertTrue(any("Внимание: Отчёт сохранён" in line for line in etim_importer.result_lines(result)))

    def test_broken_file_stops_before_backup_and_tables(self) -> None:
        broken = self.root / "etim.xlsx"
        broken.write_bytes(b"not a workbook")
        with self.assertRaises(Exception):
            etim_importer.import_etim_workbook(broken, self.db, make_backup=True, report_dir=self.reports)
        self.assertEqual(["chint_mdm.sqlite"], sorted(path.name for path in self.db.parent.iterdir()))
        self.assertFalse(has_table(self.db, "dimension_imports"))


if __name__ == "__main__":
    unittest.main()
