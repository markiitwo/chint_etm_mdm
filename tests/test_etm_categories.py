import csv
import datetime as dt
import sqlite3
import tempfile
import unittest
from contextlib import closing
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from chint_etm_mdm.etm_categories import (
    category_library,
    fill_new_products,
    new_articles_since,
    read_category_template,
)
from chint_etm_mdm.value_matching import canonical
from test_import_data_integrity import make_database

YELLOW = PatternFill(fill_type="solid", fgColor="FFFFFF00")
BASE = [
    "Режим:ОписаниеТоваров", "Код производителя", "Расширенный артикул", "Артикул", "81 класс", "Название",
    "Полное название", "Страна", "Код ТН ВЭД", "Код ОКПД2", "Название упаковки", "Вес, кг",
    "Длина, м", "Ширина, м", "Высота, м",
]
CONFIG = [
    "Конфиг:Тип изделия", "Конфиг:Количество силовых полюсов", "Конфиг:Номинальный ток (А)",
    "Конфиг:Дифференциальный ток (мА)", "Конфиг:Тип срабатывания по диф.току",
]


def make_category_template(path: Path, class_code: int = 500820) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Шаблон"
    for col, header in enumerate(BASE + CONFIG, start=1):
        cell = sheet.cell(1, col, header)
        if header != "Режим:ОписаниеТоваров":
            cell.fill = YELLOW
    for row in range(2, 7):
        sheet.cell(row, 5, class_code)
    lists = workbook.create_sheet("Списки")
    columns = {
        "A": ["Выключатель автоматический дифференциального тока", "Выключатель дифференциального тока"],
        "B": [1, 2, "1P+N", "2P", "4P"],
        "C": [10, 30, 100],
        "D": ["A", "AC", "B"],
    }
    for letter, values in columns.items():
        for row, value in enumerate(values, start=1):
            lists[f"{letter}{row}"] = value
    for target, (letter, values) in zip(("P", "Q", "S", "T"), columns.items()):
        validation = DataValidation(type="list", formula1=f"'Списки'!${letter}$1:${letter}${len(values)}")
        validation.add(f"{target}2:{target}500")
        sheet.add_data_validation(validation)
    workbook.save(path)


class CategoryTemplateTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.templates = self.root / "шаблоны"
        self.templates.mkdir()
        make_category_template(self.templates / "upload_goods_category_template_500820.xlsx")
        self.db = self.root / "mdm.sqlite"
        make_database(self.db)
        with closing(sqlite3.connect(self.db)) as conn, conn:
            conn.executemany(
                "INSERT INTO products(article, name, series_name, equipment_group) VALUES (?, ?, ?, ?)",
                [
                    ("100", "АВДТ NB1L 1P+N C16 30мА тип A", "NB1L", "АВДТ"),
                    ("586772", "АВДТ NB1L-20 1P+N B6 30мА тип AC 6кА (R)", "NB1L-20", "АВДТ"),
                    ("625288", "Промежуточное реле NJDC-17/3ZS 10А 24В AC (R)", "NJDC-17", "Реле"),
                    ("777", "Изделие без серии", "", ""),
                ],
            )
            conn.execute("INSERT INTO ipro_goods(article, name, full_name, class81_code) VALUES ('100', 'x', 'x', '500820')")
            conn.execute("INSERT INTO ipro_goods(article, name, full_name, class81_code) VALUES ('900', 'y', 'y', '50341020')")
            conn.execute("INSERT INTO products(article, name, series_name) VALUES ('900', 'Реле NJDC-17/2Z', 'NJDC-17')")
            conn.executemany(
                "INSERT INTO new_products(article, name, introduced_at) VALUES (?, ?, ?)",
                [("586772", "", "29.09.2026"), ("625288", "", "04.09.2026"), ("777", "", "10.10.2026")],
            )
            conn.execute(
                "INSERT INTO product_dimensions_resolved(article, length_mm, width_mm, height_mm, weight_kg) VALUES ('586772', 77, 18, 92.5, 0.158)"
            )
            conn.executemany(
                "INSERT INTO product_attribute_values(article, attribute_name, value_raw, confidence) VALUES (?, ?, ?, 1)",
                [
                    ("586772", "Количество полюсов", "1P+N"),
                    ("586772", "Номинальный ток, А", "6"),
                    ("586772", "Номинальный ток утечки, А", "0.03"),
                    ("586772", "Тип по току утечки", "AC"),
                ],
            )

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def test_template_reads_class_and_dropdowns(self) -> None:
        category = read_category_template(self.templates / "upload_goods_category_template_500820.xlsx")
        self.assertEqual("500820", category.class_code)
        self.assertEqual(("10", "30", "100"), category.dropdowns[canonical("Дифференциальный ток (мА)")])
        self.assertEqual({"500820"}, set(category_library(self.templates)))

    def test_split_series_prefers_a_class_with_template(self) -> None:
        from chint_etm_mdm.db import fetch_products
        from chint_etm_mdm.etm_categories import suggest_classes

        with closing(sqlite3.connect(self.db)) as conn, conn:
            for idx, code in enumerate(["503599", "503599", "503599", "50350599", "50350599"]):
                conn.execute("INSERT INTO products(article, name, series_name) VALUES (?, 'Плата', 'NVF Аксессуары')", (f"a{idx}",))
                conn.execute("INSERT INTO ipro_goods(article, name, full_name, class81_code) VALUES (?, 'x', 'x', ?)", (f"a{idx}", code))
            conn.execute("INSERT INTO products(article, name, series_name) VALUES ('634412', 'Адаптер Modbus TCP', 'NVF Аксессуары')")
        products = fetch_products(self.db, ["634412"])
        self.assertEqual("503599", suggest_classes(self.db, products)["634412"].code)
        choice = suggest_classes(self.db, products, frozenset({"50350599"}))["634412"]
        self.assertEqual(("50350599", False), (choice.code, choice.sure))

    def test_series_defaults_follow_series_and_name(self) -> None:
        from chint_etm_mdm.db import fetch_products
        from chint_etm_mdm.etm_categories import (
            CategoryTemplate,
            load_series_defaults,
            series_default_values,
        )
        from chint_etm_mdm.upload_names import series_names

        with closing(sqlite3.connect(self.db)) as conn, conn:
            conn.execute(
                "INSERT INTO products(article, name, series_name) VALUES ('555', 'АВДТ NB1L-20 S 2P C25 электромеханический', 'NB1L-20')"
            )
        articles = ["586772", "555", "100"]
        products = fetch_products(self.db, articles)
        category = CategoryTemplate("500820", self.root / "x.xlsx", {})
        values = series_default_values(load_series_defaults(), category, articles, products, series_names(self.db, articles))
        self.assertEqual("Нет", values[("586772", "Конфиг:Селективное исполнение")][0])
        self.assertNotIn(("586772", "Конфиг:Тип устройства"), values)
        self.assertNotIn(("555", "Конфиг:Селективное исполнение"), values)
        self.assertEqual("Электромеханический", values[("555", "Конфиг:Тип устройства")][0])
        other = series_default_values(
            load_series_defaults(), CategoryTemplate("50341020", self.root / "x.xlsx", {}), articles, products, {}
        )
        self.assertEqual({}, other)

    def test_new_articles_are_after_the_date(self) -> None:
        self.assertEqual(["586772", "777"], new_articles_since(self.db, dt.date(2026, 9, 4)))

    def test_cards_are_filled_per_class(self) -> None:
        codes = self.root / "chz.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Price-list"
        sheet.append(["PM", None, None])
        sheet.append(["Part number", "ОКПД 2 \n(9 зн.)", "Tariff code"])
        sheet.append([586772, "27.12.22.000", 8536201007])
        workbook.save(codes)

        run = fill_new_products(
            self.db, ["586772", "625288", "777", "404"], self.templates, self.root / "out", codes_path=codes
        )
        self.assertEqual("500820", run.classes["586772"].code)
        self.assertEqual("50341020", run.classes["625288"].code)
        self.assertEqual({"50341020": ["625288"]}, run.missing_templates)
        self.assertEqual(["777"], run.unclassified)
        self.assertEqual(["404"], run.not_found)

        result, csv_path = run.results["500820"]
        # Nothing is ready to upload as is, so there is no empty upload file to send by mistake.
        self.assertEqual("Загрузить в ЭТМ", csv_path.parent.name)
        self.assertFalse(csv_path.exists())
        # The pole count is a match to check, so the card waits in the "to check" folder.
        self.assertEqual(run.output_dir / "На проверку" / f"{csv_path.stem}_check.csv", result.check_path)
        self.assertEqual((0, 1), (result.safe_rows, result.check_rows))
        with result.check_path.open(encoding="cp1251", newline="") as handle:
            rows = list(csv.DictReader(handle, delimiter=";"))
        self.assertEqual(1, len(rows))
        row = rows[0]
        self.assertEqual("АВДТ NB1L-20 1P+N B6 30мА тип AC 6кА", row["Полное название"])
        self.assertEqual("8536201007", row["Код ТН ВЭД"])
        self.assertEqual("27.12.22.000", row["Код ОКПД2"])
        self.assertEqual("0.158", row["Вес, кг"])
        self.assertEqual("0.0925", row["Высота, м"])
        self.assertEqual("Выключатель автоматический дифференциального тока", row["Конфиг:Тип изделия"])
        self.assertEqual("1P+N", row["Конфиг:Количество силовых полюсов"])
        self.assertEqual("6", row["Конфиг:Номинальный ток (А)"])
        self.assertEqual("30", row["Конфиг:Дифференциальный ток (мА)"])
        self.assertEqual("AC", row["Конфиг:Тип срабатывания по диф.току"])
        self.assertEqual("500820", row["81 класс"])
        self.assertEqual(f"upload_new_nb1l-20_{dt.date.today().isoformat()}.csv", csv_path.name)
        self.assertEqual(run.output_dir / "Служебное" / f"{csv_path.stem}.xlsx", result.output_path)
        self.assertEqual(f"{csv_path.stem}_report.xlsx", result.report_path.name)
        self.assertEqual("Служебное", result.report_path.parent.name)

        filled = load_workbook(result.output_path)["Шаблон"]
        self.assertEqual(0.158, filled["L2"].value)
        self.assertEqual("Итог.xlsx", run.summary_path.name)
        summary = load_workbook(run.summary_path)
        self.assertEqual(["Что делать", "Товары", "Дозаполнить", "Проверить", "Нет шаблона"], summary.sheetnames)
        self.assertEqual(5, summary["Товары"].max_row)
        states = {row[0]: row[4] for row in summary["Товары"].iter_rows(min_row=2, values_only=True)}
        self.assertTrue(states["586772"].startswith("Проверить перед загрузкой"))
        self.assertEqual("Нет шаблона класса: скачайте его", states["625288"])
        self.assertEqual("Артикула нет в базе", states["404"])
        self.assertEqual(["50341020"], [row[0] for row in summary["Нет шаблона"].iter_rows(min_row=2, values_only=True)])

    def test_app_worker_finds_and_fills_new_products(self) -> None:
        from chint_etm_mdm.app import NewProductsWorker
        from chint_etm_mdm.etm_categories import run_lines

        results = []
        worker = NewProductsWorker(self.db, dt.date(2026, 9, 4), self.templates, self.root / "out", None, None, None)
        worker.done.connect(results.append)
        worker.failed.connect(self.fail)
        worker.run()
        run = results[0]
        self.assertEqual({"500820"}, set(run.results))
        lines = run_lines(run)
        self.assertTrue(lines[0].startswith("Класс 500820: 1 товаров, можно грузить сразу 0; на проверку 1"))
        self.assertIn("на проверку 1 — файл upload_new_nb1l-20_", lines[0])
        self.assertIn("Не удалось подобрать класс: 1", "\n".join(lines))

        empty = []
        worker = NewProductsWorker(self.db, dt.date(2026, 12, 1), self.templates, self.root / "out2", None, None, None)
        worker.done.connect(empty.append)
        worker.run()
        self.assertEqual([None], empty)


if __name__ == "__main__":
    unittest.main()


class CardsForManagersTests(CategoryTemplateTests):
    """Empty fields go to «Дозаполнить», become questions for managers, and answers fill the next run."""

    def test_empty_fields_are_listed_and_asked(self) -> None:
        from chint_etm_mdm.etm_categories import fill_new_products

        run = fill_new_products(self.db, ["586772"], self.templates, self.root / "out")
        self.assertEqual({"Код ТН ВЭД", "Код ОКПД2"}, {row.column for row in run.missing_rows})
        self.assertEqual({("586772", "Код ТН ВЭД"), ("586772", "Код ОКПД2")}, {(n.article, n.key) for n in run.needs})
        self.assertEqual("NB1L-20", run.needs[0].series)
        empty = load_workbook(run.summary_path)["Дозаполнить"]
        self.assertEqual({"Код ТН ВЭД", "Код ОКПД2"}, {row[3] for row in empty.iter_rows(min_row=2, values_only=True)})
        todo = [row[0] for row in load_workbook(run.summary_path)["Что делать"].iter_rows(values_only=True)]
        self.assertTrue(any("Запросить у продактов" in (line or "") for line in todo))

    def test_answers_from_managers_fill_the_cards(self) -> None:
        import json

        from chint_etm_mdm.etm_categories import fill_new_products

        answers = self.root / "answers.json"
        answers.write_text(
            json.dumps(
                {
                    "version": 1,
                    "answers": [
                        {"article": "586772", "key": "Код ТН ВЭД", "value": "8536201007"},
                        {"article": "586772", "key": "Код ОКПД2", "value": "27.12.22.000"},
                    ],
                }
            ),
            encoding="utf-8",
        )
        run = fill_new_products(self.db, ["586772"], self.templates, self.root / "out", answers_file=answers)
        self.assertEqual([], run.missing_rows)
        self.assertEqual([], run.needs)
        result, _csv = run.results["500820"]
        with result.check_path.open(encoding="cp1251", newline="") as handle:
            row = next(csv.DictReader(handle, delimiter=";"))
        self.assertEqual(("8536201007", "27.12.22.000"), (row["Код ТН ВЭД"], row["Код ОКПД2"]))
        sources = {r.column: r.source for r in result.rows if r.column in {"Код ТН ВЭД", "Код ОКПД2"}}
        self.assertEqual({"pm_answer"}, set(sources.values()))

    def test_downloaded_templates_are_picked_up(self) -> None:
        import os

        from chint_etm_mdm.etm_categories import collect_templates

        downloads = self.root / "Загрузки"
        downloads.mkdir()
        make_category_template(downloads / "upload_goods_category_template_50341020.xlsx")
        (downloads / "другой файл.xlsx").write_bytes(b"")
        copied = collect_templates([downloads, self.root / "нет такой папки"], self.templates)
        self.assertEqual(["upload_goods_category_template_50341020.xlsx"], [path.name for path in copied])
        self.assertEqual([], collect_templates([downloads], self.templates))  # the same file is not copied again
        newer = downloads / "upload_goods_category_template_500820.xlsx"
        make_category_template(newer)
        old = self.templates / newer.name
        os.utime(old, (1, 1))
        self.assertEqual([old], collect_templates([downloads], self.templates))


class ManagersRoundTripTests(CategoryTemplateTests):
    def test_request_answer_and_next_run(self) -> None:
        from chint_etm_mdm.etm_categories import fill_new_products, parse_articles
        from chint_etm_mdm.pm_requests import ANSWER, HEADER_ROW, KEY, answers_path, import_answers, write_requests

        self.assertEqual(["100", "586772", "A-1"], parse_articles(" 100, 586772;\n'A-1 100 "))
        work = self.root / "work"
        run = fill_new_products(self.db, ["100"], self.templates, self.root / "out1")
        asked = {need.key: need for need in run.needs}
        poles = asked["Конфиг:Количество силовых полюсов"]
        self.assertEqual("A", poles.kind)
        self.assertIn("1P+N", poles.allowed)

        requests = write_requests(run.needs, [], work, self.root / "requests", "Карточки ЭТМ")
        request = requests.files[0].path
        workbook = load_workbook(request)
        sheet = workbook["Запрос"]
        for row in sheet.iter_rows(min_row=HEADER_ROW + 1):
            if row[KEY - 1].value == "Конфиг:Количество силовых полюсов":
                row[ANSWER - 1].value = "1p+n"  # written loosely; the program matches it to ETM's list
            if row[KEY - 1].value == "Код ТН ВЭД":
                row[ANSWER - 1].value = "8536 2010 07"
        workbook.save(request)
        answers = import_answers(request, work)
        self.assertEqual(2, answers.accepted)

        again = fill_new_products(self.db, ["100"], self.templates, self.root / "out2", answers_file=answers_path(work))
        left = {row.column for row in again.missing_rows}
        self.assertNotIn("Конфиг:Количество силовых полюсов", left)
        self.assertNotIn("Код ТН ВЭД", left)
        result, _csv = again.results["500820"]
        filled = {row.column: row.value for row in result.rows if row.source == "pm_answer"}
        self.assertEqual({"Конфиг:Количество силовых полюсов": "1P+N", "Код ТН ВЭД": "8536201007"}, filled)
