"""ETM materials: photos, technical information, certificates and descriptions as the manager filled them."""

from __future__ import annotations

import csv
import datetime as dt
import sqlite3
import tempfile
import unittest
from contextlib import closing
from pathlib import Path

from openpyxl import Workbook, load_workbook

from chint_etm_mdm import etm_materials as m
from test_import_data_integrity import make_database
from test_site_info import PAGE_URL, PRODUCT, SEARCH, FakeSite

NUMBER = "ЕАЭС N RU Д-CN.РА06.В.41573/24"
DESCRIPTION = (
    "Преобразователи частоты серии NVF7 предназначены для плавного пуска и регулирования скорости асинхронных "
    "двигателей в системах вентиляции, насосных станциях, конвейерах и другом промышленном оборудовании. "
    "Встроенный тормозной прерыватель и фильтр ЭМС; панель управления; температура работы до 40°C."
)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="cp1251", newline="") as handle:
        return list(csv.DictReader(handle, delimiter=";"))


class MaterialsTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.db = self.root / "db.sqlite"
        make_database(self.db)
        with closing(sqlite3.connect(self.db)) as conn, conn:
            conn.executemany(
                "INSERT INTO products (article, name, series_name) VALUES (?, ?, ?)",
                [("616173", "Преобразователь частоты NVF7-3.0T-S6", "NVF7"), ("777", "Без сайта", "NXB")],
            )
        self.work = self.root / "work"
        self.work.mkdir()
        self.chz = self.root / "chz.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Price-list"
        sheet.append(["Part number", "Серия", "Сертификат 1", "Ссылка на реестр 1", "Сертификат 1 от", "Сертификат 1 до"])
        sheet.append([616173, "NVF7", NUMBER, "https://pub.fsa.gov.ru/rds/declaration/view/1/common", "06.05.2024", None])
        sheet.append([777, "NXB", "не требуется", None, None, None])
        workbook.save(self.chz)
        self.site = FakeSite({"https://ensmas.ru/catalog/?q=616173": SEARCH, PAGE_URL: PRODUCT, "https://ensmas.ru/catalog/?q=777": "<html></html>"})
        self.files: list[str] = []

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def fetch_file(self, url: str) -> bytes:
        self.files.append(url)
        return b"%PDF-1.4 test"

    def run_materials(self):
        return m.make_materials(
            self.db, ["616173", "777"], self.work, self.root / "out", self.chz,
            fetch_page=self.site, fetch_file=self.fetch_file, today=dt.date(2026, 10, 22), site_delay=0,
        )

    def test_forms_files_and_checks(self) -> None:
        # The description of the series, as the user writes it once from the catalog.
        m.write_descriptions(self.work / m.DESCRIPTIONS_TABLE, {"NVF7": {"Описание": DESCRIPTION}}, {}, {})
        result = self.run_materials()

        docs = read_csv(result.files["Тех. информация"])
        self.assertEqual(
            [("616173", "Паспорт", "NVF7 - Паспорт устройства.pdf")],
            [(row["Расширенный Артикул "], row["Название дополнительной технической информации"], row["Наименование файла"]) for row in docs],
        )
        self.assertTrue((result.output_dir / m.FILES_DIR / "NVF7 - Паспорт устройства.pdf").is_file())

        certs = read_csv(result.files["Сертификаты"])
        self.assertEqual(1, len(certs))
        cert = certs[0]
        self.assertEqual(
            ("2019", "616173", NUMBER, "РА06", "ДТРЕАЭС", "06.05.2024", "06.05.2034", "ЕАЭС_N_RU_Д-CN_РА06_В_41573_24.pdf"),
            (cert["Код производителя"], cert["Расширенный артикул"], cert["Номер сертификата"], cert["Код аккредитации"],
             cert["Код типа сертификата"], cert["Начало срока действия"], cert["Окончание срока действия"], cert["Имя файла"]),
        )
        self.assertTrue((result.output_dir / m.FILES_DIR / "ЕАЭС_N_RU_Д-CN_РА06_В_41573_24.pdf").is_file())

        description = read_csv(result.files["Описание"])[0]
        self.assertNotIn(";", description["Описание"])
        self.assertIn("40 градусов Цельсия", description["Описание"])
        self.assertEqual("CHINT", description["Бренд"])

        # The test page's photo is named in Cyrillic: ETM takes it by the link all the same.
        photos = read_csv(result.files["Фото"])
        self.assertEqual(["616173"], [row["Расширенный артикул"] for row in photos])
        self.assertEqual("Фото T7 LED.jpg", photos[0]["Основное Изображение"])
        checks = {(row.article, row.form) for row in result.checks}
        self.assertNotIn(("616173", "Фото"), checks)
        self.assertIn(("777", "Фото"), checks)
        self.assertIn(("777", "Описание"), checks)  # series NXB has no text yet
        # "3D NVF7.rar" has no ETM type: listed for a look.
        self.assertTrue(any("3D NVF7.rar" in row.problem for row in result.checks))

        summary = load_workbook(result.summary_path)
        self.assertEqual(["Что делать", "Проверить"], summary.sheetnames)
        table = m.read_descriptions(self.work / m.DESCRIPTIONS_TABLE)
        self.assertIn("NXB", table)  # the new series got an empty row to fill

        # A second run takes the files from the cache, not from the site again.
        asked = len(self.files)
        self.run_materials()
        self.assertEqual(asked, len(self.files))

    def test_latin_photo_goes_by_link(self) -> None:
        page = PRODUCT.replace("%D0%A4%D0%BE%D1%82%D0%BE%20T7%20LED.jpg", "NVF7_1.jpg")
        self.site.pages[PAGE_URL] = page
        result = self.run_materials()
        photos = read_csv(result.files["Фото"])
        self.assertEqual("NVF7_1.jpg", photos[0]["Основное Изображение"])
        self.assertTrue(photos[0]["Ссылка на скачивание файла"].startswith("https://ensmas.ru/upload/"))
        self.assertEqual("Добавление материалов", photos[0]["Статус"])

    def test_bad_description_goes_to_check(self) -> None:
        m.write_descriptions(self.work / m.DESCRIPTIONS_TABLE, {"NVF7": {"Описание": "Коротко."}}, {}, {})
        result = self.run_materials()
        self.assertNotIn("Описание", result.files)
        self.assertEqual("На проверку", result.files["Описание на проверку"].parent.name)
        self.assertTrue(any("от 250 до 1000" in row.problem for row in result.checks))


class PiecesTests(unittest.TestCase):
    def test_certificate_codes(self) -> None:
        self.assertEqual("СТРЕАЭС", m.certificate_type("ЕАЭС RU С-CN.НВ54.В.02409/22"))
        self.assertEqual("ДТРЕАЭС", m.certificate_type("ЕАЭС N RU Д-CN.РА04.В.00526/25"))
        self.assertEqual("ПОТКАЗ", m.certificate_type("ПИСЬМО 777-0111223"))
        self.assertEqual("", m.certificate_type("не требуется"))
        self.assertEqual("РА04", m.accreditation_code("ЕАЭС N RU Д-CN.PA04.B.00526/25"))  # Latin look-alikes
        self.assertEqual("ЕАЭС_RU_С-CN_НВ54_В_02409_22.pdf", m.certificate_file_name("ЕАЭС RU С-CN.НВ54.В.02409/22"))

    def test_document_types(self) -> None:
        cases = {
            "NVF7 - Каталог.pdf": "Каталог производителя",
            "NVF7 - Руководство по эксплуатации.pdf": "Руководство по эксплуатации",
            "RT28 - Паспорт устройства.pdf": "Паспорт",
            "Оборудование для защиты и управления.pdf": "Техническая информация",
            "Вспомогательный контакт F4-20.dwg": "2D чертеж (формат dwg)",
            "Вспомогательный контакт F4-20.stp": "3D модель (формат stp)",
            "3D NVF7.rar": "",
        }
        for name, kind in cases.items():
            self.assertEqual(kind, m.document_type(name), name)

    def test_clean_text(self) -> None:
        self.assertEqual("До 40 градусов Цельсия, ток - 16 А ''Чинт''", m.clean_text("До 40°C; ток: 16 А\n«Чинт»"))

    def test_import_old_description_file(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            old = root / "old.csv"
            text = DESCRIPTION.replace(";", ",")  # files sent to ETM have no ";" inside texts
            old.write_bytes(
                ("Статус;Код производителя;Расширенный артикул;Бренд;Описание;Преимущества\r\n"
                 f"Добавление материалов;2019;593627;CHINT;{text};\r\n").encode("cp1251")
            )
            table = root / "Описания серий.xlsx"
            m.write_descriptions(table, {"ND16": {"Описание": ""}, "NXB": {"Описание": "Своё"}}, {}, {})
            taken = m.import_description_file(old, table, lambda articles: {"593627": "ND16"})
            self.assertEqual(1, taken)
            texts = m.read_descriptions(table)
            self.assertEqual(text, texts["ND16"]["Описание"])
            self.assertEqual("Своё", texts["NXB"]["Описание"])
            self.assertEqual("old.csv", texts["ND16"]["Откуда"])


if __name__ == "__main__":
    unittest.main()


class EtmReportTests(MaterialsTests):
    """ETM's goods list and certificates report: materials only for what the cards lack."""

    def goods_list(self) -> Path:
        path = self.root / "goodsList_05-10-26_732.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "goodsList"
        sheet.append([
            "Код Товара", "Наименование", "Артикул", "Расширенный Артикул", "Поставляемый", "Наличие Фото",
            "Наличие файлов тех.информации", "Наличие описания", "Наличие 2D/3D модели", "Наличие ТНВЭД", "Наличие ОКПД 2",
        ])
        sheet.append([1, "NVF7", 616173, 616173, "Да", "Да", "Да", "Нет", "Нет", "Нет", "Да"])
        sheet.append([2, "Без сайта", 777, 777, "Да", "Нет", "Да", "Да", "Да", "Да", "Да"])
        sheet.append([3, "Снят", 888, 888, "Нет", "Нет", "Нет", "Нет", "Нет", "Нет", "Нет"])
        sheet.append([4, "Без артикула", None, None, "Да", "Нет", "Нет", "Нет", "Нет", "Нет", "Нет"])
        workbook.save(path)
        return path

    def sert_report(self) -> Path:
        path = self.root / "report_sert_19-10-26_2019.csv"
        path.write_bytes(
            ("Код товара;Полное наименование;Артикул;Расширенный артикул;Код производителя;Сертификат;Причина\r\n"
             "1;NVF7;616173;616173;2019;" + NUMBER + ";У действующего сертификата отсутствует обязательная URL-ссылка на запись в реестре\r\n"
             "2;Без сайта;777;777;2019;отсутствует;нет сертификата (или нет ни одного обязательного)\r\n").encode("cp1251")
        )
        return path

    def test_reports_are_read(self) -> None:
        needs = m.ReportNeeds()
        needs.add(m.read_etm_report(self.goods_list()))
        needs.add(m.read_etm_report(self.sert_report()))
        self.assertEqual({"616173"}, needs.wanted["Описание"])
        self.assertEqual({"616173"}, needs.wanted["2D/3D"])
        self.assertEqual({"777"}, needs.wanted["Фото"])  # 888 is no longer supplied, 4 has no article
        self.assertNotIn("Тех. информация", needs.wanted)
        self.assertEqual({"616173", "777"}, needs.wanted["Сертификаты"])
        self.assertEqual({"616173": ["Код ТН ВЭД"]}, needs.codes_missing)
        self.assertIn("Товаров без фото: 1", m.report_lines(needs))
        with self.assertRaisesRegex(ValueError, "не знаю такой отчёт"):
            other = self.root / "other.csv"
            other.write_text("Расширенный артикул;Цена\n1;2\n", encoding="utf-8")
            m.read_etm_report(other)

    def test_materials_only_for_what_is_missing(self) -> None:
        m.write_descriptions(self.work / m.DESCRIPTIONS_TABLE, {"NVF7": {"Описание": DESCRIPTION}}, {}, {})
        needs = m.ReportNeeds()
        needs.add(m.read_etm_report(self.goods_list()))
        needs.add(m.read_etm_report(self.sert_report()))
        result = m.make_materials(
            self.db, [], self.work, self.root / "out", self.chz, fetch_page=self.site, fetch_file=self.fetch_file,
            today=dt.date(2026, 10, 22), site_delay=0, wanted=needs.wanted, codes_missing=needs.codes_missing,
        )
        # 616173 has its documents in ETM: only drawings and models would go, and the test page has none.
        self.assertNotIn("Тех. информация", result.files)
        self.assertEqual(["616173"], [row["Расширенный артикул"] for row in read_csv(result.files["Описание"])])
        self.assertEqual(["616173"], [row["Расширенный артикул"] for row in read_csv(result.files["Сертификаты"])])
        # 777: no photo on the site; "не требуется" in ЧЗ while ETM asks for a document.
        checks = {(row.article, row.form): row.problem for row in result.checks}
        self.assertIn(("777", "Фото"), checks)
        self.assertIn("ЭТМ просит документ", checks[("777", "Сертификаты")])
        self.assertNotIn(("616173", "Фото"), checks)  # the card has a photo: nothing asked
        codes = load_workbook(result.summary_path)["Коды в карточке"]
        self.assertEqual(("616173", "Код ТН ВЭД"), (codes["A2"].value, codes["C2"].value))


def make_cert_folder(folder: Path, rows: list[list[object]], pdfs: list[str]) -> None:
    """The certification team's folder: «Перечень документов….xlsx» and the PDFs in subfolders."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "LV"
    sheet.append(["Производитель", "Серия", "Тип документа", "Номер сертификата ", "Выдан", "Действ. до ",
                  "Название документа в папке VALID CERTIFICATES"])
    for row in rows:
        sheet.append(row)
    folder.mkdir(parents=True, exist_ok=True)
    workbook.save(folder / "Перечень документов upd 2026-10-12.xlsx")
    for name in pdfs:
        path = folder / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(b"%PDF-1.4 " + name.encode())


class CertificatesFolderTests(MaterialsTests):
    def test_dates_and_pdfs_from_the_certificates_folder(self) -> None:
        workbook = load_workbook(self.chz)
        workbook.active.append([777, "NXB", "ОР-25/06-0035", None, None, None])  # a refusal letter, no dates anywhere else
        workbook.save(self.chz)
        folder = self.root / "Сертификаты"
        make_cert_folder(
            folder,
            [
                ["CHINT", "NVF7", "ДС 004", NUMBER, dt.datetime(2024, 5, 6), dt.datetime(2029, 5, 5), "ДС 004 NVF7"],
                ["CHINT", "LMI", "ОР 004", "ОР-25/06-0035", dt.datetime(2025, 6, 30), "бессрочно", "ОР_004 LMI механические блокировки"],
            ],
            ["Актуальные декларации ТР ТС/ДС 004 NVF7.pdf", "Отрицательные решения/ОР_004 LMI механические блокировки.pdf"],
        )
        result = m.make_materials(
            self.db, ["616173", "777"], self.work, self.root / "out", self.chz, fetch_page=self.site,
            fetch_file=self.fetch_file, today=dt.date(2026, 10, 22), site_delay=0, cert_folder=folder,
        )
        rows = {row["Расширенный артикул"]: row for row in read_csv(result.files["Сертификаты"])}
        self.assertEqual(("06.05.2024", "05.05.2029"), (rows["616173"]["Начало срока действия"], rows["616173"]["Окончание срока действия"]))
        letter = rows["777"]
        self.assertEqual(("ПОТКАЗ", "", ""), (letter["Код типа сертификата"], letter["Код аккредитации"], letter["Ссылка на запись документа в реестре"]))
        self.assertEqual(("30.06.2025", "30.06.2035"), (letter["Начало срока действия"], letter["Окончание срока действия"]))  # «бессрочно»: +10 years
        files = result.output_dir / m.FILES_DIR
        self.assertIn("ОР_004 LMI".encode(), (files / letter["Имя файла"]).read_bytes())
        self.assertIn("ДС 004 NVF7".encode(), (files / rows["616173"]["Имя файла"]).read_bytes())  # the folder's file, not the site's
        self.assertFalse(any(row.form == "Сертификаты" for row in result.checks))
