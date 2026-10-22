"""FIL-02/03 and MAN-01: unknown products, hand-entered values and putting the base's value back."""

from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from chint_etm_mdm.db import ProductRecord
from chint_etm_mdm.filler import MISSING_CELL_FILL, fill_xlsx_template
from chint_etm_mdm.manual_values import import_manual_completions, load_manual_values

YELLOW = PatternFill(fill_type="solid", fgColor="FFFF00")
HEADERS = ["Артикул", "Название", "Полное название", "Вес, кг"]


def write_book(path: Path, rows: list[list[object]]) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.fill = YELLOW
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()
    return path


def write_manual(path: Path, values: list[tuple[str, str, str]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps({"version": 1, "values": [{"article": a, "field": f, "value": v} for a, f, v in values]}),
        encoding="utf-8",
    )
    return path


class UnknownProductTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def fill(self, template: Path, manual: Path | None = None):
        with patch("chint_etm_mdm.filler.fetch_products", return_value={}):
            return fill_xlsx_template(self.root / "db.sqlite", template, self.root / "out", None, manual)

    def test_filled_cells_of_an_unknown_product_stay_as_they_are(self) -> None:
        template = self.root / "template.xlsx"
        write_book(template, [["X-9", "Уже есть", None, None]])
        workbook = load_workbook(template)
        workbook.active["B2"].font = Font(name="Arial", size=14, bold=True)
        workbook.active["C2"].font = Font(name="Arial", size=14, bold=True)
        workbook.save(template)
        workbook.close()

        result = self.fill(template)

        sheet = load_workbook(result.output_path).active
        self.assertEqual("Уже есть", sheet["B2"].value)
        self.assertNotEqual(MISSING_CELL_FILL.fgColor.rgb, sheet["B2"].fill.fgColor.rgb)
        self.assertEqual(MISSING_CELL_FILL.fgColor.rgb, sheet["C2"].fill.fgColor.rgb)
        # The red marking keeps the template's typeface.
        self.assertEqual(("Arial", 14, True), (sheet["C2"].font.name, sheet["C2"].font.sz, sheet["C2"].font.b))
        self.assertEqual("FF9C0006", sheet["C2"].font.color.rgb)
        self.assertEqual(2, result.missing_cells)  # "Полное название" and "Вес", not the filled name

    def test_hand_entered_value_fills_an_unknown_product(self) -> None:
        template = write_book(self.root / "template.xlsx", [["X-9", None, None, None]])
        manual = write_manual(self.root / "rules" / "manual_values.json", [("X-9", "Название", "Ручное имя")])

        result = self.fill(template, manual)

        sheet = load_workbook(result.output_path).active
        self.assertEqual("Ручное имя", sheet["B2"].value)
        self.assertNotEqual(MISSING_CELL_FILL.fgColor.rgb, sheet["B2"].fill.fgColor.rgb)
        self.assertIsNone(sheet["C2"].value)
        self.assertEqual(1, result.filled_cells)


class PutBaseValueBackTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        (self.root / "db.sqlite").write_bytes(b"")
        self.manual = write_manual(self.root / "rules" / "manual_values.json", [("A-1", "Название", "Ручное")])
        self.products = {"A-1": ProductRecord(article="A-1", name="Из базы")}

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def import_book(self, name_value: object):
        filled = write_book(self.root / "filled.xlsx", [["A-1", name_value, None, None]])
        with patch("chint_etm_mdm.manual_values.fetch_products", return_value=self.products):
            return import_manual_completions(self.root / "db.sqlite", filled, self.manual, self.root / "out")

    def test_base_value_typed_back_removes_the_old_hand_value(self) -> None:
        result = self.import_book("Из базы")
        self.assertEqual(1, result.removed_values)
        self.assertEqual({}, load_manual_values(self.manual))
        report = load_workbook(result.report_path).active
        self.assertIn("снято", [row[4] for row in report.iter_rows(min_row=2, values_only=True)][0])

    def test_empty_cell_does_not_remove_it(self) -> None:
        result = self.import_book(None)
        self.assertEqual(0, result.removed_values)
        self.assertEqual({("A-1", "Название"): "Ручное"}, load_manual_values(self.manual))

    def test_same_file_again_is_not_new(self) -> None:
        result = self.import_book("Ручное")
        self.assertEqual((0, 1, 0), (result.imported_values, result.skipped_values, result.removed_values))
        changed = self.import_book("Другое")
        self.assertEqual(1, changed.imported_values)
        self.assertEqual({("A-1", "Название"): "Другое"}, load_manual_values(self.manual))


if __name__ == "__main__":
    unittest.main()
