"""The check of sizes already in the base: ETIM values from wrong columns are found and cleared."""

from __future__ import annotations

import sqlite3
import tempfile
import unittest
from contextlib import closing
from pathlib import Path

from openpyxl import load_workbook

from chint_etm_mdm import etim_importer
from chint_etm_mdm.dimension_audit import audit_dimensions, audit_lines, clear_lines, clear_wrong_dimensions


def make_base(path: Path) -> None:
    with closing(sqlite3.connect(path)) as conn, conn:
        conn.execute("CREATE TABLE products (article TEXT PRIMARY KEY)")
        etim_importer.ensure_etim_import_tables(conn)
        raw = [
            # What an earlier version did: cable length taken for the product length.
            ("311015", "length_mm", 1_000_000, "Лист19", "Макс. длина кабеля, м"),
            ("311015", "height_mm", 80, "Лист19", "Высота, мм"),
            ("A-2", "width_mm", 45, "Лист3", "Ширина шины, мм"),
            ("A-3", "length_mm", 120, "Лист1", "Длина, мм"),
            ("A-4", "length_mm", 50_000, "Лист1", "Длина, м"),
            ("A-5", "length_mm", 300, "Лист1", "Габаритные размеры ДхШхВ, мм"),
            ("A-8", "width_mm", 30, "Лист1", "Габаритные размеры, мм"),
            ("A-9", "height_mm", 0.08, "Лист1", "Высота"),
        ]
        for article, column, value, sheet, field in raw:
            conn.execute(
                f"INSERT INTO product_dimensions_raw (import_id, article, source_kind, dimension_kind, {column}, source_sheet, source_field) "
                "VALUES (1, ?, ?, ?, ?, ?, ?)",
                (article, etim_importer.SOURCE_KIND, column, value, sheet, field),
            )
        conn.executemany(
            "INSERT INTO product_dimensions_resolved (article, length_mm, width_mm, height_mm) VALUES (?, ?, ?, ?)",
            [
                ("311015", 1_000_000, None, 80),
                ("A-2", None, 45, None),
                ("A-3", 120, None, None),
                ("A-4", 50_000, None, None),
                ("A-5", 300, None, None),
                ("A-6", 20_000, None, None),  # not from ETIM: only reported
                ("A-7", 115, None, None),  # the base had its own value; ETIM's wrong one went to conflicts only
                ("A-8", None, 30, None),
                ("A-9", None, None, 0.08),
            ],
        )


def resolved(path: Path, article: str) -> tuple:
    with closing(sqlite3.connect(path)) as conn:
        return conn.execute(
            "SELECT length_mm, width_mm, height_mm FROM product_dimensions_resolved WHERE article = ?", (article,)
        ).fetchone()


class DimensionAuditTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.db = self.root / "base" / "chint_mdm.sqlite"
        self.db.parent.mkdir()
        make_base(self.db)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_finds_values_from_wrong_columns(self) -> None:
        result = audit_dimensions(self.db, self.root / "reports")
        found = {(f.article, f.column): f for f in result.findings}
        self.assertEqual(
            {
                ("311015", "length_mm"), ("A-2", "width_mm"), ("A-4", "length_mm"), ("A-6", "length_mm"),
                ("A-8", "width_mm"), ("A-9", "height_mm"),
            },
            set(found),
        )
        self.assertFalse(found[("A-8", "width_mm")].fixable)
        self.assertIn("Ш×В×Д", found[("A-8", "width_mm")].problem)
        self.assertFalse(found[("A-9", "height_mm")].fixable)
        self.assertTrue(found[("311015", "length_mm")].fixable)
        self.assertIn("Макс. длина кабеля, м", found[("311015", "length_mm")].source)
        self.assertTrue(found[("A-4", "length_mm")].fixable)  # 50 m from ETIM
        self.assertFalse(found[("A-6", "length_mm")].fixable)  # not from ETIM: the user decides
        self.assertEqual(3, len(result.fixable))
        self.assertEqual(9, result.products_checked)

        sheet = load_workbook(result.report_path)["Проверить"]
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        self.assertEqual(6, len(rows))
        self.assertTrue(any("Макс. длина кабеля, м" in line for line in audit_lines(result)))

    def test_check_only_reads(self) -> None:
        before = self.db.read_bytes()
        audit_dimensions(self.db, self.root / "reports")
        self.assertEqual(before, self.db.read_bytes())

    def test_clearing_keeps_good_values_and_makes_a_backup(self) -> None:
        result = audit_dimensions(self.db, self.root / "reports")
        cleared = clear_wrong_dimensions(self.db, result.findings)
        self.assertEqual(3, cleared.cleared)
        self.assertTrue(cleared.backup_path.exists())
        self.assertEqual((None, None, 80.0), resolved(self.db, "311015"))
        self.assertEqual((None, None, None), resolved(self.db, "A-2"))
        self.assertEqual((120.0, None, None), resolved(self.db, "A-3"))
        self.assertEqual((300.0, None, None), resolved(self.db, "A-5"))
        self.assertEqual((20000.0, None, None), resolved(self.db, "A-6"))
        self.assertEqual((115.0, None, None), resolved(self.db, "A-7"))
        self.assertTrue(any("импортируйте ETIM" in line for line in clear_lines(cleared)))
        # Run again: nothing left to clear.
        self.assertEqual(0, len(audit_dimensions(self.db, self.root / "reports").fixable))

    def test_value_changed_after_the_check_is_left_alone(self) -> None:
        result = audit_dimensions(self.db, self.root / "reports")
        with closing(sqlite3.connect(self.db)) as conn, conn:
            conn.execute("UPDATE product_dimensions_resolved SET length_mm = 130 WHERE article = '311015'")
        cleared = clear_wrong_dimensions(self.db, result.findings)
        self.assertEqual((2, 1), (cleared.cleared, cleared.changed_since_check))
        self.assertEqual(130.0, resolved(self.db, "311015")[0])

    def test_old_report_cannot_bring_the_cable_length_back(self) -> None:
        from openpyxl import Workbook

        report = self.root / "old_report.xlsx"
        workbook = Workbook()
        summary = workbook.active
        summary.title = "Итог"
        summary.append(["Показатель", "Значение"])
        summary.append(["Импорт id", 1])
        sheet = workbook.create_sheet("Конфликты")
        sheet.append(["Решение", "Артикул", "Поле", "Сейчас в базе", "Пришло из ETIM"])
        sheet.append([etim_importer.ACCEPT_ETIM_DECISION, "A-7", "Длина, мм", "115", "1000000"])
        workbook.save(report)
        with closing(sqlite3.connect(self.db)) as conn, conn:
            conn.execute("INSERT INTO products VALUES ('A-7')")
            conn.execute(
                "INSERT INTO product_dimensions_raw (import_id, article, source_kind, dimension_kind, length_mm, source_field) "
                "VALUES (1, 'A-7', ?, 'length_mm', 1000000, 'Макс. длина кабеля, м')",
                (etim_importer.SOURCE_KIND,),
            )
        result = etim_importer.apply_etim_conflict_decisions(report, self.db)
        self.assertEqual(0, result.accepted_rows)
        self.assertIn("не размер товара", result.rejected[0])
        self.assertEqual((115.0, None, None), resolved(self.db, "A-7"))


if __name__ == "__main__":
    unittest.main()


class AuditWindowTests(unittest.TestCase):
    def test_workers_report_and_clear(self) -> None:
        from chint_etm_mdm.app import DimensionAuditWorker, DimensionClearWorker

        with tempfile.TemporaryDirectory() as tmp:
            db = Path(tmp) / "chint_mdm.sqlite"
            make_base(db)
            audits = []
            worker = DimensionAuditWorker(db, Path(tmp) / "reports")
            worker.done.connect(audits.append)
            worker.failed.connect(self.fail)
            worker.run()
            cleared = []
            clear = DimensionClearWorker(db, audits[0].fixable)
            clear.done.connect(cleared.append)
            clear.failed.connect(self.fail)
            clear.run()
            self.assertEqual(3, cleared[0].cleared)
