import datetime as dt
import sqlite3
import tempfile
import unittest
from contextlib import closing
from pathlib import Path

from openpyxl import Workbook, load_workbook

from chint_etm_mdm.certificates import (
    FROM_DB,
    FROM_FILE,
    FROM_USER,
    chz_date_columns,
    number_key,
    table_dates,
    table_lines,
    update_table,
)
from test_import_data_integrity import make_database

DECL = "ЕАЭС N RU Д-CN.РА06.В.41573/24"
CERT = "ЕАЭС RU С-CN.НВ12.В.00413/23"


def make_chz(path: Path, dates: bool = False) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Price-list"
    sheet.append(["PM", None, None, "Спец. по сертификации"])
    header = ["Part number", "Russian description", "Серия", "Сертификат 1", "Ссылка на реестр 1"]
    if dates:
        header += ["Сертификат 1 от", "Сертификат 1 до"]
    sheet.append(header)
    rows = [
        [173585, "Пускатель NKB1-45", "NKB1", CERT, "https://pub.fsa.gov.ru/rss/certificate/view/3365945/baseInfo"],
        [173584, "Пускатель NKB1-45 Q", "NKB1", CERT.replace("C", "С"), None],
        [391935, "Блокировка LMI-12", "NXB", "ОР-25/06-0035", None],
        [210487, "Дроссель ACL", None, "не требуется", None],
    ]
    for row in rows:
        if dates:
            row += [dt.datetime(2023, 5, 1), dt.datetime(2028, 4, 30)] if row[3] == CERT else [None, None]
        sheet.append(row)
    workbook.save(path)


class CertificateTableTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.db = self.root / "mdm.sqlite"
        make_database(self.db)
        with closing(sqlite3.connect(self.db)) as conn, conn:
            conn.executemany(
                "INSERT INTO product_attribute_values(article, attribute_name, value_raw, confidence) VALUES (?, ?, ?, 1)",
                [
                    ("616173", "Номер сертификата", DECL),
                    ("616173", "Дата начала", "2024-07-25 00:00:00"),
                    ("616173", "Дата окончания", "2029-07-25 00:00:00"),
                    ("616172", "Номер сертификата", DECL),
                    ("173585", "Номер сертификата", CERT),
                ],
            )
        self.chz = self.root / "Список товаров под ЧЗ_19.10.2026.xlsx"
        self.table = self.root / "work" / "Сертификаты.xlsx"

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def rows(self) -> dict[str, tuple]:
        sheet = load_workbook(self.table)["Сертификаты"]
        return {row[0]: row for row in sheet.iter_rows(min_row=2, values_only=True)}

    def test_table_once_per_document(self) -> None:
        make_chz(self.chz)
        result = update_table(self.db, self.chz, self.table, today=dt.date(2026, 10, 22))
        rows = self.rows()
        # Two ЕАЭС documents: the refusal letter and "не требуется" are not certificates.
        self.assertEqual({DECL, CERT}, set(rows))
        self.assertEqual(("Декларация ЕАС", dt.datetime(2024, 7, 25), dt.datetime(2029, 7, 25), FROM_DB, 2), rows[DECL][1:6])
        self.assertEqual(("Сертификат ЕАС", None, None, None, 2, "NKB1"), rows[CERT][1:7])  # Latin "C" is the same document
        self.assertEqual((2, 1, 1, 2), (result.documents, result.with_dates, result.without_dates, result.added))
        self.assertTrue(any("жёлтые клетки" in line for line in table_lines(result)))

        # The user types the dates of the certificate; the next update keeps them and adds nothing.
        workbook = load_workbook(self.table)
        sheet = workbook["Сертификаты"]
        for row in sheet.iter_rows(min_row=2):
            if row[0].value == CERT:
                row[2].value, row[3].value = dt.datetime(2023, 5, 1), dt.datetime(2026, 12, 1)
        workbook.save(self.table)
        again = update_table(self.db, self.chz, self.table, today=dt.date(2026, 10, 22))
        rows = self.rows()
        self.assertEqual((dt.datetime(2023, 5, 1), dt.datetime(2026, 12, 1), FROM_USER), rows[CERT][2:5])
        self.assertEqual(0, again.added)
        self.assertEqual([(CERT, dt.date(2026, 12, 1))], again.expiring)  # within 90 days
        self.assertTrue(any(line.startswith("Истекает 01.12.2026") for line in table_lines(again)))
        self.assertEqual((dt.date(2023, 5, 1), dt.date(2026, 12, 1)), table_dates(self.table)[number_key(CERT)])

    def test_dates_from_chz_columns(self) -> None:
        make_chz(self.chz, dates=True)
        update_table(self.db, self.chz, self.table)
        self.assertEqual((dt.datetime(2023, 5, 1), dt.datetime(2028, 4, 30), "файл ЧЗ"), self.rows()[CERT][2:5])

    def test_dates_from_an_earlier_file(self) -> None:
        make_chz(self.chz)
        earlier = self.root / "Запрос сертификатов_Таблица_заполнено.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Обязательные поля для заполнения", None, None, None, None, None, None, "Обязательные поля для заполнения"])
        sheet.append(["Код РС (если известен)", "Наименование", "Артикул", "Номер сертификата", "Начало действия сертификата",
                      "Окончание действия сертификата", "Тип сертификата", "Имя файла из архива или ссылка для скачивания",
                      "Прямая ссылка на страницу документа в реестре Росаккредитации", "Код сертифицирующего органа"])
        for article in (173585, 173584):
            sheet.append([article, "Пускатель", article, CERT, dt.datetime(2023, 12, 22), dt.datetime(2028, 12, 21), "Сертификат"])
        sheet.append([616173, "Частотник", 616173, DECL, dt.datetime(2024, 7, 1), dt.datetime(2029, 7, 1), "Декларация"])
        workbook.save(earlier)

        # The user already typed the declaration's dates by hand: they stay.
        update_table(self.db, self.chz, self.table)
        workbook = load_workbook(self.table)
        for row in workbook["Сертификаты"].iter_rows(min_row=2):
            if row[0].value == DECL:
                row[2].value, row[3].value = dt.datetime(2024, 7, 26), dt.datetime(2029, 7, 26)
        workbook.save(self.table)

        result = update_table(self.db, self.chz, self.table, dates_file=earlier)
        rows = self.rows()
        self.assertEqual((dt.datetime(2023, 12, 22), dt.datetime(2028, 12, 21), FROM_FILE), rows[CERT][2:5])
        self.assertEqual((dt.datetime(2024, 7, 26), dt.datetime(2029, 7, 26), FROM_USER), rows[DECL][2:5])
        self.assertEqual(1, result.from_file)
        self.assertTrue(any("Из выбранного файла" in line for line in table_lines(result)))

        # Later updates without the file keep its dates.
        update_table(self.db, self.chz, self.table)
        self.assertEqual((dt.datetime(2023, 12, 22), dt.datetime(2028, 12, 21), FROM_FILE), self.rows()[CERT][2:5])

    def test_date_column_names(self) -> None:
        def found(names):
            return chz_date_columns([name.casefold() for name in names])

        base = ["part number", "сертификат 1", "ссылка на реестр 1", "сертификат 2"]
        self.assertEqual({"start": 4, "end": 5}, found(base + ["Сертификат 1 от", "Сертификат 1 до"]))
        self.assertEqual({"start": 4, "end": 5}, found(base + ["Дата начала 1", "Дата окончания 1"]))
        self.assertEqual({"start": 4, "end": 5}, found(base + ["Действует с", "Действует до"]))
        self.assertEqual({}, found(base + ["Сертификат 2 от", "Сертификат 2 до", "Складской статус"]))


class RaekUsesTableTests(unittest.TestCase):
    def test_raek_dates_from_table(self) -> None:
        from chint_etm_mdm.db import ProductRecord
        from chint_etm_mdm.raek import Sources, canonical, certificate

        sources = Sources(
            products={},
            chz={("173585", "cert_number"): CERT},
            series={},
            coded={},
            settings={},
            learned={},
            cert_dates={number_key(CERT): (dt.date(2023, 5, 1), dt.date(2028, 4, 30))},
        )
        cert = certificate("173585", ProductRecord("173585", attributes={}), sources)
        self.assertEqual(dt.datetime(2023, 5, 1), cert[canonical("Сертификат от")].value)
        self.assertEqual(dt.datetime(2028, 4, 30), cert[canonical("Сертификат до")].value)
        empty = certificate("173585", ProductRecord("173585", attributes={}), Sources({}, sources.chz, {}, {}, {}, {}))
        self.assertIn("таблицу сертификатов", empty[canonical("Сертификат от")].note)


if __name__ == "__main__":
    unittest.main()
