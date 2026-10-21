from __future__ import annotations

import math
import sqlite3
import tempfile
import unittest
from contextlib import closing
from pathlib import Path

from openpyxl import Workbook

from chint_etm_mdm import db, etim_importer
from chint_etm_mdm.number_parsing import parse_number
from chint_etm_mdm.price_importer import PRICE_SHEETS, import_price_workbook, read_price_workbook, result_lines


SCHEMA = """
CREATE TABLE products (
    article TEXT PRIMARY KEY, name TEXT, series_name TEXT, equipment_type TEXT,
    equipment_group TEXT, price_collection TEXT, unit TEXT, stock_status TEXT,
    is_new TEXT, is_promo TEXT, updated_at TEXT
);
CREATE TABLE price_snapshots (
    id INTEGER PRIMARY KEY, snapshot_date TEXT, source_url TEXT, local_file TEXT, workbook_name TEXT
);
CREATE TABLE price_snapshot_items (
    id INTEGER PRIMARY KEY, snapshot_id INTEGER, article TEXT, name TEXT, price_with_vat TEXT,
    price_without_vat TEXT, unit TEXT, price_collection TEXT, stock_status TEXT, is_removed TEXT,
    is_new TEXT, is_promo TEXT, equipment_type TEXT, equipment_group TEXT, series_name TEXT,
    pack_multiple TEXT, min_shipment TEXT, gross_weight_transport TEXT, net_weight_transport TEXT,
    transport_volume TEXT, gross_weight_min TEXT, net_weight_min TEXT, min_volume TEXT,
    gross_weight_unit TEXT, net_weight_unit TEXT, unit_volume TEXT, outgoing_stock TEXT,
    analog_article TEXT, analog_name TEXT
);
CREATE TABLE new_products (article TEXT PRIMARY KEY, name TEXT, introduced_at TEXT, first_snapshot_id INTEGER);
CREATE TABLE assortment_outgoing (article TEXT, name TEXT, removed_at TEXT, snapshot_id INTEGER);
CREATE TABLE price_change_history (
    article TEXT, name TEXT, change_type TEXT, old_value TEXT, new_value TEXT, changed_at TEXT, snapshot_id INTEGER
);
CREATE TABLE ipro_goods (id INTEGER PRIMARY KEY, article TEXT, name TEXT, full_name TEXT, class81_code TEXT, manufacturer TEXT);
CREATE TABLE product_etm_class_suggestions (article TEXT, class81_code TEXT, confidence REAL);
CREATE TABLE product_tnved_codes (id INTEGER PRIMARY KEY, article TEXT, tnved_code TEXT, import_id INTEGER);
CREATE TABLE product_okpd2_codes (id INTEGER PRIMARY KEY, article TEXT, okpd2_9 TEXT, okpd2_6 TEXT, import_id INTEGER);
CREATE TABLE material_data_items (
    id INTEGER PRIMARY KEY, article TEXT, gross_weight_kg REAL, outer_length_mm REAL, outer_width_mm REAL,
    outer_height_mm REAL, gtin_inner TEXT, gtin_outer TEXT
);
"""


def make_database(path: Path) -> None:
    with closing(sqlite3.connect(path)) as conn, conn:
        conn.executescript(SCHEMA)
        etim_importer.ensure_etim_import_tables(conn)


def write_price(path: Path, header: list[str], rows: list[list[object]]) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name in PRICE_SHEETS:
        sheet = workbook.create_sheet(sheet_name)
        if sheet_name == "Price-list":
            sheet.append(header)
            for row in rows:
                sheet.append(row)
        else:
            sheet.append(["Артикул", "Наименование"])
    workbook.save(path)
    workbook.close()
    return path


FULL_HEADER = [
    "Артикул", "Наименование", "Серия", "Ед.", "Группа оборудования",
    "Вес брутто за 1 ед. прод.", "Вес  нетто за 1 ед. прод.", "Объём, за 1 ед. прод.",
]


class PartialPriceTests(unittest.TestCase):
    """PRI-02: a column missing from the price must not wipe product fields."""

    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)
        self.database = self.root / "db.sqlite"
        make_database(self.database)

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def product(self, article: str = "A-1") -> sqlite3.Row:
        with closing(sqlite3.connect(self.database)) as conn:
            conn.row_factory = sqlite3.Row
            return conn.execute("SELECT * FROM products WHERE article = ?", (article,)).fetchone()

    def test_partial_price_keeps_fields_and_weights_from_full_price(self) -> None:
        full = write_price(
            self.root / "Price-list-CHINT_2026-01-01.xlsx",
            FULL_HEADER,
            [["A-1", "Автомат", "NXB", "шт", "Модульное", 0.25, 0.2, 0.001]],
        )
        import_price_workbook(full, self.database, make_backup=False)
        partial = write_price(
            self.root / "Price-list-CHINT_2026-02-01.xlsx",
            ["Артикул", "Наименование"],
            [["A-1", "Автомат новый"]],
        )

        result = import_price_workbook(partial, self.database, make_backup=False)

        product = self.product()
        self.assertEqual("Автомат новый", product["name"])
        self.assertEqual("NXB", product["series_name"])
        self.assertEqual("шт", product["unit"])
        self.assertEqual("Модульное", product["equipment_group"])
        record = db.fetch_products(self.database, ["A-1"])["A-1"]
        self.assertEqual(0.25, record.weight_kg)
        self.assertEqual(0.001, record.volume_m3)
        self.assertIn("Серия", result.missing_columns)
        self.assertTrue(any("прежние значения в базе сохранены" in line for line in result_lines(result)))

    def test_present_column_still_updates_and_empty_name_does_not_wipe(self) -> None:
        with closing(sqlite3.connect(self.database)) as conn, conn:
            conn.execute("INSERT INTO products (article, name, series_name, unit) VALUES ('A-1', 'Автомат', 'NXB', 'шт')")
        price = write_price(
            self.root / "Price-list-CHINT_2026-03-01.xlsx",
            ["Артикул", "Наименование", "Серия"],
            [["A-1", "", "NXB2"]],
        )

        import_price_workbook(price, self.database, make_backup=False)

        product = self.product()
        self.assertEqual("Автомат", product["name"])
        self.assertEqual("NXB2", product["series_name"])
        self.assertEqual("шт", product["unit"])

    def test_header_variants_map_to_the_same_columns(self) -> None:
        price = write_price(
            self.root / "Price-list-CHINT_2026-04-01.xlsx",
            [
                "﻿Артикул", "Наименование", "Ед.\xa0",
                "Вес нетто за 1\nед. прод.", "Объем, за 1 ед. прод.", "Объём транспортной упаковки, м³",
                "Фото",
            ],
            [["A-1", "Автомат", "шт", 0.2, 0.001, 0.05]],
        )

        result = import_price_workbook(price, self.database, make_backup=False)

        record = db.fetch_products(self.database, ["A-1"])["A-1"]
        self.assertEqual(0.2, record.weight_kg)
        self.assertEqual(0.001, record.volume_m3)
        self.assertEqual(0.05, record.transport_volume_m3)
        self.assertEqual("шт", self.product()["unit"])
        self.assertEqual(("Фото",), result.unknown_columns)

    def test_duplicate_significant_columns_are_rejected_before_database_changes(self) -> None:
        price = write_price(
            self.root / "Price-list-CHINT_2026-05-01.xlsx",
            ["Артикул", "Наименование", "Серия", " Серия "],
            [["A-1", "Автомат", "NXB", "NXA"]],
        )

        with self.assertRaisesRegex(ValueError, "повторяются колонки: Серия"):
            read_price_workbook(price)
        with self.assertRaises(ValueError):
            import_price_workbook(price, self.database, make_backup=False)
        self.assertIsNone(self.product())


class PartialPriceStatisticsTests(unittest.TestCase):
    """A partial price must not become the baseline for the next import's counters."""

    def test_counters_after_partial_price_compare_with_last_full_price(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            database = root / "db.sqlite"
            make_database(database)
            header = ["Артикул", "Наименование", "Тариф с НДС, руб", "Тариф без НДС, руб"]
            rows = [["A-1", "Автомат", "120", "100"], ["A-2", "Щит", "240", "200"]]
            full = write_price(root / "Price-list-CHINT_2026-07-01.xlsx", header, rows)
            import_price_workbook(full, database, make_backup=False)
            partial = write_price(
                root / "Price-list-CHINT_2026-07-15.xlsx", ["Артикул", "Наименование"], [["A-1", "Автомат"]]
            )

            partial_result = import_price_workbook(partial, database, make_backup=False)
            next_full = write_price(
                root / "Price-list-CHINT_2026-08-01.xlsx", header, [*rows, ["A-3", "Новый", "60", "50"]]
            )
            full_result = import_price_workbook(next_full, database, make_backup=False)

            self.assertEqual(0, partial_result.changed_price_rows)
            self.assertEqual(1, full_result.new_articles)
            self.assertEqual(0, full_result.changed_price_rows)

    def test_import_adds_index_for_product_reads(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            database = root / "db.sqlite"
            make_database(database)
            price = write_price(root / "Price-list-CHINT_2026-09-01.xlsx", ["Артикул", "Наименование"], [["A-1", "Автомат"]])

            import_price_workbook(price, database, make_backup=False)

            with closing(sqlite3.connect(database)) as conn:
                plan = " ".join(
                    str(row[-1])
                    for row in conn.execute(
                        "EXPLAIN QUERY PLAN SELECT gross_weight_unit FROM price_snapshot_items "
                        "WHERE article = 'A-1' ORDER BY snapshot_id DESC, id DESC LIMIT 1"
                    )
                )
            self.assertIn("idx_price_snapshot_items_article_snapshot", plan)


class NumericTextTests(unittest.TestCase):
    """PRI-03: numbers written as text are parsed strictly, never turned into 0."""

    def test_parse_number(self) -> None:
        self.assertEqual(0.125, parse_number("0,125"))
        self.assertEqual(1234.5, parse_number("1 234,5"))
        self.assertEqual(1234.5, parse_number("1\xa0234,5"))
        self.assertEqual(1234.5, parse_number("1.234,5"))
        self.assertEqual(1234.5, parse_number("1,234.5"))
        self.assertEqual(2.0, parse_number(2))
        for bad in ("нет", "#ЗНАЧ!", "#VALUE!", "NaN", "inf", "10-20", "1 2", "", None, float("nan"), math.inf, True):
            self.assertIsNone(parse_number(bad), bad)

    def test_text_weights_are_read_as_numbers_and_words_are_not_zero(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            database = root / "db.sqlite"
            make_database(database)
            price = write_price(
                root / "Price-list-CHINT_2026-06-01.xlsx",
                FULL_HEADER,
                [
                    ["A-1", "Автомат", "NXB", "шт", "", "0,125", "0,1", "1 234,5"],
                    ["A-2", "Щит", "NXB", "шт", "", "нет", "0,3", "#ЗНАЧ!"],
                ],
            )

            result = import_price_workbook(price, database, make_backup=False)

            records = db.fetch_products(database, ["A-1", "A-2"])
            self.assertEqual(0.125, records["A-1"].weight_kg)
            self.assertEqual(1234.5, records["A-1"].volume_m3)
            # Unreadable gross weight falls back to the net weight instead of becoming 0.
            self.assertEqual(0.3, records["A-2"].weight_kg)
            self.assertIsNone(records["A-2"].volume_m3)
            self.assertEqual(2, result.invalid_numbers)
            with closing(sqlite3.connect(database)) as conn:
                stored = conn.execute(
                    "SELECT gross_weight_unit FROM price_snapshot_items WHERE article = 'A-2'"
                ).fetchone()[0]
            self.assertEqual("нет", stored)

    def test_text_already_stored_in_database_is_read_strictly(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            database = Path(tmp) / "db.sqlite"
            make_database(database)
            with closing(sqlite3.connect(database)) as conn, conn:
                conn.execute("INSERT INTO products (article, name) VALUES ('A-1', 'Автомат')")
                conn.execute(
                    "INSERT INTO price_snapshot_items (snapshot_id, article, gross_weight_unit, unit_volume, min_shipment) "
                    "VALUES (1, 'A-1', '0,125', '1 234,5', 'нет')"
                )

            record = db.fetch_products(database, ["A-1"])["A-1"]

            self.assertEqual(0.125, record.weight_kg)
            self.assertEqual(1234.5, record.volume_m3)
            self.assertIsNone(record.min_shipment)


class EtimDimensionTests(unittest.TestCase):
    """ETI-01: dimensions follow the header order and units; ambiguity is not guessed."""

    def test_chain_order_comes_from_header(self) -> None:
        parse = etim_importer.parse_overall_dimensions
        self.assertEqual(
            {"length_mm": 100.0, "width_mm": 200.0, "height_mm": 300.0},
            parse("100x200x300", "Габаритные размеры ДхШхВ, мм"),
        )
        self.assertEqual(
            {"width_mm": 100.0, "height_mm": 200.0, "length_mm": 300.0},
            parse("10х20х30 см", "Габариты (Ш x В x Г)"),
        )
        self.assertEqual(
            {"width_mm": 12.5, "height_mm": 20.0, "length_mm": 30.0},
            parse("12,5x20x30", "Габаритные размеры (ширина x высота x глубина), мм"),
        )
        self.assertEqual({"length_mm": 1000.0, "width_mm": 500.0}, parse("1x0,5", "Габарит Д×Ш, м"))

    def test_unambiguous_legacy_chain_is_unchanged(self) -> None:
        self.assertEqual(
            {"width_mm": 100.0, "height_mm": 200.0, "length_mm": 300.0},
            etim_importer.parse_overall_dimensions("100x200x300", "Габаритные размеры"),
        )

    def test_ambiguous_chains_are_not_guessed(self) -> None:
        parse = etim_importer.parse_overall_dimensions
        self.assertEqual({}, parse("100x200", "Габаритные размеры"))
        self.assertEqual({}, parse("1x2x3x4", "Габаритные размеры"))
        self.assertEqual({}, parse("100x200x300", "Габаритные размеры ДхШ, мм"))
        self.assertEqual({}, parse("100-200x300x10", "Габаритные размеры"))
        self.assertEqual({}, parse("0x200x300", "Габаритные размеры"))
        self.assertEqual({}, parse("100x200x300 и 10x20x30", "Габаритные размеры"))

    def test_single_dimension_units_and_invalid_values(self) -> None:
        cell = etim_importer.parse_dimension_cell
        self.assertEqual(120.0, cell(12, "Габаритная длина, см", "length_mm"))
        self.assertEqual(500.0, cell("0,5 м", "Высота", "height_mm"))
        self.assertEqual(45.5, cell("45,5", "Ширина, мм", "width_mm"))
        for bad in ("-5", "10-20", "до 100", "нет", "nan"):
            self.assertIsNone(cell(bad, "Высота, мм", "height_mm"), bad)

    def test_explicit_columns_win_over_overall_chain(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "etim.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Артикул", "Габаритные размеры ДхШхВ, мм", "Длина, мм", "Ширина, мм", "Высота, мм"])
            sheet.append(["A-1", "100x200x300", 101, 201, 301])
            sheet.append(["Артикул", "Габаритная длина, см", "Габаритные размеры ДхШхВ, мм"])
            sheet.append(["A-2", 12, "100x200x300"])
            workbook.save(path)
            workbook.close()

            found = etim_importer.scan_etim_workbook(path, ["A-1", "A-2"])

            self.assertEqual({"length_mm": 101.0, "width_mm": 201.0, "height_mm": 301.0}, found["A-1"].dimensions)
            self.assertEqual({"length_mm": 120.0, "width_mm": 200.0, "height_mm": 300.0}, found["A-2"].dimensions)
            self.assertEqual("Габаритная длина, см", found["A-2"].dimension_sources["length_mm"])


class EtimDecisionTests(unittest.TestCase):
    """ETI-03: a stale or foreign conflict report must not overwrite newer data."""

    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)
        self.database = self.root / "db.sqlite"
        make_database(self.database)
        with closing(sqlite3.connect(self.database)) as conn, conn:
            for article in ("A-1", "A-2", "A-3"):
                conn.execute("INSERT INTO products (article) VALUES (?)", (article,))
            conn.execute("INSERT INTO dimension_imports (id, source_kind) VALUES (7, ?)", (etim_importer.SOURCE_KIND,))
            for article, length in (("A-1", 200), ("A-2", 250), ("A-3", 300), ("GHOST", 50)):
                conn.execute(
                    "INSERT INTO product_dimensions_raw (import_id, article, source_kind, dimension_kind, length_mm) "
                    "VALUES (7, ?, ?, 'length_mm', ?)",
                    (article, etim_importer.SOURCE_KIND, length),
                )
            conn.executemany(
                "INSERT INTO product_dimensions_resolved (article, length_mm) VALUES (?, ?)",
                [("A-1", 150), ("A-2", 100), ("A-3", 100)],
            )

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def write_report(self, rows: list[list[object]], import_id: int | None = 7, with_current: bool = True) -> Path:
        path = self.root / "report.xlsx"
        workbook = Workbook()
        summary = workbook.active
        summary.title = "Итог"
        summary.append(["Показатель", "Значение"])
        if import_id is not None:
            summary.append(["Импорт id", import_id])
        sheet = workbook.create_sheet("Конфликты")
        header = ["Решение", "Артикул", "Поле", "Сейчас в базе", "Пришло из ETIM"]
        if not with_current:
            header.remove("Сейчас в базе")
            rows = [row[:3] + row[4:] for row in rows]
        sheet.append(header)
        for row in rows:
            sheet.append(row)
        workbook.save(path)
        workbook.close()
        return path

    def length(self, article: str) -> float | None:
        with closing(sqlite3.connect(self.database)) as conn:
            row = conn.execute(
                "SELECT length_mm FROM product_dimensions_resolved WHERE article = ?", (article,)
            ).fetchone()
        return row[0] if row else None

    def test_stale_report_does_not_overwrite_newer_value(self) -> None:
        accept = etim_importer.ACCEPT_ETIM_DECISION
        report = self.write_report([[accept, "A-1", "Длина, мм", "100", "200"]])

        result = etim_importer.apply_etim_conflict_decisions(report, self.database)

        self.assertEqual(150, self.length("A-1"))
        self.assertEqual(0, result.accepted_rows)
        self.assertEqual(1, len(result.rejected))
        self.assertIn("база изменилась", result.rejected[0])

    def test_valid_rows_apply_and_bad_rows_are_rejected(self) -> None:
        accept = etim_importer.ACCEPT_ETIM_DECISION
        report = self.write_report(
            [
                [accept, "A-2", "Длина, мм", "100", "250"],          # valid
                [accept, "A-3", "Длина, мм", "100", "999"],          # edited proposal
                [accept, "GHOST", "Длина, мм", "", "50"],            # product not in this DB
                [accept, "A-3", "Длина, мм", "100", "nan"],          # not a number
                [etim_importer.KEEP_DB_DECISION, "A-1", "Длина, мм", "150", "200"],
            ]
        )

        result = etim_importer.apply_etim_conflict_decisions(report, self.database)

        self.assertEqual(250, self.length("A-2"))
        self.assertEqual(100, self.length("A-3"))
        self.assertIsNone(self.length("GHOST"))
        self.assertEqual(1, result.accepted_rows)
        self.assertEqual(1, result.skipped_rows)
        self.assertEqual(3, len(result.rejected))

    def test_contradicting_duplicates_are_rejected_and_repeat_is_noop(self) -> None:
        accept = etim_importer.ACCEPT_ETIM_DECISION
        report = self.write_report(
            [
                [accept, "A-2", "Длина, мм", "100", "250"],
                [accept, "A-3", "Длина, мм", "100", "300"],
                [accept, "A-3", "Длина, мм", "100", "250"],
            ]
        )

        first = etim_importer.apply_etim_conflict_decisions(report, self.database)
        second = etim_importer.apply_etim_conflict_decisions(report, self.database)

        self.assertEqual(100, self.length("A-3"))
        self.assertEqual(1, first.accepted_rows)
        self.assertIn("противоречивые", " ".join(first.rejected))
        self.assertEqual(0, second.accepted_rows)
        self.assertEqual(1, second.already_applied_rows)
        self.assertEqual(250, self.length("A-2"))

    def test_report_from_another_import_is_rejected(self) -> None:
        accept = etim_importer.ACCEPT_ETIM_DECISION
        report = self.write_report([[accept, "A-2", "Длина, мм", "100", "250"]], import_id=99)

        result = etim_importer.apply_etim_conflict_decisions(report, self.database)

        self.assertEqual(100, self.length("A-2"))
        self.assertEqual(1, len(result.rejected))

    def test_report_without_current_value_column_must_be_recreated(self) -> None:
        accept = etim_importer.ACCEPT_ETIM_DECISION
        report = self.write_report([[accept, "A-2", "Длина, мм", "100", "250"]], with_current=False)

        with self.assertRaisesRegex(ValueError, "пересоздайте отчет"):
            etim_importer.apply_etim_conflict_decisions(report, self.database)
        self.assertEqual(100, self.length("A-2"))


if __name__ == "__main__":
    unittest.main()
