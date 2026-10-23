import datetime as dt
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

from chint_etm_mdm.etm_categories import export_upload_csv
from chint_etm_mdm.etm_materials import (
    DESCRIPTIONS_TABLE,
    clean_text,
    read_descriptions,
    update_descriptions_table,
    write_csv,
)
from chint_etm_mdm.upload_names import cp1251_text, run_folder


class Cp1251Tests(unittest.TestCase):
    """ETM takes Win-1251 files: signs it lacks used to turn into "?" in the card."""

    def test_signs_are_spelled(self) -> None:
        self.assertEqual("2,5 мм2, ток <= 63 А, 18x80 мм, Ф6 мм", cp1251_text("2,5 мм², ток ≤ 63 А, 18×80 мм, Ø6 мм"))
        # As they come in the ЧЗ file.
        self.assertEqual("Звонок ND16-22F Ф22 мм", cp1251_text("Звонок ND16-22F Φ22 мм"))
        self.assertEqual("AC380~415В, 50 кВАр(CHINT), 1/2", cp1251_text("AC380～415В, 50 кВАр（CHINT）, ½"))
        self.assertEqual("обычный текст ±5 °C", cp1251_text("обычный текст ±5 °C"))

    def test_description_file(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "desc.csv"
            write_csv(path, ["Описание"], [[clean_text("Сечение до 2,5 мм², ток ≤ 63 А")]])
            text = path.read_bytes().decode("cp1251")
        self.assertNotIn("?", text)
        self.assertIn("2,5 мм2, ток <= 63 А", text)

    def test_cards_file(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Шаблон"
            sheet.append(["Артикул", "Сечение"])
            sheet.append(["586772", "1,5–6 мм²"])
            xlsx = Path(tmp) / "cards.xlsx"
            workbook.save(xlsx)
            csv_path = Path(tmp) / "cards.csv"
            self.assertEqual(1, export_upload_csv(xlsx, csv_path))
            self.assertEqual("Артикул;Сечение\r\n586772;1,5–6 мм2\r\n", csv_path.read_bytes().decode("cp1251"))


class RunFolderTests(unittest.TestCase):
    def test_two_runs_in_one_minute_get_two_folders(self) -> None:
        when = dt.datetime(2026, 10, 23, 9, 15)
        with tempfile.TemporaryDirectory() as tmp:
            first = run_folder(Path(tmp), "ЭТМ материалы", when)
            first.mkdir()
            second = run_folder(Path(tmp), "ЭТМ материалы", when)
        self.assertEqual("ЭТМ материалы 2026-10-23 09-15", first.name)
        self.assertEqual("ЭТМ материалы 2026-10-23 09-15 (2)", second.name)


class DescriptionsTableTests(unittest.TestCase):
    def test_counts_and_examples_survive_and_open_file_is_left_alone(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            table = Path(tmp) / DESCRIPTIONS_TABLE
            self.assertEqual(2, update_descriptions_table(table, {"1": "NXB", "2": "NXB", "3": "NH40"}, {"1": "Выключатель NXB"}))
            # A run about another series keeps what the table says about NXB.
            self.assertEqual(1, update_descriptions_table(table, {"4": "NB2"}, {"4": "Адаптер NB2"}))
            texts = read_descriptions(table)
            self.assertEqual("2", texts["NXB"]["Товаров"])
            self.assertEqual("Выключатель NXB", texts["NXB"]["Пример товара"])
            # Nothing new: the file is not touched, so Excel holding it is no problem.
            with mock.patch("chint_etm_mdm.etm_materials.write_descriptions", side_effect=PermissionError) as write:
                self.assertEqual(0, update_descriptions_table(table, {"1": "NXB"}, {}))
            write.assert_not_called()
            self.assertEqual(4, load_workbook(table).worksheets[0].max_row)



class PlainErrorTests(unittest.TestCase):
    def test_file_open_in_excel(self) -> None:
        from chint_etm_mdm.app import plain_error

        details = "Traceback (most recent call last):\n  ...\nPermissionError: [Errno 13] Permission denied: 'Сертификаты.xlsx'\n"
        self.assertTrue(plain_error(details).startswith("Файл занят другой программой"))
        self.assertIn("Сертификаты.xlsx", plain_error(details))
        self.assertEqual("ValueError: x", plain_error("ValueError: x"))


if __name__ == "__main__":
    unittest.main()
