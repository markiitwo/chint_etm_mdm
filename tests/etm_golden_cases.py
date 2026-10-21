"""Fixed inputs for the ETM golden test.

The expected outputs live in fixtures/etm_golden.json. To rebuild them after an
intentional change of the ETM output, run:

    PYTHONPATH=src python3 tests/etm_golden_cases.py
"""
from __future__ import annotations

import csv
import json
import sys
import tempfile
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from chint_etm_mdm.db import ProductRecord
from chint_etm_mdm import filler

FIXTURE_PATH = Path(__file__).with_name("fixtures") / "etm_golden.json"

RULES = {
    "version": 1,
    "class_rules": [
        {
            "class81_code": "50301005",
            "template_field": "Конфиг:Напряжение, В",
            "source_attributes": ["Напряжение лампы, В", "Напряжение лампы"],
            "confidence": "approved_class_rule",
        },
        {
            "class81_code": "50301005",
            "template_field": "Конфиг:Цвет свечения",
            "source_attributes": ["Цвет"],
            "confidence": "approved_class_rule",
        },
    ],
    "rejected_class_rules": [],
}

HEADERS = [
    "Артикул",
    "Расширенный артикул",
    "Код производителя",
    "Страна",
    "Название упаковки",
    "81 класс",
    "Название",
    "Полное название",
    "Краткое имя WMS",
    "Вес, кг",
    "Длина, м",
    "Ширина, м",
    "Высота, м",
    "Объем, м3",
    "Код ТН ВЭД",
    "Код ОКПД2",
    "Упак3 Название",
    "Упак3 Емкость",
    "Упак3 Вес, кг",
    "Упак3 Длина, м",
    "Упак3 Ширина, м",
    "Упак3 Высота, м",
    "Упак3 Объем, м3",
    "Штрих-код3",
    "Штрих-код3уп",
    "Отгрузка кратно УпЗавод",
    "Упак5 Емкость",
    "Упак5 Вес, кг",
    "Упак5 Длина, м",
    "Упак5 Ширина, м",
    "Упак5 Высота, м",
    "Упак5 Объем, м3",
    "Штрих-код5",
    "Штрих-код5уп",
    "Конфиг:Напряжение, В",
    "Конфиг:Цвет свечения",
    "Конфиг:Количество полюсов",
    "Описание",
    "Упак3 Что-то новое",
    "Упак5 Что-то новое",
]
NOT_YELLOW = {"Описание"}
EXTRA_HEADERS = ["Конфиг: Количество полюсов ", "Конфиг:", "", "Артикул"]

PRODUCTS = {
    "FULL-1": ProductRecord(
        article="FULL-1",
        name="Выключатель   автоматический NXB-63 1P 16А характеристика C 6кА очень длинное имя",
        full_name="Выключатель автоматический NXB-63 1P 16А C 6кА (CHINT)",
        class81_code="50101001",
        unit="шт",
        length_mm=81.5,
        width_mm=17.6,
        height_mm=74,
        weight_kg=0.105,
        volume_m3=0.000106,
        tnved_code="8536209008",
        okpd2_code="27.12.22.000",
        min_shipment=12,
        pack_multiple=120.0,
        package_min_weight_kg=1.32,
        package_min_volume_m3=0.0013,
        transport_weight_kg=13.5,
        transport_volume_m3=0.0145,
        transport_length_mm=390,
        transport_width_mm=280,
        transport_height_mm=133,
        gtin_inner="6943307012345",
        gtin_outer="16943307012342",
        attributes={"Количество полюсов": "1", "Цвет": "белый"},
    ),
    "LAMP-1": ProductRecord(
        article="LAMP-1",
        name="Лампа сигнальная ND16-22DS",
        class81_code="50301005",
        weight_kg=0.02,
        min_shipment=2.5,
        pack_multiple=None,
        gtin_outer="16943307099999",
        attributes={"Напряжение лампы": "230", "Цвет": "красный"},
    ),
    "EMPTY-1": ProductRecord(article="EMPTY-1"),
    "OVERRIDE-1": ProductRecord(
        article="OVERRIDE-1",
        full_name="Только полное имя",
        attributes={"Напряжение лампы, В": "24", "Цвет свечения": "зелёный"},
    ),
    "1001": ProductRecord(article="1001", name="Числовой артикул", gtin_inner="460000000001"),
}

# Template rows: article, prefilled cells by header, 81 класс override.
ROWS = [
    ("FULL-1", {}, ""),
    ("LAMP-1", {"Название": "Своё имя", "Вес, кг": 0}, ""),
    ("EMPTY-1", {}, ""),
    ("MISSING-1", {}, ""),
    ("OVERRIDE-1", {}, "50301005"),
    ("", {}, ""),
    (1001, {"Описание": "не трогать"}, ""),
]


def write_xlsx(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADERS)
    yellow = PatternFill(fill_type="solid", fgColor="FFFF00")
    for cell in sheet[1]:
        if cell.value not in NOT_YELLOW:
            cell.fill = yellow
    for article, prefilled, class_override in ROWS:
        row = [None] * len(HEADERS)
        row[0] = article
        if class_override:
            row[HEADERS.index("81 класс")] = class_override
        for header, value in prefilled.items():
            row[HEADERS.index(header)] = value
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def write_csv(path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh, delimiter=";")
        writer.writerow(HEADERS)
        for article, prefilled, class_override in ROWS:
            row = [""] * len(HEADERS)
            row[0] = str(article)
            if class_override:
                row[HEADERS.index("81 класс")] = class_override
            for header, value in prefilled.items():
                row[HEADERS.index(header)] = str(value)
            writer.writerow(row)


def read_report(path: Path) -> dict[str, list[list[object]]]:
    workbook = load_workbook(path)
    try:
        return {
            sheet.title: [list(row) for row in sheet.iter_rows(values_only=True)]
            for sheet in workbook.worksheets
        }
    finally:
        workbook.close()


def read_xlsx_output(path: Path) -> dict[str, object]:
    workbook = load_workbook(path)
    try:
        sheet = workbook.active
        values = [list(row) for row in sheet.iter_rows(values_only=True)]
        red = sorted(
            cell.coordinate
            for row in sheet.iter_rows(min_row=2)
            for cell in row
            if cell.fill.fill_type == "solid" and str(cell.fill.fgColor.rgb).upper() == "FFFFC7CE"
        )
    finally:
        workbook.close()
    return {"values": values, "red_cells": red}


def result_counts(result: filler.FillResult) -> dict[str, int]:
    return {
        "total_rows": result.total_rows,
        "found_articles": result.found_articles,
        "filled_cells": result.filled_cells,
        "suggested_cells": result.suggested_cells,
        "missing_cells": result.missing_cells,
        "missing_articles": result.missing_articles,
    }


def build_snapshot() -> dict[str, object]:
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        rules_path = root / "rules.json"
        rules_path.write_text(json.dumps(RULES, ensure_ascii=False), encoding="utf-8")
        manual_path = root / "manual.json"
        manual_path.write_text(
            json.dumps(
                {
                    "version": 1,
                    "values": [
                        {
                            "article": "EMPTY-1",
                            "field": "Код ТН ВЭД",
                            "value": "8536000000",
                        }
                    ],
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )

        direct = {}
        for article, product in PRODUCTS.items():
            direct[article] = {
                header: list(filler.product_value(product, header, rules_path))
                for header in HEADERS + EXTRA_HEADERS
            }
        headers_info = {
            header: {
                "reportable": filler.is_reportable_header(header),
                "missing_note": filler.missing_note(header),
            }
            for header in HEADERS + EXTRA_HEADERS
        }

        xlsx_template = root / "upload_goods.xlsx"
        csv_template = root / "upload_goods.csv"
        write_xlsx(xlsx_template)
        write_csv(csv_template)
        with patch("chint_etm_mdm.filler.fetch_products", side_effect=fetch_stub):
            xlsx_result = filler.fill_template(
                root / "db.sqlite", xlsx_template, root / "out_xlsx", rules_path, manual_path
            )
            csv_result = filler.fill_template(
                root / "db.sqlite", csv_template, root / "out_csv", rules_path, manual_path
            )
        with csv_result.output_path.open(newline="", encoding="utf-8-sig") as fh:
            csv_rows = list(csv.reader(fh, delimiter=";"))
        return {
            "product_value": direct,
            "headers": headers_info,
            "xlsx": {
                "counts": result_counts(xlsx_result),
                "output": read_xlsx_output(xlsx_result.output_path),
                "report": read_report(xlsx_result.report_path),
            },
            "csv": {
                "counts": result_counts(csv_result),
                "output": csv_rows,
                "report": read_report(csv_result.report_path),
            },
        }


def fetch_stub(db_path: Path, articles: list[str]) -> dict[str, ProductRecord]:
    return {article: PRODUCTS[article] for article in articles if article in PRODUCTS}


def normalize(snapshot: dict[str, object]) -> object:
    return json.loads(json.dumps(snapshot, ensure_ascii=False, default=str))


if __name__ == "__main__":
    data = normalize(build_snapshot())
    FIXTURE_PATH.parent.mkdir(parents=True, exist_ok=True)
    FIXTURE_PATH.write_text(json.dumps(data, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(f"Written {FIXTURE_PATH}", file=sys.stderr)
