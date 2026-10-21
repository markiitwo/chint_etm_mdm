import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from chint_etm_mdm.db import ProductRecord
from chint_etm_mdm.filler import fill_template
from chint_etm_mdm.manual_values import import_manual_completions
from chint_etm_mdm.value_matching import import_etm_reference

from test_value_matching import write_reference_xlsx

HEAD = [
    "ТОВАРЫ НЕДОСТАЮЩИЕ ХАРАКТЕРИСТИКИ (АВТО, #SIT-4542)",
    " ДАТА ОТЧЕТА: 21.10.2026 ВРЕМЯ ОТЧЕТА: 04:44",
    "Код производителя: 0000002019 CHINT      ",
    "Товары с учетом 27-го кл-ра: Все (где-нибудь не N)    ",
    "Товарный класс (3):      ",
    "Класс 81-го кл-ра:      ",
    "Класс81     ;Наименование класс 81     ;Наименование товара     ;Код товара;Артикул     ;Расширенный артикул     ;"
    "Код производителя;Производитель     ;наименование характеристики     ;Поле для заполнения характеристики     ;",
]
ROWS = [
    "'50329025;Катушки;Катушка AC220;1;'A-1;'A-1;2019;CHINT;Род тока; ",
    "'50329025;Катушки;Катушка AC220;1;'A-1;'A-1;2019;CHINT;Напряжение, В; ",
    "'50329025;Катушки;Катушка DC24;2;'A-2;'A-2;2019;CHINT;Род тока; ",
    "'50329025;Катушки;Катушка DC24;2;'A-2;'A-2;2019;CHINT;Номинальный ток (А); ",
    "'50329025;Катушки;Нет в базе;3;'X-9;'X-9;2019;CHINT;Род тока; ",
]
PRODUCTS = {
    "A-1": ProductRecord(article="A-1", attributes={"Род тока": "Переменный ток (AC)", "Напряжение, В": "230 В"}),
    "A-2": ProductRecord(article="A-2", attributes={"Род тока": "Трехфазный", "Номинальный ток (А)": "2,5 А"}),
}


class CharacteristicsReportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        rules_dir = self.root / "rules"
        rules_dir.mkdir()
        self.rules_path = rules_dir / "attribute_mappings.json"
        self.rules_path.write_text(json.dumps({"version": 1, "class_rules": []}), encoding="utf-8")
        write_reference_xlsx(self.root / "ref.xlsx")
        import_etm_reference(self.root / "ref.xlsx", rules_dir / "etm_reference.json")
        self.report = self.root / "rxreport_21-10-26_2019.csv"
        self.report.write_bytes("\n".join(HEAD + ROWS).encode("cp1251"))
        (self.root / "db.sqlite").write_bytes(b"")

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def fill(self):
        manual_path = self.root / "rules" / "manual_values.json"
        with patch("chint_etm_mdm.etm_characteristics.fetch_products", return_value=PRODUCTS):
            return fill_template(self.root / "db.sqlite", self.report, self.root / "out", self.rules_path, manual_path)

    def test_fills_only_the_value_column(self) -> None:
        result = self.fill()
        self.assertEqual((5, 2, 1, 3), (result.total_rows, result.found_articles, result.missing_articles, result.filled_cells))
        lines = result.output_path.read_bytes().decode("cp1251").split("\n")
        self.assertEqual(HEAD, lines[: len(HEAD)])
        self.assertEqual(
            [
                "'50329025;Катушки;Катушка AC220;1;'A-1;'A-1;2019;CHINT;Род тока;Переменный (AC)",
                "'50329025;Катушки;Катушка AC220;1;'A-1;'A-1;2019;CHINT;Напряжение, В;230",
                "'50329025;Катушки;Катушка DC24;2;'A-2;'A-2;2019;CHINT;Род тока; ",
                "'50329025;Катушки;Катушка DC24;2;'A-2;'A-2;2019;CHINT;Номинальный ток (А);2.5",
                ROWS[4],
            ],
            lines[len(HEAD) :],
        )

    def test_report_lists_values_to_confirm_once(self) -> None:
        result = self.fill()
        workbook = load_workbook(result.report_path)
        try:
            rows = list(workbook["Сопоставления"].iter_rows(values_only=True))
        finally:
            workbook.close()
        self.assertEqual(
            [
                ("Конфиг:Род тока", "Переменный ток (AC)", "Переменный (AC)", 1, None, "Переменный (AC)"),
                ("Конфиг:Род тока", "Трехфазный", None, 1, None, None),
            ],
            rows[1:],
        )

    def test_confirmations_and_filled_report_are_remembered(self) -> None:
        result = self.fill()
        workbook = load_workbook(result.report_path)
        workbook["Сопоставления"]["F3"] = "Переменный (AC)"
        workbook.save(result.report_path)
        workbook.close()
        manual_path = self.root / "rules" / "manual_values.json"
        confirmed = import_manual_completions(self.root / "db.sqlite", result.report_path, manual_path, self.root / "out", self.rules_path)
        self.assertEqual(2, confirmed.learned_matches)

        second = self.fill()
        lines = second.output_path.read_bytes().decode("cp1251").split("\n")
        self.assertTrue(lines[len(HEAD) + 2].endswith(";Род тока;Переменный (AC)"))

        # A value typed straight into the ETM report is kept for that product.
        filled = self.root / "rxreport_filled.csv"
        filled.write_bytes("\n".join(HEAD + [ROWS[1].replace("; ", ";220")]).encode("cp1251"))
        with patch("chint_etm_mdm.manual_values.fetch_products", return_value=PRODUCTS):
            imported = import_manual_completions(self.root / "db.sqlite", filled, manual_path, self.root / "out", self.rules_path)
        self.assertEqual(1, imported.imported_values)
        third = self.fill()
        self.assertIn(";Напряжение, В;220", third.output_path.read_bytes().decode("cp1251"))


    def test_characteristic_name_mapping_is_learned(self) -> None:
        products = dict(PRODUCTS)
        products["A-3"] = ProductRecord(
            article="A-3", attributes={"Первичный ток, А": "100 А", "Номинальное напряжение изоляции, В": "690", "Номинальный ток, А": "16 А", "Первичное напряжение 1, В": "380 В"}
        )
        self.report.write_bytes(
            "\n".join(
                HEAD
                + [
                    "'50329025;Трансформаторы;ОСМ1;3;'A-3;'A-3;2019;CHINT;Номинальный первичный ток (А); ",
                    "'50329025;Трансформаторы;ОСМ1;3;'A-3;'A-3;2019;CHINT;Номинальный ток (А); ",
                    "'50329025;Трансформаторы;ОСМ1;3;'A-3;'A-3;2019;CHINT;Номинальное напряжение первичной обмотки, В; ",
                ]
            ).encode("cp1251")
        )
        manual_path = self.root / "rules" / "manual_values.json"
        with patch("chint_etm_mdm.etm_characteristics.fetch_products", return_value=products):
            first = fill_template(self.root / "db.sqlite", self.report, self.root / "out", self.rules_path, manual_path)
        # Same name up to the unit notation: filled, but flagged for a check; the shipped
        # "первичная обмотка" pair fills directly.
        self.assertEqual((2, 1), (first.filled_cells, first.suggested_cells))
        workbook = load_workbook(first.report_path)
        names = workbook["Названия"]
        self.assertEqual(("Номинальный первичный ток (А)", 1, "Первичный ток, А", 1, "100 А", None), tuple(cell.value for cell in names[2]))
        names["F2"] = "да"
        workbook.save(first.report_path)
        workbook.close()
        imported = import_manual_completions(self.root / "db.sqlite", first.report_path, manual_path, self.root / "out", self.rules_path)
        self.assertEqual(1, imported.learned_names)
        with patch("chint_etm_mdm.etm_characteristics.fetch_products", return_value=products):
            second = fill_template(self.root / "db.sqlite", self.report, self.root / "out", self.rules_path, manual_path)
        self.assertEqual(3, second.filled_cells)
        text = second.output_path.read_bytes().decode("cp1251")
        self.assertIn(";Номинальный первичный ток (А);100\n", text)
        self.assertIn(";Номинальный ток (А);16\n", text)
        self.assertTrue(text.endswith(";Номинальное напряжение первичной обмотки, В;380"))


if __name__ == "__main__":
    unittest.main()
