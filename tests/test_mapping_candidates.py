"""MAP-03: every candidate reaches the user, also for a field with an approved source."""

from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from chint_etm_mdm.analyzer import analyze_config_field, write_mapping_report
from chint_etm_mdm.db import ProductRecord
from chint_etm_mdm.mapping_rules import add_approved_class_rule


class CandidatesTests(unittest.TestCase):
    def test_all_candidates_and_the_uncovered_part(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            rules = Path(tmp) / "attribute_mappings.json"
            rules.write_text(json.dumps({"version": 1, "class_rules": []}), encoding="utf-8")
            add_approved_class_rule(rules, "C1", "Конфиг:Номинальный ток", "Ток, А")
            products = [
                ProductRecord(article="1", attributes={"Ток, А": "16"}),
                ProductRecord(article="2", attributes={"Номинальный ток расцепителя": "25", "Номинальный ток (In)": "25 А"}),
            ]
            coverage, examples, issues = analyze_config_field("C1", "Конфиг:Номинальный ток", products, rules)
            self.assertEqual(1, coverage.approved_rule_count)
            self.assertEqual(1, coverage.candidate_count)
            self.assertEqual(
                {"Номинальный ток расцепителя", "Номинальный ток (In)"}, set(coverage.candidate_sources)
            )

            report = Path(tmp) / "mapping_review.xlsx"
            write_mapping_report(report, [coverage], examples, issues)
            rows = list(load_workbook(report)["Правила"].iter_rows(min_row=2, values_only=True))
            self.assertEqual("уже утверждено", rows[0][6])
            self.assertEqual(
                {"Номинальный ток расцепителя", "Номинальный ток (In)"}, {row[2] for row in rows[1:]}
            )
