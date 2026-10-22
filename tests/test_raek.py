import datetime as dt
import json
import sqlite3
import tempfile
import unittest
import zipfile
from contextlib import closing
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

from chint_etm_mdm.raek import (
    CHECK,
    EMPTY,
    FILLED,
    Cell,
    Sources,
    certificate,
    fill_raek_zip,
    is_raek_zip,
    latest_raek_zip,
    mandatory_certification,
    remember_name_answers,
    result_lines,
)
from test_import_data_integrity import make_database

BASE = ["RAEC_ID", "Код поставщика", "Название от поставщика", "Бренд", "Статус поставщика", "ABC Приоритет"]
RELAY = ("3244400", "631929", "Промежуточное реле NJDC-17/4ZS(MD)L 4 конт. 6А 24В AC (R)")
DRIVE = ("3249224", "616173", "Преобразователь частоты NVF7-3.0T-S6, 690В")
UNKNOWN = ("3249999", "999999", "Чего-то нет в базе")


def plain_file(column, products):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(BASE + (column if isinstance(column, list) else [column]))
    for raec_id, article, name in products:
        sheet.append([raec_id, article, name, "ЧИНТ (CHINT)", "Z", "A"])
    return workbook


def features_file():
    workbook = Workbook()
    relay = workbook.active
    relay.title = "EC001437"
    relay.append([None] * 8 + ["EV000460 - AC (перемен.)\nEV000461 - DC (постоян.)", None, None, "EV006405 - IP20\nEV006415 - IP40", None, None])
    relay.append(
        [None] * 6
        + [
            "<обязательное>\nR - Номин. напряжение питания цепи управления Us AC 50 Гц, В",
            "<обязательное>\nR - Номин. напряжение питания цепи управления Us постоян. тока DC, В",
            "<обязательное>\nA - Тип напряжения управления",
            "<обязательное>\nN - Количество переключающих (перекидных) контактов",
            "<обязательное>\nL - В комплекте с розеткой (цоколем)",
            "<обязательное>\nA - Степень защиты (IP)",
            "<обязательное>\nN - Ширина, мм",
            "<обязательное>\nN - Номин. ток, А",
        ]
    )
    relay.append(BASE + ["EF003978", "EF003980", "EF008242", "EF008257", "EF003359", "EF005474", "EF000008", "EF000001"])
    relay.append(list(RELAY) + ["ЧИНТ (CHINT)", "Z", "A"])
    drive = workbook.create_sheet("EC001857")
    drive.append([None] * 7 + ["EV000599 - 1\nEV000069 - 3\nEV000154 - Прочее", None, None, "EV008013 - 50 Гц\nEV008017 - 60 Гц\nEV008015 - 50/60 Гц"])
    drive.append(
        [None] * 6
        + [
            "<обязательное>\nR - Сетевое напряжение, В",
            "<обязательное>\nA - Количество вход. фаз",
            "<обязательное>\nN - Макс. частота на выходе, Гц",
            "<обязательное>\nN - Высота, мм",
            "<обязательное>\nA - Частота сети",
        ]
    )
    drive.append(BASE + ["EF006780", "EF007117", "EF007350", "EF000040", "EF006783"])
    drive.append(list(DRIVE) + ["ЧИНТ (CHINT)", "Z", "A"])
    return workbook


def make_zip(path: Path) -> None:
    files = {
        "ChINT (CHINT)_COUNTRY_ID.xlsx": plain_file("Страна происхождения", [RELAY, DRIVE]),
        "ChINT (CHINT)_THVED.xlsx": plain_file("Код ТН ВЭД", [DRIVE, UNKNOWN]),
        "ChINT (CHINT)_EAN.xlsx": plain_file("Штрих-код EAN", [DRIVE, RELAY]),
        "ChINT (CHINT)_OKPD2_ID.xlsx": plain_file(["Код ОКПД2", "Название ОКПД2"], [DRIVE]),
        "ChINT (CHINT)_SERIES_ID.xlsx": plain_file("Серия", [DRIVE]),
        "ChINT (CHINT)_IMAGE.xlsx": plain_file("Изображение", [DRIVE]),
        "ChINT (CHINT)_certificate.xlsx": plain_file(
            ["Тип сертификата", "Номер сертификата", "Сертификат от", "Сертификат до", "Файл", "Ссылка на реестр"], [DRIVE, RELAY]
        ),
        "ChINT (CHINT)_required_features.xlsx": features_file(),
        "ChINT (CHINT)_MANDATORY_CERTIFICATION.xlsx": plain_file("Подлежит обязательному подтверждению соответствия", [DRIVE, RELAY]),
    }
    with zipfile.ZipFile(path, "w") as archive:
        for name, workbook in files.items():
            buffer = BytesIO()
            workbook.save(buffer)
            archive.writestr(name, buffer.getvalue())
        archive.writestr("readme.txt", "оставить как есть")


def make_codes(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Price-list"
    sheet.append(["PM"])
    sheet.append(["Part number", "ОКПД 2 \n(9 зн.)", "Серия", "ШК / GTIN", "Сертификат 1", "Ссылка на реестр 1", "Tariff code"])
    sheet.append([616173, "26.51.70.190", "NVF7", "6941498900011", "ЕАЭС N RU Д-CN.РА06.В.41573/24", "https://pub.fsa.gov.ru/x", 9032890000])
    workbook.save(path)


def cells(workbook_bytes: bytes, sheet: str | None = None) -> list[tuple]:
    workbook = load_workbook(BytesIO(workbook_bytes))
    target = workbook[sheet] if sheet else workbook.active
    return list(target.iter_rows(values_only=True))


class RaekTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.db = self.root / "mdm.sqlite"
        make_database(self.db)
        with closing(sqlite3.connect(self.db)) as conn, conn:
            conn.executemany(
                "INSERT INTO products(article, name, series_name) VALUES (?, ?, ?)",
                [(RELAY[1], RELAY[2], "NJDC-17"), (DRIVE[1], DRIVE[2], "NVF7")],
            )
            conn.executemany(
                "INSERT INTO product_attribute_values(article, attribute_name, value_raw, confidence) VALUES (?, ?, ?, 1)",
                [
                    (RELAY[1], "Номинальное напряжение питания цепи управления Us перемен. тока АС при 50 Гц, В", "24"),
                    (RELAY[1], "Род тока включения", "Переменный ток (AC)"),
                    (RELAY[1], "Количество переключающих контактов", "4"),
                    (RELAY[1], "В комплекте с розеткой", "Нет"),
                    (RELAY[1], "Степень защиты (IP)", "IP40"),
                    (RELAY[1], "Габаритные и установочные размеры (ШхВхГ), мм", "27,8×21,8×44"),
                    (RELAY[1], "Номинальный ток, А", "6"),
                    (RELAY[1], "Номер сертификата", "ЕАЭС N RU Д-CN.PA03.B.02078/25"),
                    (RELAY[1], "Дата начала", "2025-03-21 00:00:00"),
                    (RELAY[1], "Дата окончания", "2030-03-20 00:00:00"),
                    (DRIVE[1], "Напряжение сети, В", "525-690"),
                    (DRIVE[1], "Количество входных фаз", "3"),
                    (DRIVE[1], "Диапазон входной частоты, Гц", "47-63"),
                    (DRIVE[1], "Диапазон выходной частоты, Гц", "0-500"),
                    (DRIVE[1], "Габаритные и установочные размеры (ШхВхГ), мм", "270х466х268.1"),
                    (DRIVE[1], "Номер сертификата", "ЕАЭС N RU Д-CN.РА06.В.41573/24"),
                    (DRIVE[1], "Дата начала", "2024-07-25 00:00:00"),
                    (DRIVE[1], "Дата окончания", "2029-07-25 00:00:00"),
                ],
            )
            conn.execute(
                "INSERT INTO material_data_items(article, gtin_inner) VALUES (?, '16941498941592')", (RELAY[1],)
            )
        self.zip = self.root / "CHINT.zip"
        make_zip(self.zip)
        self.codes = self.root / "Список товаров под ЧЗ_19.10.2026.xlsx"
        make_codes(self.codes)
        self.rules = self.root / "rules" / "attribute_mappings.json"
        self.rules.parent.mkdir()

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def fill(self):
        return fill_raek_zip(self.db, self.zip, self.root / "out", self.codes, self.rules)

    def member(self, result, name: str) -> bytes:
        with zipfile.ZipFile(result.output_path) as archive:
            return archive.read(name)

    def test_plain_columns(self) -> None:
        result = self.fill()
        self.assertEqual(f"CHINT_заполнено_{dt.date.today().isoformat()}.zip", result.output_path.name)
        with zipfile.ZipFile(result.output_path) as archive:
            self.assertEqual(
                [
                    "ChINT (CHINT)_COUNTRY_ID.xlsx", "ChINT (CHINT)_THVED.xlsx", "ChINT (CHINT)_EAN.xlsx",
                    "ChINT (CHINT)_OKPD2_ID.xlsx", "ChINT (CHINT)_SERIES_ID.xlsx", "ChINT (CHINT)_IMAGE.xlsx",
                    "ChINT (CHINT)_certificate.xlsx", "ChINT (CHINT)_required_features.xlsx",
                    "ChINT (CHINT)_MANDATORY_CERTIFICATION.xlsx", "readme.txt",
                ],
                archive.namelist(),
            )
            self.assertEqual("оставить как есть", archive.read("readme.txt").decode())
        self.assertEqual(["Китай", "Китай"], [row[6] for row in cells(self.member(result, "ChINT (CHINT)_COUNTRY_ID.xlsx"))[1:]])
        self.assertEqual(["9032890000", None], [row[6] for row in cells(self.member(result, "ChINT (CHINT)_THVED.xlsx"))[1:]])
        ean = cells(self.member(result, "ChINT (CHINT)_EAN.xlsx"))
        self.assertEqual("6941498900011", ean[1][6])
        self.assertIsNone(ean[2][6])  # a pack barcode is a different code: never used for the product
        self.assertEqual(("26.51.70.190", None), cells(self.member(result, "ChINT (CHINT)_OKPD2_ID.xlsx"))[1][6:8])
        self.assertEqual("NVF7", cells(self.member(result, "ChINT (CHINT)_SERIES_ID.xlsx"))[1][6])
        cert = cells(self.member(result, "ChINT (CHINT)_certificate.xlsx"))
        self.assertEqual(
            ("Декларация ЕАС", "ЕАЭС N RU Д-CN.РА06.В.41573/24", dt.datetime(2024, 7, 25), dt.datetime(2029, 7, 25), None, "https://pub.fsa.gov.ru/x"),
            cert[1][6:12],
        )
        self.assertEqual(("Декларация ЕАС", "ЕАЭС N RU Д-CN.PA03.B.02078/25", dt.datetime(2025, 3, 21)), cert[2][6:9])
        self.assertIsNone(cert[2][11])  # no registry link without the ЧЗ file

        # From the ЧЗ file: the drive has a declaration ЕАЭС; the relay is not in the file, so no guess.
        self.assertEqual(["Да", None], [row[6] for row in cells(self.member(result, "ChINT (CHINT)_MANDATORY_CERTIFICATION.xlsx"))[1:]])

        report = load_workbook(result.report_path)
        empty = [row for row in report["Нечем заполнить"].iter_rows(min_row=2, values_only=True)]
        self.assertIn(
            ("ChINT (CHINT)_MANDATORY_CERTIFICATION.xlsx", "631929"),
            [(row[0], row[3]) for row in empty if "файле ЧЗ" in row[7]],
        )
        self.assertIn(("ChINT (CHINT)_THVED.xlsx", "999999", "Артикула нет в базе"), [(row[0], row[3], row[7]) for row in empty])
        self.assertTrue(any(row[5] == "Изображение" and "ensmas" in row[7] for row in empty))
        checks = [row for row in report["Проверить"].iter_rows(min_row=2, values_only=True)]
        self.assertTrue(any(row[5] == "Штрих-код EAN" and row[3] == "631929" for row in empty))
        self.assertFalse(any(row[5] == "Штрих-код EAN" for row in checks))

        # The summary speaks of products and their data, not of cells.
        summary = list(report["Итог"].iter_rows(values_only=True))
        self.assertEqual(("Товаров в архиве", 3), summary[0][:2])
        self.assertIn(("Штрих-код EAN", 1, "Нет штрих-кода товара в файле ЧЗ"), [row[:3] for row in summary])
        lines = "\n".join(result_lines(result))
        self.assertIn("Товаров в архиве: 3.", lines)
        self.assertNotIn("ячеек", lines)

    def test_features(self) -> None:
        result = self.fill()
        relay = cells(self.member(result, "ChINT (CHINT)_required_features.xlsx"), "EC001437")
        self.assertEqual(("24...24", "-...-", "AC (перемен.)", "4", "false", "IP40", "27,8", "6"), relay[3][6:14])
        drive = cells(self.member(result, "ChINT (CHINT)_required_features.xlsx"), "EC001857")
        self.assertEqual(("525...690", "3", "500", "466", "50/60 Гц"), drive[3][6:11])
        report = load_workbook(result.report_path)
        checks = [row for row in report["Проверить"].iter_rows(min_row=2, values_only=True)]
        self.assertTrue(any(row[5].startswith("EF007350") and "верхний предел" in row[7] for row in checks))
        lines = "\n".join(result_lines(result))
        self.assertIn("Архив для РАЭК: CHINT_заполнено_", lines)
        self.assertIn("Список товаров под ЧЗ_19.10.2026.xlsx", lines)

    def test_filled_cells_stay(self) -> None:
        with zipfile.ZipFile(self.zip) as archive:
            data = archive.read("ChINT (CHINT)_required_features.xlsx")
        workbook = load_workbook(BytesIO(data))
        workbook["EC001437"]["G4"] = "220...220"
        buffer = BytesIO()
        workbook.save(buffer)
        with zipfile.ZipFile(self.zip, "w") as archive:
            archive.writestr("ChINT (CHINT)_required_features.xlsx", buffer.getvalue())
        result = self.fill()
        relay = cells(self.member(result, "ChINT (CHINT)_required_features.xlsx"), "EC001437")
        self.assertEqual("220...220", relay[3][6])
        self.assertEqual(1, result.kept)

    def test_similar_names_are_checked_and_answers_remembered(self) -> None:
        with closing(sqlite3.connect(self.db)) as conn, conn:
            conn.execute(
                "INSERT INTO product_attribute_values(article, attribute_name, value_raw, confidence) VALUES (?, ?, ?, 1)",
                (DRIVE[1], "Высота корпуса, мм", "470"),
            )
            conn.execute("DELETE FROM product_attribute_values WHERE attribute_name LIKE 'Габарит%' AND article = ?", (DRIVE[1],))
        result = self.fill()
        drive = cells(self.member(result, "ChINT (CHINT)_required_features.xlsx"), "EC001857")
        self.assertEqual("470", drive[3][9])
        self.assertEqual([["EC001857", "EF000040", "Высота, мм", "Высота корпуса, мм", 1, "470", ""]], result.names)

        report = load_workbook(result.report_path)
        report["Названия"]["G2"] = "нет"
        report.save(result.report_path)
        self.assertEqual((0, 1), remember_name_answers(result.report_path, self.rules))
        learned = json.loads((self.rules.parent / "raek_learned.json").read_text(encoding="utf-8"))
        self.assertEqual({"EF000040": ["Высота корпуса, мм"]}, learned["name_rejects"])
        again = self.fill()
        self.assertIsNone(cells(self.member(again, "ChINT (CHINT)_required_features.xlsx"), "EC001857")[3][9])

    def test_finds_the_newest_raek_archive(self) -> None:
        other = self.root / "photos.zip"
        with zipfile.ZipFile(other, "w") as archive:
            archive.writestr("a.jpg", b"")
        self.assertTrue(is_raek_zip(self.zip))
        self.assertFalse(is_raek_zip(other))
        self.assertEqual(self.zip, latest_raek_zip([self.root]))

    def test_app_worker(self) -> None:
        from chint_etm_mdm.app import RaekWorker

        results = []
        worker = RaekWorker(self.db, self.zip, self.root / "out", self.codes, self.rules)
        worker.done.connect(results.append)
        worker.failed.connect(self.fail)
        worker.run()
        self.assertTrue(results[0].output_path.exists())


if __name__ == "__main__":
    unittest.main()


    def test_site_fills_links_and_pictures(self) -> None:
        """Links, pictures and the certificate number from ensmas.ru (a fake site here)."""
        from test_site_info import PAGE_URL, PRODUCT, SEARCH, FakeSite

        site = FakeSite({"https://ensmas.ru/catalog/?q=616173": SEARCH, PAGE_URL: PRODUCT, "https://ensmas.ru/catalog/?q=631929": "<html></html>"})
        site.pages["https://ensmas.ru/catalog/?q=999999"] = "<html></html>"
        lines: list[str] = []
        cache = self.root / "work" / "cache" / "ensmas_site.json"
        result = fill_raek_zip(self.db, self.zip, self.root / "out", None, self.rules, cache, True, site, lines.append)
        image = cells(self.member(result, "ChINT (CHINT)_IMAGE.xlsx"))
        self.assertTrue(image[1][6].startswith("https://ensmas.ru/upload/iblock/fd6/"))
        # No ЧЗ file this time: the drive's declaration on the site still says certification is mandatory, flagged.
        mandatory = cells(self.member(result, "ChINT (CHINT)_MANDATORY_CERTIFICATION.xlsx"))
        self.assertEqual(["Да", None], [row[6] for row in mandatory[1:]])
        self.assertTrue(cache.exists())
        self.assertTrue(any("Сайт ensmas.ru: спросил 3, найдено 1, не найдено 2" in line for line in result_lines(result)))
        self.assertTrue(any("616173 найден" in line for line in lines))

        report = load_workbook(result.report_path)
        checks = [(row[3], row[5]) for row in report["Проверить"].iter_rows(min_row=2, values_only=True)]
        self.assertIn(("616173", "Подлежит обязательному подтверждению соответствия"), checks)

        # The next run takes the answers from memory and does not ask the site again.
        asked = len(site.asked)
        again = fill_raek_zip(self.db, self.zip, self.root / "out", None, self.rules, cache, True, site)
        self.assertEqual(asked, len(site.asked))
        self.assertTrue(cells(self.member(again, "ChINT (CHINT)_IMAGE.xlsx"))[1][6])


class MandatoryCertificationTests(unittest.TestCase):
    def sources(self, document: str | None) -> Sources:
        chz = {("100", "series"): "NKB1"}
        if document is not None:
            chz[("100", "cert_number")] = document
        return Sources(products={}, chz=chz, series={}, coded={}, settings={}, learned={})

    def test_from_chz_certificate_column(self) -> None:
        cases = {
            "ЕАЭС RU С-CN.НВ12.В.00413/23": Cell(FILLED, "Да"),
            "EAЭС N RU Д-CN.РА06.В.41573/24": Cell(FILLED, "Да"),  # Latin letters mixed in
            "не требуется ": Cell(FILLED, "Нет"),
        }
        for document, expected in cases.items():
            self.assertEqual(expected, mandatory_certification("100", self.sources(document)), document)
        letter = mandatory_certification("100", self.sources("ОР-25/06-0035"))
        self.assertEqual((CHECK, "Нет"), (letter.status, letter.value))
        self.assertIn("ОР-25/06-0035", letter.note)

    def test_unknown_stays_empty(self) -> None:
        self.assertEqual(EMPTY, mandatory_certification("100", self.sources(None)).status)
        self.assertEqual(EMPTY, mandatory_certification("200", self.sources("ЕАЭС RU С-CN.1")).status)
        no_file = Sources(products={}, chz={}, series={}, coded={}, settings={}, learned={})
        self.assertEqual(EMPTY, mandatory_certification("100", no_file).status)

    def test_refusal_letter_is_not_a_certificate(self) -> None:
        from chint_etm_mdm.db import ProductRecord

        product = ProductRecord("100", attributes={})
        cert = certificate("100", product, self.sources("30.04.2026-21ССП"))
        number = [cell for key, cell in cert.items() if "номер" in key][0]
        self.assertEqual(EMPTY, number.status)
        self.assertIn("30.04.2026-21ССП", number.note)
