import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from chint_etm_mdm.db import ProductRecord
from chint_etm_mdm.filler import fill_xlsx_template, product_value
from chint_etm_mdm.manual_values import import_manual_completions
from chint_etm_mdm.value_matching import Characteristic, import_etm_reference, match_value

CURRENT = Characteristic("Род тока", values=["Переменный (AC)", "Переменный/Постоянный (AC/DC)", "Постоянный (DC)"])
VOLTAGE = Characteristic("Напряжение, В", values=["12", "24", "220", "230", "220/380"])
LOGICAL = Characteristic("Наличие индикации", values=["Да", "Нет"])
POLES = Characteristic("Количество силовых полюсов", values=["1P", "2P", "3P", "3P+N", "4P"])


class MatchValueTests(unittest.TestCase):
    def test_exact_and_normalized(self) -> None:
        self.assertEqual(("3P", "exact"), (match_value("3P", POLES).value, match_value("3P", POLES).kind))
        result = match_value(" 3p + n ", POLES)
        self.assertEqual(("3P+N", "normalized"), (result.value, result.kind))

    def test_number_with_unit(self) -> None:
        result = match_value("230 В", VOLTAGE)
        self.assertEqual(("230", "number"), (result.value, result.kind))
        self.assertTrue(result.confident)

    def test_logical_words(self) -> None:
        self.assertEqual("Да", match_value("есть", LOGICAL).value)
        self.assertEqual("Нет", match_value("отсутствует", LOGICAL).value)

    def test_same_words_in_other_order_needs_check(self) -> None:
        result = match_value("AC (перемен.)", CURRENT)
        self.assertEqual(("Переменный (AC)", "words"), (result.value, result.kind))
        self.assertFalse(result.confident)

    def test_ambiguous_value_is_not_guessed(self) -> None:
        result = match_value("переменный 50 Гц", CURRENT)
        self.assertEqual(("", "none"), (result.value, result.kind))
        self.assertIn("Переменный (AC)", result.candidates)
        self.assertEqual("none", match_value("1P+N", POLES).kind)

    def test_closest_fit_by_words(self) -> None:
        self.assertEqual("Переменный (AC)", match_value("Переменный ток (AC)", CURRENT).value)
        self.assertEqual("Переменный/Постоянный (AC/DC)", match_value("AC/DC", CURRENT).value)
        self.assertEqual("words", match_value("Постоянный ток (DC)", CURRENT).kind)

    def test_learned_match_wins(self) -> None:
        result = match_value("~230", VOLTAGE, {"~230": "230"})
        self.assertEqual(("230", "learned"), (result.value, result.kind))

    def test_open_characteristic(self) -> None:
        self.assertEqual("open", match_value("что угодно", Characteristic("Серия")).kind)


def write_reference_xlsx(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Справочник"
    sheet.append(["Код характеристики", "Название характеристики", "Формат хар-ки", "Возможные значения", "Комментарий", "Дата актуализации значений"])
    for value in CURRENT.values:
        sheet.append(["68", "Род тока", "Текстовая", value, None, None])
    for value in [12, 24, 220, 230]:
        sheet.append(["132", "Напряжение, В", "Текстовая", value, None, None])
    sheet.append(["132", "Напряжение, В", "Текстовая", "220/380", None, None])
    sheet.append(["900", "Серия", "Текстовая", None, None, None])
    workbook.save(path)


class ReferenceInFillTests(unittest.TestCase):
    def test_fill_uses_allowed_values_and_learns_manual_choice(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            rules_dir = root / "rules"
            rules_path = rules_dir / "attribute_mappings.json"
            rules_dir.mkdir()
            rules_path.write_text(json.dumps({"version": 1, "class_rules": []}), encoding="utf-8")
            reference_xlsx = root / "spravochnik.xlsx"
            write_reference_xlsx(reference_xlsx)
            self.assertEqual((3, 8), import_etm_reference(reference_xlsx, rules_dir / "etm_reference.json"))

            products = {
                "A-1": ProductRecord(article="A-1", attributes={"Род тока": "Переменный 50 Гц", "Напряжение, В": "230 V", "Серия": "NXB"}),
                "A-2": ProductRecord(article="A-2", attributes={"Род тока": "AC (перемен.)", "Напряжение, В": "230-240"}),
            }
            self.assertEqual(("230", "filled", "product_attribute_values.attribute_name"), product_value(products["A-1"], "Конфиг:Напряжение, В", rules_path))
            self.assertEqual(("NXB", "filled", "product_attribute_values.attribute_name"), product_value(products["A-1"], "Конфиг:Серия", rules_path))
            self.assertEqual("value_not_allowed", product_value(products["A-1"], "Конфиг:Род тока", rules_path)[1])
            self.assertEqual(("Переменный (AC)", "suggested"), product_value(products["A-2"], "Конфиг:Род тока", rules_path)[:2])

            template = root / "upload_goods.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Артикул", "Конфиг:Род тока", "Конфиг:Напряжение, В"])
            for cell in sheet[1]:
                cell.fill = PatternFill(fill_type="solid", fgColor="FFFF00")
            sheet.append(["A-1", None, None])
            sheet.append(["A-2", None, None])
            workbook.save(template)

            with patch("chint_etm_mdm.filler.fetch_products", return_value=products):
                result = fill_xlsx_template(root / "db.sqlite", template, root / "out", rules_path)
            self.assertEqual((2, 1, 2), (result.filled_cells, result.suggested_cells, result.missing_cells))
            output = load_workbook(result.output_path)
            try:
                self.assertEqual([None, "230", "Переменный (AC)", None], [output.active["B2"].value, output.active["C2"].value, output.active["B3"].value, output.active["C3"].value])
            finally:
                output.close()
            report = load_workbook(result.report_path)
            try:
                notes = [row[6] for row in report["Детали"].iter_rows(min_row=2, values_only=True) if row[3] == "Значения нет в списке партнёра"]
            finally:
                report.close()
            self.assertEqual(2, len(notes))
            self.assertIn("«Переменный 50 Гц»", notes[0])
            self.assertIn("Переменный (AC)", notes[0])

            # The user picks the allowed values by hand and keeps the suggestion; the program remembers all three.
            filled = load_workbook(result.output_path)
            filled.active["B2"] = "Переменный (AC)"
            filled.active["C3"] = "230"
            filled.save(root / "filled.xlsx")
            filled.close()
            (root / "db.sqlite").write_bytes(b"")
            with patch("chint_etm_mdm.manual_values.fetch_products", return_value=products):
                manual = import_manual_completions(root / "db.sqlite", root / "filled.xlsx", rules_dir / "manual_values.json", root / "out", rules_path)
            self.assertEqual(3, manual.learned_matches)
            other = ProductRecord(article="B-1", attributes={"Род тока": "переменный  50 гц", "Напряжение, В": "230-240"})
            self.assertEqual(("Переменный (AC)", "filled"), product_value(other, "Конфиг:Род тока", rules_path)[:2])
            self.assertEqual(("230", "filled"), product_value(other, "Конфиг:Напряжение, В", rules_path)[:2])


if __name__ == "__main__":
    unittest.main()


class NameCandidateTests(unittest.TestCase):
    ATTRIBUTES = [
        {
            "Номинальное напряжение изоляции, В": "415",
            "Напряжение питания лампы, В": "220",
            "Вторичное напряжение 1, В": "24",
            "Первичное напряжение 1, В": "380",
            "Количество замыкающих контактов": "1",
            "Количество размыкающих контактов": "0",
            "Количество фаз": "3",
            "Номинальное напряжение питания цепи управления Us перемен. тока АС при 50 Гц, В": "230",
            "Номинальное напряжение питания цепи управления Us постоянного тока DC, В": "24",
            "Паспорт устройства": "E-com/NB1L - Паспорт устройства.pdf",
            "Номинальная реактивная мощность при 400 В, 50 Гц, кВАр": "10",
            "С подсветкой": "Да",
            "Тип присоединения главной электрической цепи": "Винтовое соединение",
        }
    ]

    def names(self, name: str) -> list[str]:
        from chint_etm_mdm.value_matching import name_candidates

        return [item[0] for item in name_candidates(name, self.ATTRIBUTES)]

    def test_good_candidates(self) -> None:
        self.assertEqual(["Напряжение питания лампы, В"], self.names("Напряжение подсветки, В"))
        self.assertEqual(["Вторичное напряжение 1, В"], self.names("Номинальное напряжение основной вторичной обмотки, В"))
        self.assertEqual(["Количество размыкающих контактов"], self.names("Количество НЗ контактов"))
        self.assertEqual(
            "Номинальное напряжение питания цепи управления Us постоянного тока DC, В",
            self.names("Номинальное напряжение питания постоянного тока, В")[0],
        )

    def test_false_friends_are_filtered(self) -> None:
        self.assertEqual([], self.names("Коммутируемое напряжение, В"))
        self.assertEqual([], self.names("Количество аналоговых входов"))
        self.assertEqual([], self.names("Тип устройства"))
        self.assertEqual([], self.names("Номинальная мощность электродвигателя, кВт"))
        self.assertEqual(["С подсветкой"], self.names("Наличие подсветки"))


class NumericCharacteristicTests(unittest.TestCase):
    BACKLIGHT = Characteristic("Напряжение подсветки, В", formats=["Числовая"], values=["AC 220", "AC/DC 24"])

    def test_numeric_list_is_only_an_example(self) -> None:
        voltage = Characteristic("Напряжение, В", formats=["Числовая", "Текстовая"], values=["12", "24", "220/380"])
        self.assertEqual(("230", "open"), (match_value("230 В", voltage).value, match_value("230 В", voltage).kind))

    def test_current_kind_is_added(self) -> None:
        self.assertEqual("AC/DC 230", match_value("230", self.BACKLIGHT, current_kind="AC/DC").value)
        self.assertEqual("none", match_value("230", self.BACKLIGHT).kind)

    def test_kind_and_voltage_from_product_name(self) -> None:
        from chint_etm_mdm.export_templates import load_partner_template

        template = load_partner_template("etm")
        product = ProductRecord(article="B", name="Кнопка с подсветкой NP2 AC/DC230В", attributes={"Род тока": "AC"})
        self.assertEqual(("230", "AC/DC", "product_name_current"), template.with_current_kind(product, "220", "db"))
        plain = ProductRecord(article="L", name="Лампа ND16", attributes={"Род тока": "Переменный ток (AC)"})
        self.assertEqual(("220", "AC", "db"), template.with_current_kind(plain, "220", "db"))

    def test_range_is_written_with_a_dash(self) -> None:
        from chint_etm_mdm.value_matching import clean_open_value

        self.assertEqual("380-415", clean_open_value("Номинальное напряжение питания переменного тока 50 Гц, В", "380…415"))
