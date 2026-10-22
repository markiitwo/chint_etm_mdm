import subprocess
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from chint_etm_mdm import etim_importer
from chint_etm_mdm.excel_files import ENCRYPTED_MARKER, OLE_MAGIC, file_kind, find_soffice, readable_xlsx


class ExcelFilesTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def write(self, name: str, data: bytes) -> Path:
        path = self.root / name
        path.write_bytes(data)
        return path

    def test_kinds_and_messages(self) -> None:
        book = self.root / "ok.xlsx"
        Workbook().save(book)
        self.assertEqual(book, readable_xlsx(book))
        cases = {
            "empty.xlsx": (b"", "empty", "пустой"),
            "locked.xlsx": (OLE_MAGIC + b"\0" * 100 + ENCRYPTED_MARKER, "encrypted", "защищён паролем"),
            "old.xlsx": (OLE_MAGIC + b"\0" * 100, "ole", "старом формате"),
            "page.xlsx": (b"\xef\xbb\xbf<html><table></table></html>", "html", "старом формате"),
            "junk.xlsx": (b"hello", "unknown", "не похож"),
        }
        for name, (data, kind, message) in cases.items():
            path = self.write(name, data)
            self.assertEqual(kind, file_kind(path))
            with self.assertRaises(ValueError) as caught:
                readable_xlsx(path, "ETIM-файл", soffice="")
            self.assertIn(message, str(caught.exception))
            self.assertIn(f"ETIM-файл «{name}»", str(caught.exception))

    def test_unreadable_etim_file_makes_no_backup(self) -> None:
        locked = self.write("ETIM.xlsx", OLE_MAGIC + b"\0" * 100 + ENCRYPTED_MARKER)
        database = self.write("db.sqlite", b"")
        with patch("chint_etm_mdm.etim_importer.backup_database") as backup:
            with self.assertRaisesRegex(ValueError, "защищён паролем"):
                etim_importer.import_etim_workbook(locked, database)
        backup.assert_not_called()

    @unittest.skipUnless(find_soffice(), "LibreOffice is not installed")
    def test_old_excel_is_converted(self) -> None:
        book = self.root / "src.xlsx"
        workbook = Workbook()
        workbook.active.append(["Артикул", "Ширина, мм"])
        workbook.active.append(["A-1", 45])
        workbook.save(book)
        subprocess.run([find_soffice(), "--headless", "--convert-to", "xls", "--outdir", str(self.root), str(book)], check=True, capture_output=True, timeout=300)
        old = self.root / "ETIM.xlsx"  # an old book under the new extension, as Excel's compatibility mode leaves it
        (self.root / "src.xls").rename(old)
        self.assertEqual("ole", file_kind(old))
        readable = readable_xlsx(old, out_dir=self.root / "converted")
        rows = list(load_workbook(readable, read_only=True).active.iter_rows(values_only=True))
        self.assertEqual([("Артикул", "Ширина, мм"), ("A-1", 45)], rows)


if __name__ == "__main__":
    unittest.main()
