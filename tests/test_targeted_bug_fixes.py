from __future__ import annotations

import io
import os
import ssl
import sqlite3
import tempfile
import unittest
import urllib.error
import zipfile
from contextlib import closing
from email.message import Message
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from chint_etm_mdm import db, etim_importer
from chint_etm_mdm.db import ProductRecord
from chint_etm_mdm.filler import fill_xlsx_template
from chint_etm_mdm.price_importer import (
    PRICE_SHEETS,
    download_price_file,
    find_latest_price_url,
    filename_from_response,
    read_sheet_rows,
    read_price_workbook,
)


class FakeResponse(io.BytesIO):
    def __init__(self, data: bytes, content_disposition: str = "", content_length: int | None = None) -> None:
        super().__init__(data)
        self.headers = Message()
        if content_disposition:
            self.headers["Content-Disposition"] = content_disposition
        if content_length is not None:
            self.headers["Content-Length"] = str(content_length)

    def __enter__(self):
        return self

    def __exit__(self, *_args):
        self.close()


def workbook_bytes() -> bytes:
    workbook = Workbook()
    workbook.active.title = "Sheet1"
    workbook.active.append(["value"])
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def price_workbook(path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name in PRICE_SHEETS:
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(["Артикул", "Наименование"])
        sheet.append(["A-1", "Товар"])
    workbook.save(path)
    workbook.close()


def rewrite_relationship_targets(path: Path, absolute: bool) -> None:
    replacement = path.with_suffix(".replacement")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(replacement, "w") as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "xl/_rels/workbook.xml.rels":
                if absolute:
                    data = data.replace(b'Target="xl/worksheets/', b'Target="/xl/worksheets/')
                else:
                    data = data.replace(b'Target="/xl/worksheets/', b'Target="worksheets/')
            target.writestr(info, data)
    os.replace(replacement, path)


class PricePackageTests(unittest.TestCase):
    def test_reads_relative_and_absolute_opc_sheet_targets(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            for absolute in (False, True):
                path = Path(tmp) / f"price-{absolute}.xlsx"
                price_workbook(path)
                rewrite_relationship_targets(path, absolute)
                rows = read_price_workbook(path)
                self.assertEqual("A-1", rows["Price-list"][1][0])
                self.assertEqual("Товар", rows["Price-list"][1][1])

    def test_price_page_retries_with_bundled_certificates(self) -> None:
        page = b'<a href="/upload/Price-list-CHINT.xlsx">\xd0\x9f\xd1\x80\xd0\xb0\xd0\xb9\xd1\x81</a>'
        untrusted = urllib.error.URLError(ssl.SSLCertVerificationError("self-signed certificate in certificate chain"))
        with patch("urllib.request.urlopen", side_effect=[untrusted, FakeResponse(page)]) as opened:
            url = find_latest_price_url("https://ensmas.ru/x/")
        self.assertEqual("https://ensmas.ru/upload/Price-list-CHINT.xlsx", url)
        self.assertIsInstance(opened.call_args_list[1].kwargs["context"], ssl.SSLContext)

    def test_download_validates_before_publish_and_uses_unique_name(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            downloads = Path(tmp)
            existing = downloads / "price.xlsx"
            existing.write_bytes(b"old")
            disposition = "attachment; filename=ignored.xlsx; filename*=UTF-8''price.xlsx"
            with patch("urllib.request.urlopen", return_value=FakeResponse(workbook_bytes(), disposition)):
                result = download_price_file("https://example.test/file", downloads)
            self.assertEqual("price_1.xlsx", result.name)
            self.assertEqual(b"old", existing.read_bytes())
            self.assertTrue(zipfile.is_zipfile(result))
            self.assertEqual([], list(downloads.glob("*.part")))

    def test_download_rejects_html_and_cleans_temporary_file(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            downloads = Path(tmp)
            with patch("urllib.request.urlopen", return_value=FakeResponse(b"<html>error</html>")):
                with self.assertRaisesRegex(ValueError, "корректной книгой Excel"):
                    download_price_file("http://example.test/price.xlsx", downloads)
            self.assertEqual([], list(downloads.iterdir()))

    def test_download_rejects_short_response_and_cleans_temporary_file(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            downloads = Path(tmp)
            data = workbook_bytes()
            with patch("urllib.request.urlopen", return_value=FakeResponse(data, content_length=len(data) + 1)):
                with self.assertRaisesRegex(ValueError, "короче"):
                    download_price_file("https://example.test/price.xlsx", downloads)
            self.assertEqual([], list(downloads.iterdir()))

    def test_download_enforces_size_limit_with_and_without_content_length(self) -> None:
        data = workbook_bytes()
        for declared_length in (len(data), None):
            with self.subTest(content_length=declared_length), tempfile.TemporaryDirectory() as tmp:
                downloads = Path(tmp)
                response = FakeResponse(data, content_length=declared_length)
                with (
                    patch("urllib.request.urlopen", return_value=response),
                    patch("chint_etm_mdm.price_importer.MAX_PRICE_DOWNLOAD_BYTES", len(data) - 1),
                ):
                    with self.assertRaisesRegex(ValueError, "допустимый размер"):
                        download_price_file("https://example.test/price.xlsx", downloads)
                self.assertEqual([], list(downloads.iterdir()))

    def test_publish_failure_cleans_download_without_changing_existing_file(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            downloads = Path(tmp)
            existing = downloads / "price.xlsx"
            existing.write_bytes(b"previous download")
            with (
                patch("urllib.request.urlopen", return_value=FakeResponse(workbook_bytes())),
                patch("chint_etm_mdm.price_importer.publish_unique_path", side_effect=PermissionError("locked")),
            ):
                with self.assertRaises(PermissionError):
                    download_price_file("https://example.test/price.xlsx", downloads)
            self.assertEqual([existing], list(downloads.iterdir()))
            self.assertEqual(b"previous download", existing.read_bytes())

    def test_parser_rejects_corrupt_zip_member(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "price.xlsx"
            price_workbook(path)
            stored = Path(tmp) / "stored.xlsx"
            with zipfile.ZipFile(path) as source, zipfile.ZipFile(stored, "w", compression=zipfile.ZIP_STORED) as target:
                for info in source.infolist():
                    target.writestr(info.filename, source.read(info.filename), compress_type=zipfile.ZIP_STORED)
            data = bytearray(stored.read_bytes())
            marker = data.find(b"<worksheet")
            self.assertNotEqual(-1, marker)
            data[marker] ^= 1
            stored.write_bytes(data)
            downloads = Path(tmp) / "downloads"
            with patch("urllib.request.urlopen", return_value=FakeResponse(stored.read_bytes())):
                with self.assertRaises((zipfile.BadZipFile, ValueError)):
                    download_price_file("https://example.test/price.xlsx", downloads)
            self.assertEqual([], list(downloads.iterdir()))

    def test_reads_shared_and_inline_strings_from_real_package(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "strings.xlsx"
            workbook = Workbook()
            workbook.active.title = "Data"
            workbook.active.append(["inline"])
            workbook.save(path)
            workbook.close()
            replacement = path.with_suffix(".replacement")
            with zipfile.ZipFile(path) as source, zipfile.ZipFile(replacement, "w") as target:
                for info in source.infolist():
                    data = source.read(info.filename)
                    if info.filename == "xl/worksheets/sheet1.xml":
                        data = data.replace(b't="inlineStr"><is><t>inline</t></is>', b't="s"><v>0</v>')
                    target.writestr(info, data)
                target.writestr(
                    "xl/sharedStrings.xml",
                    b'<?xml version="1.0"?><sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><t>shared</t></si></sst>',
                )
            os.replace(replacement, path)
            self.assertEqual([["shared"]], read_sheet_rows(path, "Data"))

    def test_reports_missing_sheet_relationship(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "missing-rel.xlsx"
            price_workbook(path)
            replacement = path.with_suffix(".replacement")
            with zipfile.ZipFile(path) as source, zipfile.ZipFile(replacement, "w") as target:
                for info in source.infolist():
                    data = source.read(info.filename)
                    if info.filename == "xl/_rels/workbook.xml.rels":
                        data = data.replace(b'Id="rId1"', b'Id="removed"')
                    target.writestr(info, data)
            os.replace(replacement, path)
            with self.assertRaisesRegex(ValueError, "не найдена связь"):
                read_sheet_rows(path, PRICE_SHEETS[0])

    def test_download_cleans_partial_file_when_read_fails(self) -> None:
        class BrokenResponse(FakeResponse):
            def read(self, size=-1):
                if self.tell() > 0:
                    raise TimeoutError("interrupted")
                return super().read(10 if size < 0 else min(size, 10))

        with tempfile.TemporaryDirectory() as tmp:
            downloads = Path(tmp)
            with patch("urllib.request.urlopen", return_value=BrokenResponse(workbook_bytes())):
                with self.assertRaises(TimeoutError):
                    download_price_file("https://example.test/broken.xlsx", downloads)
            self.assertEqual([], list(downloads.iterdir()))

    def test_sanitizes_external_windows_filename(self) -> None:
        self.assertEqual("outside.xlsx", filename_from_response("https://x/f.xlsx", 'attachment; filename="..\\outside.xlsx"'))
        self.assertEqual("_CON.xlsx", filename_from_response("https://x/f.xlsx", "attachment; filename*=UTF-8''CON"))
        self.assertEqual("цена_09.xlsx", filename_from_response("https://x/f.xlsx", "attachment; filename=old.xlsx; filename*=UTF-8''%D1%86%D0%B5%D0%BD%D0%B0%3A09.xlsx"))


class TrackingConnection(sqlite3.Connection):
    closed = False

    def close(self) -> None:
        self.closed = True
        super().close()


class ConnectionLifetimeTests(unittest.TestCase):
    def test_db_reader_closes_connection_after_exception(self) -> None:
        conn = sqlite3.connect(":memory:", factory=TrackingConnection)
        with patch("chint_etm_mdm.db.connect", return_value=conn):
            with self.assertRaises(sqlite3.OperationalError):
                db.get_stats(Path("unused.sqlite"))
        self.assertTrue(conn.closed)

    def test_fetch_products_closes_connection_after_exception(self) -> None:
        conn = sqlite3.connect(":memory:", factory=TrackingConnection)
        with patch("chint_etm_mdm.db.connect", return_value=conn):
            with self.assertRaises(sqlite3.OperationalError):
                db.fetch_products(Path("unused.sqlite"), ["A-1"])
        self.assertTrue(conn.closed)

    def test_etim_writer_rolls_back_and_closes_after_exception(self) -> None:
        connections: list[TrackingConnection] = []

        def connect(path):
            connection = sqlite3.connect(path, factory=TrackingConnection)
            connections.append(connection)
            return connection

        def fail_after_write(connection, *_args):
            connection.execute("INSERT INTO marker VALUES ('uncommitted')")
            raise RuntimeError("boom")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "etim.xlsx"
            database = root / "db.sqlite"
            source.touch()
            with closing(sqlite3.connect(database)) as setup, setup:
                setup.execute("CREATE TABLE marker(value TEXT)")

            def ensure_tables(connection):
                connection.execute(
                    "CREATE TABLE IF NOT EXISTS dimension_imports (id INTEGER PRIMARY KEY, source_kind TEXT, local_file TEXT, workbook_name TEXT, notes TEXT)"
                )
            with (
                patch("chint_etm_mdm.etim_importer.connect_database", side_effect=connect),
                patch("chint_etm_mdm.etim_importer.ensure_etim_import_tables", side_effect=ensure_tables),
                patch("chint_etm_mdm.etim_importer.fetch_all_articles", return_value=set()),
                patch("chint_etm_mdm.etim_importer.scan_etim_workbook", return_value={}),
                patch("chint_etm_mdm.etim_importer.build_dimension_report_rows", return_value=(0, 0, [], [])),
                patch("chint_etm_mdm.etim_importer.insert_etim_data", side_effect=fail_after_write),
            ):
                with self.assertRaisesRegex(RuntimeError, "boom"):
                    etim_importer.import_etim_workbook(source, database, make_backup=False)
            with closing(sqlite3.connect(database)) as verify:
                self.assertEqual(0, verify.execute("SELECT COUNT(*) FROM marker").fetchone()[0])
        self.assertEqual(2, len(connections))
        self.assertTrue(all(connection.closed for connection in connections))

    def test_successful_etim_import_closes_both_connections(self) -> None:
        connections: list[TrackingConnection] = []

        def connect(path):
            connection = sqlite3.connect(path, factory=TrackingConnection)
            connections.append(connection)
            return connection

        def ensure_tables(connection):
            connection.execute(
                "CREATE TABLE IF NOT EXISTS dimension_imports (id INTEGER PRIMARY KEY, source_kind TEXT, local_file TEXT, workbook_name TEXT, notes TEXT)"
            )

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "etim.xlsx"
            database = root / "db.sqlite"
            source.touch()
            database.touch()
            with (
                patch("chint_etm_mdm.etim_importer.connect_database", side_effect=connect),
                patch("chint_etm_mdm.etim_importer.ensure_etim_import_tables", side_effect=ensure_tables),
                patch("chint_etm_mdm.etim_importer.fetch_all_articles", return_value=set()),
                patch("chint_etm_mdm.etim_importer.scan_etim_workbook", return_value={}),
                patch("chint_etm_mdm.etim_importer.build_dimension_report_rows", return_value=(0, 0, [], [])),
                patch("chint_etm_mdm.etim_importer.insert_etim_data", return_value=(0, 0)),
                patch("chint_etm_mdm.etim_importer.write_etim_report"),
            ):
                result = etim_importer.import_etim_workbook(source, database, make_backup=False)
            self.assertEqual(1, result.import_id)
            os.replace(database, root / "renamed.sqlite")
        self.assertEqual(2, len(connections))
        self.assertTrue(all(connection.closed for connection in connections))


    def test_applying_etim_decisions_commits_and_closes_connection(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            database = root / "db.sqlite"
            with closing(sqlite3.connect(database)) as setup, setup:
                setup.execute("CREATE TABLE products(article TEXT PRIMARY KEY)")
                setup.execute("INSERT INTO products VALUES ('A-1')")
                etim_importer.ensure_etim_import_tables(setup)
                setup.execute(
                    "INSERT INTO product_dimensions_raw (import_id, article, source_kind, dimension_kind, length_mm) "
                    "VALUES (1, 'A-1', ?, 'length_mm', 123)",
                    (etim_importer.SOURCE_KIND,),
                )
            report = root / "decisions.xlsx"
            workbook = Workbook()
            workbook.active.title = "Конфликты"
            workbook.active.append(["Решение", "Артикул", "Поле", "Сейчас в базе", "Пришло из ETIM"])
            workbook.active.append([etim_importer.ACCEPT_ETIM_DECISION, "A-1", "Длина, мм", "", 123])
            workbook.save(report)
            workbook.close()

            connection = sqlite3.connect(database, factory=TrackingConnection)
            with patch("chint_etm_mdm.etim_importer.connect_database", return_value=connection):
                result = etim_importer.apply_etim_conflict_decisions(report, database)
            self.assertTrue(connection.closed)
            self.assertEqual(1, result.accepted_rows)
            self.assertTrue(result.backup_path.is_file())
            with closing(sqlite3.connect(database)) as verify:
                self.assertEqual(
                    123,
                    verify.execute(
                        "SELECT length_mm FROM product_dimensions_resolved WHERE article = 'A-1'"
                    ).fetchone()[0],
                )
            os.replace(database, root / "renamed.sqlite")


class EmptyOnlyTests(unittest.TestCase):
    def test_xlsx_preserves_zero_false_and_formula(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Артикул", "Вес, кг", "Название", "Полное название", "Краткое имя WMS"])
            for cell in sheet[1]:
                cell.fill = PatternFill(fill_type="solid", fgColor="FFFF00")
            sheet.append(["A-1", 0, False, "=1+1", None])
            workbook.save(template)
            workbook.close()

            product = ProductRecord(article="A-1", name="Новое", full_name="Полное", weight_kg=12)
            with patch("chint_etm_mdm.filler.fetch_products", return_value={"A-1": product}):
                result = fill_xlsx_template(root / "db.sqlite", template, root / "out")

            output = load_workbook(result.output_path, data_only=False)
            try:
                values = [output.active.cell(2, column).value for column in range(2, 6)]
            finally:
                output.close()
            self.assertEqual([0, False, "=1+1", "Новое"], values)
            original = load_workbook(template, data_only=False)
            try:
                self.assertEqual(0, original.active["B2"].value)
            finally:
                original.close()


if __name__ == "__main__":
    unittest.main()
