import datetime as dt
import json
import sqlite3
import unittest
from contextlib import closing
from email import message_from_bytes
from pathlib import Path

from openpyxl import Workbook, load_workbook

from chint_etm_mdm.pm_requests import (
    ANSWER,
    CONTACTS_NAME,
    HEADER_ROW,
    KEY,
    Need,
    answer_lines,
    answers_path,
    import_answers,
    is_request_file,
    load_assignment,
    manager_for,
    request_lines,
    write_requests,
)
import test_raek
from test_raek import DRIVE, cells


def make_assignment(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Закрепление ПМ"
    sheet.append(["Низковольтное", None, None, None, "Средневольтное", None, None])
    sheet.append(["Серия", "Ответственный ПМ", "Примечание", None, "Серия", "Ответственный ПМ", "Примечание"])
    sheet.append(["Преобразователи частоты", None, None, None, "Реле", None, None])
    sheet.append(["NVF7", "Левин И.", None, None, "NJDC-17", "Калашникова Н./Левин И.", None])
    sheet.append(["NB1L", "Власова С.", "Выводятся из асортимента", None, None, None, None])
    workbook.save(path)


class AssignmentTests(unittest.TestCase):
    def test_series_to_manager(self) -> None:
        import tempfile

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "Закрепление ПМ.xlsx"
            make_assignment(path)
            assignment = load_assignment(path)
        self.assertEqual(
            [("NVF7", "Левин И.", "Преобразователи частоты"), ("NB1L", "Власова С.", "Преобразователи частоты"), ("NJDC-17", "Калашникова Н./Левин И.", "Реле")],
            [(item.series, item.manager, item.line) for item in assignment],
        )
        self.assertEqual("Левин И.", manager_for("NVF7", assignment))
        self.assertEqual("Власова С.", manager_for("NB1L-20", assignment))  # the longest series it starts with
        self.assertEqual("Калашникова Н./Левин И.", manager_for("njdc 17", assignment))
        self.assertEqual("", manager_for("NXB", assignment))


class RequestRoundTripTests(unittest.TestCase):
    setUp = test_raek.RaekTests.setUp
    tearDown = test_raek.RaekTests.tearDown
    fill = test_raek.RaekTests.fill
    member = test_raek.RaekTests.member

    def test_request_answer_and_fill(self) -> None:
        with closing(sqlite3.connect(self.db)) as conn, conn:
            conn.execute("DELETE FROM product_attribute_values WHERE article = ? AND attribute_name = 'Количество входных фаз'", (DRIVE[1],))
        rejects = ["Диапазон входной частоты, Гц", "Диапазон выходной частоты, Гц", "Напряжение сети, В"]
        (self.rules.parent / "raek_learned.json").write_text(json.dumps({"name_rejects": {"EF007117": rejects}}), encoding="utf-8")
        first = self.fill()
        self.assertIn(
            Need(DRIVE[1], DRIVE[2], "NVF7", "EF007117", "Количество вход. фаз", "A", ("1", "3", "Прочее")),
            first.needs,
        )
        assignment_file = self.root / "Закрепление ПМ.xlsx"
        make_assignment(assignment_file)
        out = self.root / "requests"
        today = dt.date(2026, 10, 22)
        result = write_requests(first.needs, load_assignment(assignment_file), self.root, out, "Заполняем файлы для РАЭК", today)
        levin = next(item for item in result.files if item.manager == "Левин И.")
        self.assertTrue(is_request_file(levin.path))
        lists = load_workbook(levin.path)["Запрос"].data_validations.dataValidation
        self.assertTrue(lists and all(not item.formula1.startswith("=") for item in lists))  # LibreOffice ignores "=…"
        self.assertIn("Левин И.", result.missing_emails)
        self.assertTrue((self.root / CONTACTS_NAME).exists())
        draft = message_from_bytes(levin.draft.read_bytes())
        self.assertEqual("1", draft["X-Unsent"])
        self.assertIn("Запрос продакту Левин И.", "\n".join(request_lines(result)))

        # Asking again soon after does not repeat open questions.
        again = write_requests(first.needs, load_assignment(assignment_file), self.root, out, "Заполняем файлы для РАЭК", today)
        self.assertEqual([], again.files)
        self.assertEqual(len(first.needs), again.skipped_open)

        # Asked to make the files anew (lost, broken, emails added later): every open question again.
        anew = write_requests(first.needs, load_assignment(assignment_file), self.root, out, "Заполняем файлы для РАЭК", today, again=True)
        self.assertEqual(len(result.files), len(anew.files))
        self.assertEqual(0, anew.skipped_open)
        self.assertEqual([], write_requests(first.needs, load_assignment(assignment_file), self.root, out, "Заполняем файлы для РАЭК", today).files)

        # The manager answers in their own words; the program matches the partner's list.
        workbook = load_workbook(levin.path)
        sheet = workbook["Запрос"]
        for row in range(HEADER_ROW + 1, sheet.max_row + 1):
            if sheet.cell(row, KEY).value == "EF007117":
                sheet.cell(row, ANSWER).value = 3
        workbook.save(levin.path)
        answers = import_answers(levin.path, self.root)
        self.assertEqual(1, answers.accepted, answer_lines(answers))
        stored = json.loads(answers_path(self.root).read_text(encoding="utf-8"))["answers"]
        self.assertEqual([(DRIVE[1], "EF007117", "3")], [(a["article"], a["key"], a["value"]) for a in stored])

        from chint_etm_mdm.raek import fill_raek_zip

        filled = fill_raek_zip(self.db, self.zip, self.root / "out", self.codes, self.rules, answers_path=answers_path(self.root))
        drive = cells(self.member(filled, "ChINT (CHINT)_required_features.xlsx"), "EC001857")
        self.assertEqual("3", drive[3][7])
        self.assertNotIn("EF007117", [need.key for need in filled.needs])

    def test_wrong_answers_are_not_kept(self) -> None:
        needs = [
            Need(DRIVE[1], DRIVE[2], "NVF7", "кодтнвэд", "Код ТН ВЭД"),
            Need(DRIVE[1], DRIVE[2], "NVF7", "EF000040", "Высота, мм", "N"),
            Need("", "Характеристика «Высота, мм» — в базе «Высота корпуса, мм»", "NVF7", "name:EF000040:Высота корпуса, мм", "Это одна и та же характеристика?"),
        ]
        result = write_requests(needs, [], self.root, self.root / "requests", "Проба", dt.date(2026, 10, 22))
        path = result.files[0].path
        workbook = load_workbook(path)
        sheet = workbook["Запрос"]
        answers = {"кодтнвэд": "8504 31", "EF000040": "высокий", "name:EF000040:Высота корпуса, мм": "да"}
        for row in range(HEADER_ROW + 1, sheet.max_row + 1):
            sheet.cell(row, ANSWER).value = answers[sheet.cell(row, KEY).value]
        workbook.save(path)
        imported = import_answers(path, self.root)
        self.assertEqual((0, 1), (imported.accepted, imported.names_kept))
        self.assertEqual(["нужно 10 цифр", "нужно число"], [item[3] for item in imported.rejected])
        learned = json.loads((self.root / "rules" / "raek_learned.json").read_text(encoding="utf-8"))
        self.assertEqual({"EF000040": ["Высота корпуса, мм"]}, learned["name_aliases"])


if __name__ == "__main__":
    unittest.main()
